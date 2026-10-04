"""Import seeds/foods_catalog_perfect_v1.xlsx (Foods_Catalog_Master) into foods."""

from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.entities import Food, FoodAlias, FoodCategory

MASTER_SHEET = "Foods_Catalog_Master"
CURATED_SHEET = "Dishes_Curated"

# Excel column indices (0-based) for Foods_Catalog_Master.
COL_ID = 0
COL_SLUG = 1
COL_NAME_VI = 2
COL_NAME_EN = 3
COL_GROUP = 4
COL_KIND = 5
COL_PREP = 6
COL_SERVING_SIZE = 7
COL_SERVING_GRAMS = 8
COL_CALORIES = 9
COL_PROTEIN = 10
COL_CARBS = 11
COL_FAT = 12
COL_FIBER = 13
COL_KCAL_100 = 14
COL_PROTEIN_100 = 15
COL_CARBS_100 = 16
COL_FAT_100 = 17
COL_FIBER_100 = 18
COL_YIELD = 19
COL_VERIFIED = 20
COL_AI_ELIGIBLE = 21
COL_DEFAULT_AI = 22
COL_ALIASES = 23
COL_SOURCE = 24

ZERO_MACRO_OK = frozenset({"nuoc-loc", "muoi"})


@dataclass
class PerfectFoodRow:
    id: int | None
    slug: str
    name_vi: str
    name_en: str | None
    group_name: str
    food_kind: str
    prep_state: str | None
    serving_size: str
    serving_grams: float
    calories: float
    protein_g: float
    carbs_g: float
    fat_g: float
    fiber_g: float
    kcal_100g: float
    protein_100g: float
    carbs_100g: float
    fat_100g: float
    fiber_100g: float
    yield_factor: float | None
    is_verified: bool
    ai_eligible: bool
    default_for_ai: bool
    aliases: list[str]
    source_ref: str | None


@dataclass
class CuratedDishRow:
    id: int
    name_vi: str
    serving_size: str
    serving_grams: float
    calories: float
    protein_g: float
    carbs_g: float
    fat_g: float
    kcal_100g: float
    ai_note: str | None


@dataclass
class ImportStats:
    updated: int = 0
    skipped_missing: int = 0
    aliases_added: int = 0
    missing_slugs: list[str] | None = None

    def __post_init__(self) -> None:
        if self.missing_slugs is None:
            self.missing_slugs = []


def default_perfect_xlsx_path(project_root: Path) -> Path:
    return project_root / "seeds" / "foods_catalog_perfect_v1.xlsx"


def _as_float(value: Any, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _as_optional_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _as_bool(value: Any, default: bool = False) -> bool:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    text = str(value).strip().lower()
    if text in {"1", "true", "yes", "y", "có", "co"}:
        return True
    if text in {"0", "false", "no", "n", "không", "khong"}:
        return False
    return default


def _parse_aliases(value: Any) -> list[str]:
    if value is None or value == "":
        return []
    if isinstance(value, list):
        return [str(x).strip() for x in value if str(x).strip()]
    parts = []
    for chunk in str(value).replace(";", ",").split(","):
        alias = chunk.strip()
        if alias:
            parts.append(alias)
    return parts


def load_perfect_master_rows(xlsx_path: Path) -> list[PerfectFoodRow]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        if MASTER_SHEET not in wb.sheetnames:
            raise FileNotFoundError(f"missing sheet {MASTER_SHEET} in {xlsx_path}")
        ws = wb[MASTER_SHEET]
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if not header or str(header[COL_SLUG] or "").strip().lower() != "slug":
            raise ValueError(f"unexpected header in {MASTER_SHEET}: {header}")

        out: list[PerfectFoodRow] = []
        for raw in rows_iter:
            if not raw or not raw[COL_SLUG]:
                continue
            slug = str(raw[COL_SLUG]).strip()
            if not slug:
                continue
            food_id = None
            if raw[COL_ID] is not None and str(raw[COL_ID]).strip() != "":
                try:
                    food_id = int(raw[COL_ID])
                except (TypeError, ValueError):
                    food_id = None
            out.append(
                PerfectFoodRow(
                    id=food_id,
                    slug=slug,
                    name_vi=str(raw[COL_NAME_VI] or slug).strip(),
                    name_en=(str(raw[COL_NAME_EN]).strip() if raw[COL_NAME_EN] else None),
                    group_name=str(raw[COL_GROUP] or "").strip(),
                    food_kind=str(raw[COL_KIND] or "ingredient").strip() or "ingredient",
                    prep_state=(str(raw[COL_PREP]).strip() if raw[COL_PREP] else None),
                    serving_size=str(raw[COL_SERVING_SIZE] or "100g").strip() or "100g",
                    serving_grams=_as_float(raw[COL_SERVING_GRAMS], 100.0),
                    calories=_as_float(raw[COL_CALORIES]),
                    protein_g=_as_float(raw[COL_PROTEIN]),
                    carbs_g=_as_float(raw[COL_CARBS]),
                    fat_g=_as_float(raw[COL_FAT]),
                    fiber_g=_as_float(raw[COL_FIBER]),
                    kcal_100g=_as_float(raw[COL_KCAL_100]),
                    protein_100g=_as_float(raw[COL_PROTEIN_100]),
                    carbs_100g=_as_float(raw[COL_CARBS_100]),
                    fat_100g=_as_float(raw[COL_FAT_100]),
                    fiber_100g=_as_float(raw[COL_FIBER_100]),
                    yield_factor=_as_optional_float(raw[COL_YIELD]),
                    is_verified=_as_bool(raw[COL_VERIFIED], False),
                    ai_eligible=_as_bool(raw[COL_AI_ELIGIBLE], True),
                    default_for_ai=_as_bool(raw[COL_DEFAULT_AI], False),
                    aliases=_parse_aliases(raw[COL_ALIASES]),
                    source_ref=(str(raw[COL_SOURCE]).strip() if raw[COL_SOURCE] else None),
                )
            )
        return out
    finally:
        wb.close()


def load_curated_dish_rows(xlsx_path: Path) -> list[CuratedDishRow]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        if CURATED_SHEET not in wb.sheetnames:
            return []
        ws = wb[CURATED_SHEET]
        started = False
        out: list[CuratedDishRow] = []
        for raw in ws.iter_rows(values_only=True):
            if not raw:
                continue
            if raw[1] == "ID":
                started = True
                continue
            if not started or raw[1] is None:
                continue
            try:
                food_id = int(raw[1])
            except (TypeError, ValueError):
                continue
            out.append(
                CuratedDishRow(
                    id=food_id,
                    name_vi=str(raw[2] or "").strip(),
                    serving_size=str(raw[3] or "").strip(),
                    serving_grams=_as_float(raw[4]),
                    calories=_as_float(raw[5]),
                    protein_g=_as_float(raw[6]),
                    carbs_g=_as_float(raw[7]),
                    fat_g=_as_float(raw[8]),
                    kcal_100g=_as_float(raw[9]),
                    ai_note=(str(raw[10]).strip() if raw[10] else None),
                )
            )
        return out
    finally:
        wb.close()


def apply_perfect_catalog(
    db: Session,
    rows: list[PerfectFoodRow],
    *,
    category_by_name: dict[str, FoodCategory] | None = None,
) -> ImportStats:
    """Update existing public foods from perfect Master rows. Does not change status."""
    if category_by_name is None:
        category_by_name = {
            c.name_vi: c for c in db.query(FoodCategory).all() if c.name_vi
        }

    by_slug = {
        f.slug: f
        for f in db.query(Food).filter(Food.owner_user_id.is_(None)).all()
    }
    by_id = {f.id: f for f in by_slug.values()}

    purged: set[str] = set()
    purged_path = Path(__file__).resolve().parents[3] / "seeds" / "purged_food_slugs.json"
    if purged_path.is_file():
        try:
            raw = json.loads(purged_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            raw = []
        if isinstance(raw, dict):
            raw = raw.get("slugs") or []
        if isinstance(raw, list):
            purged = {str(s).strip().lower() for s in raw if str(s).strip()}

    stats = ImportStats()
    for item in rows:
        if item.slug.strip().lower() in purged:
            continue
        row = by_slug.get(item.slug)
        if row is None and item.id is not None:
            row = by_id.get(item.id)
        if row is None:
            stats.skipped_missing += 1
            stats.missing_slugs.append(item.slug)
            continue

        if item.food_kind == "dish" or (row.food_kind or "") == "dish":
            continue

        cat = category_by_name.get(item.group_name)
        if cat is not None:
            row.category_id = cat.id

        row.name_vi = item.name_vi
        row.name_en = item.name_en
        row.food_kind = item.food_kind
        row.prep_state = item.prep_state
        row.serving_size = item.serving_size
        row.serving_grams = item.serving_grams
        row.calories = item.calories
        row.protein_g = item.protein_g
        row.carbs_g = item.carbs_g
        row.fat_g = item.fat_g
        row.fiber_g = item.fiber_g
        row.kcal_100g = item.kcal_100g
        row.protein_100g = item.protein_100g
        row.carbs_100g = item.carbs_100g
        row.fat_100g = item.fat_100g
        row.fiber_100g = item.fiber_100g
        row.yield_factor = item.yield_factor
        row.is_verified = item.is_verified
        row.ai_eligible = item.ai_eligible
        row.default_for_ai = item.default_for_ai
        if item.source_ref:
            row.source_ref = item.source_ref
        if item.is_verified:
            row.confidence = "reference"
        elif item.food_kind == "dish":
            row.confidence = "estimated"

        if item.aliases:
            existing = {
                a.alias.casefold(): a
                for a in db.query(FoodAlias).filter(FoodAlias.food_id == row.id).all()
            }
            for alias in item.aliases:
                key = alias.casefold()
                if key in existing:
                    continue
                if key == (row.name_vi or "").casefold():
                    continue
                db.add(FoodAlias(food_id=row.id, alias=alias))
                existing[key] = None  # type: ignore[assignment]
                stats.aliases_added += 1

        stats.updated += 1

    db.commit()
    return stats


def dish_rows_for_seed(rows: list[PerfectFoodRow]) -> list[dict[str, Any]]:
    """Build traditional-dish seed payloads from Master dish rows (preserve extra keys later)."""
    dishes: list[dict[str, Any]] = []
    for item in rows:
        if item.food_kind != "dish":
            continue
        dishes.append(
            {
                "slug": item.slug,
                "name_vi": item.name_vi,
                "name_en": item.name_en,
                "serving_size": item.serving_size,
                "serving_grams": item.serving_grams,
                "calories": item.calories,
                "protein_g": item.protein_g,
                "carbs_g": item.carbs_g,
                "fat_g": item.fat_g,
                "fiber_g": item.fiber_g,
                "prep_state": item.prep_state or "cooked",
                "is_verified": item.is_verified,
                "is_common": True,
                "source_ref": item.source_ref or "excel+specialty:viet-nam-dishes",
                "macro_roles": ["carb", "protein"],
                "ai_priority": 25,
            }
        )
    return dishes
