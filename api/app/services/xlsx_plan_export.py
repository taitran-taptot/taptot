"""Simple Excel workout planner: start date + sessions with exercises and meals."""

from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

SECTION_LABEL = {
    "warmup": "Khởi động",
    "main": "Tập chính",
    "cooldown": "Giãn cơ",
    "cardio": "Cardio",
}
MEAL_LABEL = {
    "breakfast": "Sáng",
    "lunch": "Trưa",
    "dinner": "Tối",
    "snack": "Phụ",
}
ALWAYS_SECTIONS = ("warmup", "main", "cooldown", "cardio")
WEEK_RE = re.compile(r"[Tt]u[aàâầ]n\s+(\d+)")

GREEN = "059669"
GREEN_DARK = "047857"
INK = "334155"
MUTED = "64748B"
YELLOW = "FEF3C7"
ROW_ALT = "F0FDFA"
WHITE = "FFFFFF"
LINE = "D1D5DB"

thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(size=11, bold=False, color=INK, name="Calibri") -> Font:
    return Font(name=name, size=size, bold=bold, color=color)


def _parse_start(options: dict[str, Any] | None, detail: dict[str, Any]) -> date:
    opts = options or {}
    raw = opts.get("start_date") or detail.get("start_date")
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    if isinstance(raw, str) and raw.strip():
        try:
            return date.fromisoformat(raw.strip()[:10])
        except ValueError:
            pass
    return date.today()


def _week_of(day: dict[str, Any]) -> int | None:
    m = WEEK_RE.search(str(day.get("title_vi") or ""))
    return int(m.group(1)) if m else None


def _group_weeks(days: list[dict[str, Any]]) -> list[tuple[int, list[dict[str, Any]]]]:
    if not days:
        return []
    titled = all(_week_of(d) is not None for d in days)
    if not titled:
        return [(1, list(days))]
    buckets: dict[int, list[dict[str, Any]]] = {}
    for day in days:
        w = _week_of(day) or 1
        buckets.setdefault(w, []).append(day)
    return [(w, buckets[w]) for w in sorted(buckets)]


def _session_offsets(count: int) -> list[int]:
    presets = {
        1: [0],
        2: [0, 3],
        3: [0, 2, 4],
        4: [0, 2, 4, 6],
        5: [0, 1, 3, 4, 6],
        6: [0, 1, 2, 4, 5, 6],
        7: [0, 1, 2, 3, 4, 5, 6],
    }
    return presets.get(count, list(range(max(1, count))))


def _exercises_by_section(exercises: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = {sec: [] for sec in ALWAYS_SECTIONS}
    for ex in exercises or []:
        sec = ex.get("section") or "main"
        grouped.setdefault(sec, []).append(ex)
    return grouped


def _meals_by_type(meals: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = {
        "breakfast": [],
        "lunch": [],
        "dinner": [],
        "snack": [],
    }
    for meal in meals or []:
        mt = meal.get("meal_type") or "snack"
        grouped.setdefault(mt, []).append(meal)
    return grouped


def _work_label(ex: dict[str, Any]) -> str:
    raw = ex.get("reps")
    if raw is None or raw == "":
        return "—"
    return str(raw)


def _rest_label(ex: dict[str, Any]) -> str:
    sec = int(ex.get("rest_seconds") or 0)
    if sec <= 0:
        return "—"
    if sec >= 60 and sec % 60 == 0:
        return f"{sec // 60} phút"
    return f"{sec} giây"


def _weekday_formula(cell: str) -> str:
    return (
        f'CHOOSE(WEEKDAY({cell},2),'
        '"Thứ 2","Thứ 3","Thứ 4","Thứ 5","Thứ 6","Thứ 7","Chủ nhật")'
    )


def _style_header_row(ws: Worksheet, row: int, cols: int, fill: PatternFill, font: Font) -> None:
    for col in range(1, cols + 1):
        cell = ws.cell(row, col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin


def build_plan_xlsx(detail: dict[str, Any], options: dict[str, Any] | None = None) -> bytes:
    opts = options or {}
    start = _parse_start(opts, detail)
    days = list(detail.get("days") or [])
    weeks = _group_weeks(days)
    customer = (opts.get("customer_name") or "").strip()
    header_text = (opts.get("header_text") or "").strip()
    footer_text = (opts.get("footer_text") or "").strip()
    title = str(detail.get("title_vi") or "Lịch tập TAPTOT")

    wb = Workbook()
    ws_setup = wb.active
    ws_setup.title = "Huong dan"
    ws_overview = wb.create_sheet("Tong quan")
    ws_plan = wb.create_sheet("Lich tap")

    _write_setup(ws_setup, detail, title, customer, header_text, footer_text, start)
    date_cells = _write_overview(ws_overview, weeks)
    _write_plan(ws_plan, weeks, date_cells, footer_text)

    try:
        wb.defined_names.add(DefinedName(name="NgayBatDau", attr_text="'Huong dan'!$B$5"))
    except Exception:
        pass

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_setup(
    ws: Worksheet,
    detail: dict[str, Any],
    title: str,
    customer: str,
    header_text: str,
    footer_text: str,
    start: date,
) -> None:
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 18

    ws.merge_cells("A1:C1")
    ws["A1"] = "TAPTOT — Lịch tập"
    ws["A1"].font = _font(20, bold=True, color=WHITE)
    ws["A1"].fill = _fill(GREEN)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32

    ws["A3"] = "Tên lịch"
    ws["B3"] = title
    ws["B3"].font = _font(13, bold=True)

    ws["A4"] = "Khách hàng"
    ws["B4"] = customer or "—"

    ws["A5"] = "Ngày bắt đầu"
    ws["B5"] = start
    ws["B5"].number_format = "DD/MM/YYYY"
    ws["B5"].fill = _fill(YELLOW)
    ws["B5"].font = _font(12, bold=True)
    ws["B5"].border = thin
    ws["C5"] = "← Đổi ô này, mọi buổi sẽ tự chạy ngày"
    ws["C5"].font = _font(10, color=MUTED)

    ws["A6"] = "Hướng dẫn"
    ws["B6"] = (
        "1) Sửa Ngày bắt đầu (ô vàng). "
        "2) Nếu tập lệch ngày, sửa ô Ngày trên sheet Tong quan. "
        "3) Sheet Lich tap in bài tập + thực đơn đúng từng buổi."
    )
    ws["B6"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("B6:C6")
    ws.row_dimensions[6].height = 48

    kcal = detail.get("target_calories")
    ws["A8"] = "Mục tiêu calo"
    ws["B8"] = f"{kcal} kcal/ngày" if kcal else "—"
    ws["A9"] = "Đạm / Tinh bột / Béo"
    p, c, f = detail.get("target_protein_g"), detail.get("target_carbs_g"), detail.get("target_fat_g")
    if p is not None or c is not None or f is not None:
        ws["B9"] = (
            f"{p or 0:g}g  ·  {c or 0:g}g  ·  {f or 0:g}g"
        )
    else:
        ws["B9"] = "—"

    share = detail.get("share_url_path") or ""
    ws["A10"] = "Link xem lại"
    ws["B10"] = share or "—"

    if header_text:
        ws["A12"] = "Lời chào"
        ws["B12"] = header_text
        ws["B12"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells("B12:C13")

    if footer_text:
        ws["A15"] = "Ghi chú cuối"
        ws["B15"] = footer_text
        ws["B15"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells("B15:C16")

    desc = (detail.get("description_vi") or "").strip()
    if desc:
        ws["A18"] = "Mô tả"
        ws["B18"] = desc
        ws["B18"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells("B18:C20")

    for row in (3, 4, 8, 9, 10):
        ws.cell(row, 1).font = _font(10, bold=True, color=MUTED)


def _write_overview(
    ws: Worksheet,
    weeks: list[tuple[int, list[dict[str, Any]]]],
) -> list[str]:
    """Return Excel cell refs (e.g. 'Tong quan'!B3) for each session date, in order."""
    ws.freeze_panes = "A2"
    widths = {"A": 10, "B": 14, "C": 14, "D": 36, "E": 12, "F": 14}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    headers = ["STT", "Ngày", "Thứ", "Buổi", "Số bài", "Calo thực đơn"]
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
    _style_header_row(ws, 1, 6, _fill(GREEN), _font(11, bold=True, color=WHITE))

    date_cells: list[str] = []
    row = 2
    stt = 1
    for week_num, week_days in weeks:
        offsets = _session_offsets(len(week_days))
        for idx, day in enumerate(week_days):
            offset = (week_num - 1) * 7 + offsets[min(idx, len(offsets) - 1)]
            date_ref = f"B{row}"
            abs_ref = f"'Tong quan'!{date_ref}"
            date_cells.append(abs_ref)

            ws.cell(row, 1, stt).alignment = Alignment(horizontal="center")
            cell_date = ws.cell(row, 2)
            cell_date.value = f"='Huong dan'!$B$5+{offset}"
            cell_date.number_format = "DD/MM/YYYY"
            cell_date.fill = _fill(YELLOW)
            cell_date.font = _font(11, bold=True)
            ws.cell(row, 3, f"={_weekday_formula(date_ref)}")
            ws.cell(row, 4, day.get("title_vi") or f"Ngày {day.get('day_number') or stt}")
            ws.cell(row, 5, len(day.get("exercises") or [])).alignment = Alignment(horizontal="center")
            meals = day.get("meals") or []
            cals = 0
            for m in meals:
                try:
                    cals += float(m.get("calories") or 0)
                except (TypeError, ValueError):
                    pass
            cal_cell = ws.cell(row, 6, round(cals) if cals else None)
            cal_cell.alignment = Alignment(horizontal="center")

            for col in range(1, 7):
                ws.cell(row, col).border = thin
                if row % 2 == 0 and col != 2:
                    ws.cell(row, col).fill = _fill(ROW_ALT)

            row += 1
            stt += 1

    if not date_cells:
        ws.cell(2, 1, "Chưa có buổi tập trong lịch.")

    ws.auto_filter.ref = f"A1:F{max(1, row - 1)}"
    return date_cells


def _write_plan(
    ws: Worksheet,
    weeks: list[tuple[int, list[dict[str, Any]]]],
    date_cells: list[str],
    footer_text: str,
) -> None:
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    widths = {"A": 16, "B": 38, "C": 12, "D": 16, "E": 14}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws.merge_cells("A1:E1")
    ws["A1"] = "Lịch tập & thực đơn theo buổi"
    ws["A1"].font = _font(16, bold=True, color=WHITE)
    ws["A1"].fill = _fill(GREEN)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26

    cursor = 3
    session_i = 0
    for week_num, week_days in weeks:
        origin = ws.cell(cursor, 1, f"Tuần {week_num}")
        origin.font = _font(13, bold=True, color=WHITE)
        origin.fill = _fill(GREEN_DARK)
        origin.alignment = Alignment(vertical="center")
        ws.merge_cells(start_row=cursor, start_column=1, end_row=cursor, end_column=5)
        cursor += 2

        for day in week_days:
            date_ref = date_cells[session_i] if session_i < len(date_cells) else None
            cursor = _write_day_block(ws, cursor, day, date_ref)
            cursor += 1
            session_i += 1

    if footer_text:
        origin = ws.cell(cursor, 1, footer_text)
        origin.alignment = Alignment(wrap_text=True, vertical="top")
        origin.font = _font(10, color=MUTED)
        ws.merge_cells(start_row=cursor, start_column=1, end_row=cursor + 1, end_column=5)


def _write_day_block(ws: Worksheet, row: int, day: dict[str, Any], date_ref: str | None) -> int:
    title = day.get("title_vi") or f"Ngày {day.get('day_number') or ''}"
    head = ws.cell(row, 1, title)
    head.font = _font(12, bold=True, color=WHITE)
    head.fill = _fill(GREEN)
    head.alignment = Alignment(vertical="center")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    row += 1

    ws.cell(row, 1, "Ngày tập").font = _font(10, bold=True, color=MUTED)
    date_cell = ws.cell(row, 2)
    if date_ref:
        date_cell.value = f"={date_ref}"
    else:
        date_cell.value = "—"
    date_cell.number_format = "DD/MM/YYYY"
    date_cell.fill = _fill(YELLOW)
    date_cell.font = _font(11, bold=True)
    date_cell.border = thin
    ws.cell(row, 3, "Thứ").font = _font(10, bold=True, color=MUTED)
    weekday = ws.cell(row, 4, f"={_weekday_formula(date_cell.coordinate)}" if date_ref else "—")
    weekday.border = thin
    row += 2

    for i, h in enumerate(["Phần", "Bài tập", "Hiệp", "Khối lượng", "Nghỉ"], 1):
        cell = ws.cell(row, i, h)
        cell.font = _font(10, bold=True, color=WHITE)
        cell.fill = _fill(GREEN_DARK)
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin
    row += 1

    by_sec = _exercises_by_section(day.get("exercises") or [])
    section_notes = day.get("section_notes") or {}
    wrote_ex = False
    for sec in ALWAYS_SECTIONS:
        items = by_sec.get(sec) or []
        label = SECTION_LABEL.get(sec, sec)
        if not items:
            _data_row(ws, row, [label, "Không có", "", "", ""])
            row += 1
            wrote_ex = True
            continue
        for ex in items:
            _data_row(
                ws,
                row,
                [
                    label,
                    ex.get("name_vi") or "—",
                    ex.get("sets") or "",
                    _work_label(ex),
                    _rest_label(ex),
                ],
            )
            row += 1
            wrote_ex = True
        note = section_notes.get(sec)
        if note and str(note).strip():
            _note_row(ws, row, str(note).strip())
            row += 1
    if not wrote_ex:
        _data_row(ws, row, ["—", "Chưa có bài tập", "", "", ""])
        row += 1

    row += 1
    meal_head = ws.cell(row, 1, "Thực đơn")
    meal_head.font = _font(11, bold=True, color=WHITE)
    meal_head.fill = _fill(GREEN)
    meal_head.alignment = Alignment(vertical="center")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    row += 1
    for i, h in enumerate(["Bữa", "Món", "Khẩu phần", "kcal", ""], 1):
        cell = ws.cell(row, i, h)
        cell.font = _font(10, bold=True, color=WHITE)
        cell.fill = _fill(GREEN_DARK)
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin
    row += 1

    by_meal = _meals_by_type(day.get("meals") or [])
    meal_notes = day.get("meal_notes") or {}
    wrote_meal = False
    for mt in ("breakfast", "lunch", "dinner", "snack"):
        items = by_meal.get(mt) or []
        label = MEAL_LABEL.get(mt, mt)
        if not items:
            continue
        for meal in items:
            servings = meal.get("servings") or 1
            try:
                kcal = round(float(meal.get("calories") or 0))
            except (TypeError, ValueError):
                kcal = ""
            _data_row(
                ws,
                row,
                [label, meal.get("name_vi") or "—", servings, kcal, ""],
            )
            row += 1
            wrote_meal = True
        if mt != "snack":
            note = meal_notes.get(mt)
            if note and str(note).strip():
                _note_row(ws, row, str(note).strip())
                row += 1
    if not wrote_meal:
        _data_row(ws, row, ["—", "Chưa có thực đơn", "", "", ""])
        row += 1

    free = day.get("notes_vi")
    if free and str(free).strip():
        row += 1
        note_cell = ws.cell(row, 1, f"Lưu ý ngày: {str(free).strip()}")
        note_cell.fill = _fill("FFFBEB")
        note_cell.alignment = Alignment(wrap_text=True, vertical="center")
        note_cell.font = _font(10)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        row += 1

    return row + 1


def _note_row(ws: Worksheet, row: int, text: str) -> None:
    ws.cell(row, 1, "Lưu ý").font = _font(10, bold=True)
    ws.cell(row, 1).fill = _fill("FFFBEB")
    ws.cell(row, 1).border = thin
    body = ws.cell(row, 2, text)
    body.fill = _fill("FFFBEB")
    body.alignment = Alignment(wrap_text=True, vertical="center")
    body.border = thin
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)


def _data_row(ws: Worksheet, row: int, values: list[Any]) -> None:
    for col, val in enumerate(values, 1):
        cell = ws.cell(row, col, val)
        cell.border = thin
        cell.alignment = Alignment(vertical="center", wrap_text=col == 2)
        if col != 2:
            cell.alignment = Alignment(horizontal="center", vertical="center")
    title = day.get("title_vi") or f"Ngày {day.get('day_number') or ''}"
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    head = ws.cell(row, 1, title)
    head.font = _font(12, bold=True, color=WHITE)
    head.fill = _fill(GREEN)
    for col in range(1, 6):
        ws.cell(row, col).fill = _fill(GREEN)
        ws.cell(row, col).border = thin
    row += 1

    ws.cell(row, 1, "Ngày tập").font = _font(10, bold=True, color=MUTED)
    date_cell = ws.cell(row, 2)
    if date_ref:
        date_cell.value = f"={date_ref}"
    else:
        date_cell.value = "—"
    date_cell.number_format = "DD/MM/YYYY"
    date_cell.fill = _fill(YELLOW)
    date_cell.font = _font(11, bold=True)
    date_cell.border = thin
    ws.cell(row, 3, "Thứ").font = _font(10, bold=True, color=MUTED)
    weekday = ws.cell(row, 4, f"={_weekday_formula(date_cell.coordinate)}" if date_ref else "—")
    weekday.border = thin
    row += 2

    for i, h in enumerate(["Phần", "Bài tập", "Hiệp", "Khối lượng", "Nghỉ"], 1):
        cell = ws.cell(row, i, h)
        cell.font = _font(10, bold=True, color=WHITE)
        cell.fill = _fill(GREEN_DARK)
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin
    row += 1

    by_sec = _exercises_by_section(day.get("exercises") or [])
    section_notes = day.get("section_notes") or {}
    wrote_ex = False
    for sec in ALWAYS_SECTIONS:
        items = by_sec.get(sec) or []
        label = SECTION_LABEL.get(sec, sec)
        if not items:
            _data_row(ws, row, [label, "Không có", "", "", ""])
            row += 1
            wrote_ex = True
            continue
        for ex in items:
            _data_row(
                ws,
                row,
                [
                    label,
                    ex.get("name_vi") or "—",
                    ex.get("sets") or "",
                    _work_label(ex),
                    _rest_label(ex),
                ],
            )
            row += 1
            wrote_ex = True
        note = section_notes.get(sec)
        if note and str(note).strip():
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
            _data_row(ws, row, ["Lưu ý", str(note).strip(), "", "", ""])
            ws.cell(row, 1).fill = _fill("FFFBEB")
            ws.cell(row, 2).fill = _fill("FFFBEB")
            row += 1
    if not wrote_ex:
        _data_row(ws, row, ["—", "Chưa có bài tập", "", "", ""])
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    meal_head = ws.cell(row, 1, "Thực đơn")
    meal_head.font = _font(11, bold=True, color=WHITE)
    meal_head.fill = _fill(GREEN)
    for col in range(1, 6):
        ws.cell(row, col).fill = _fill(GREEN)
    row += 1
    for i, h in enumerate(["Bữa", "Món", "Khẩu phần", "kcal", ""], 1):
        cell = ws.cell(row, i, h)
        cell.font = _font(10, bold=True, color=WHITE)
        cell.fill = _fill(GREEN_DARK)
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin
    row += 1

    by_meal = _meals_by_type(day.get("meals") or [])
    meal_notes = day.get("meal_notes") or {}
    wrote_meal = False
    for mt in ("breakfast", "lunch", "dinner", "snack"):
        items = by_meal.get(mt) or []
        label = MEAL_LABEL.get(mt, mt)
        if not items:
            continue
        for meal in items:
            servings = meal.get("servings") or 1
            try:
                kcal = round(float(meal.get("calories") or 0))
            except (TypeError, ValueError):
                kcal = ""
            _data_row(
                ws,
                row,
                [label, meal.get("name_vi") or "—", servings, kcal, ""],
            )
            row += 1
            wrote_meal = True
        if mt != "snack":
            note = meal_notes.get(mt)
            if note and str(note).strip():
                ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
                _data_row(ws, row, ["Lưu ý", str(note).strip(), "", "", ""])
                row += 1
    if not wrote_meal:
        _data_row(ws, row, ["—", "Chưa có thực đơn", "", "", ""])
        row += 1

    free = day.get("notes_vi")
    if free and str(free).strip():
        row += 1
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        note_cell = ws.cell(row, 1, f"Lưu ý ngày: {str(free).strip()}")
        note_cell.fill = _fill("FFFBEB")
        note_cell.alignment = Alignment(wrap_text=True)
        row += 1

    return row + 1


def _data_row(ws: Worksheet, row: int, values: list[Any]) -> None:
    for col, val in enumerate(values, 1):
        cell = ws.cell(row, col, val)
        cell.border = thin
        cell.alignment = Alignment(vertical="center", wrap_text=col == 2)
        if col != 2:
            cell.alignment = Alignment(horizontal="center", vertical="center")
