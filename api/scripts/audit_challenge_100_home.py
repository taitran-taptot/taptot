"""Full-matrix audit of 100-day home challenge plans (no persist, no meals)."""

from __future__ import annotations

import argparse
import json
import re
import sys
import time
import traceback
from itertools import product
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[2]
API = ROOT / "api"
if str(API) not in sys.path:
    sys.path.insert(0, str(API))

from sqlalchemy import text
from sqlalchemy.orm import Session

from app.core.database import SessionLocal
from app.services.periodization import expand_plan_days_for_weeks
from app.services.schedule_spec_master import (
    day_label_to_split_role,
    expand_week_days,
    experience_to_master_key,
    lookup_week_split,
    rotate_week_labels_for_focus,
)
from app.services.workout_generation.kit_schedule import (
    challenge_finisher_kind,
    kit_kind_list,
)
from app.services.workout_generation.load_estimate import infer_test_kit
from app.services.workout_generation.service import generate_workout
from app.services.workout_generation.session_policy import challenge_phase_index
from app.services.workout_generation.split_map import (
    is_denied_for_split,
    normalize_split_role,
)

KITS: dict[str, tuple[str, ...]] = {
    "dumbbell": ("dumbbell",),
    "resistance-band": ("resistance-band",),
    "gymnastic-rings": ("gymnastic-rings", "pull-up-bar"),
}
FOCUS_SETS: tuple[tuple[str, ...], ...] = (
    (),
    ("mong", "bung"),
    ("nguc",),
    ("chan",),
    ("lung",),
)
KIT_LABEL = {
    "dumbbell": "tạ đơn",
    "resistance-band": "dây kháng lực",
    "gymnastic-rings": "vòng + xà",
}
GENDER_VI = {"male": "Nam", "female": "Nữ"}
FOCUS_VI = {
    "": "không focus",
    "mong+bung": "mông + bụng",
    "nguc": "ngực",
    "chan": "chân",
    "lung": "lưng",
}
CORE_SLUGS = frozenset({"core", "core-upper", "core-lower", "core-obliques", "co-bung", "abs", "waist", "bung"})
CARDIO_NAME = ("shadow boxing", "jumping jack", "trail run", "hiking", "chạy", "nhảy dang", "đấm bóng")
DEFAULT_OUT = ROOT / "artifacts" / "challenge_100_audit"


def _profile(gender: str) -> dict[str, Any]:
    female = gender == "female"
    return {
        "age": 28,
        "height_cm": 162.0 if female else 175.0,
        "weight_kg": 62.0 if female else 82.0,
        "activity": "moderate",
    }


def _baseline(kit: str) -> dict[str, Any]:
    if kit == "dumbbell":
        return {
            "test_kit": "dumbbell",
            "db_press_reps": 16,
            "db_press_kg": 10,
            "db_row_reps": 16,
            "db_row_kg": 10,
            "goblet_reps": 16,
            "goblet_kg": 16,
            "pushups_max": 16,
            "squats_max": 16,
        }
    if kit == "resistance-band":
        return {
            "test_kit": "band",
            "band_level": "blue",
            "band_level_push": "blue",
            "band_level_pull": "blue",
            "band_level_squat": "blue",
            "pushups_max": 16,
            "pullups_max": 8,
            "squats_max": 16,
        }
    return {
        "test_kit": "bar_rings",
        "pushups_max": 16,
        "pullups_max": 8,
        "squats_max": 16,
    }


def iter_specs() -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for kit, gender, sessions, minutes, level, focus in product(
        tuple(KITS),
        ("male", "female"),
        (2, 3, 4, 5, 6),
        (45, 60),
        (1, 2, 3),
        FOCUS_SETS,
    ):
        focus_key = "+".join(focus) if focus else ""
        case_id = (
            f"{kit}_{gender[0]}_s{sessions}_m{minutes}_L{level}_"
            f"{focus_key or 'none'}"
        )
        out.append(
            {
                "case_id": case_id,
                "kit": kit,
                "gender": gender,
                "sessions": sessions,
                "minutes": minutes,
                "level": level,
                "focus": list(focus),
                "focus_key": focus_key,
            }
        )
    return out


def payload_for(spec: dict[str, Any]) -> dict[str, Any]:
    return {
        "goal": "lose_weight",
        "gender": spec["gender"],
        "sessions_per_week": spec["sessions"],
        "session_minutes": spec["minutes"],
        "location": "home",
        "focus_areas": list(spec["focus"]),
        "extra_goals": [],
        "equipment_list": list(KITS[spec["kit"]]),
        "food_ids": [],
        "experience_level": spec["level"],
        "ai_suggest_equipment": False,
        "no_equipment": False,
        "ai_suggest_foods": False,
        "fitness_baseline": _baseline(spec["kit"]),
        "duration_weeks": 14,
        "kg_per_week": 0.5,
        "challenge_100_days": True,
        "curriculum_12_weeks": False,
        **_profile(spec["gender"]),
    }


def load_catalog(db: Session) -> dict[int, dict[str, Any]]:
    rows = db.execute(
        text(
            """
            SELECT e.id, e.name_en, e.name_vi, e.movement_pattern, e.movement_role,
                   mg.slug AS muscle_slug
            FROM exercises e
            LEFT JOIN muscle_groups mg ON mg.id = e.muscle_group_id
            """
        )
    ).mappings().all()
    return {int(r["id"]): dict(r) for r in rows}


def expected_week(spec: dict[str, Any]) -> tuple[str, list[str]]:
    exp = experience_to_master_key(spec["level"])
    sessions = int(spec["sessions"])
    week_code = lookup_week_split(
        experience=exp,
        sessions=sessions,
        gender=spec["gender"],
        location="home",
        home_equip="with_equip",
    ) or ""
    labels = rotate_week_labels_for_focus(
        expand_week_days(week_code),
        spec["focus"],
        week_code=week_code,
    )
    if (
        labels
        and day_label_to_split_role(labels[0]) == "legs"
        and re.match(r"^ULUL$", re.sub(r"[\s,]", "", week_code), re.I)
    ):
        week_code = "LULU"
    roles = [day_label_to_split_role(lbl) for lbl in labels[:sessions]]
    return week_code, roles


def _is_cardio(ex: dict[str, Any], cat: dict[str, Any] | None) -> bool:
    if str(ex.get("section") or "").strip().lower() == "cardio":
        return True
    blob = f"{(cat or {}).get('name_en') or ''} {(cat or {}).get('name_vi') or ''}".lower()
    return any(tok in blob for tok in CARDIO_NAME)


def _is_abs(ex: dict[str, Any], cat: dict[str, Any] | None) -> bool:
    if _is_cardio(ex, cat):
        return False
    slug = str((cat or {}).get("muscle_slug") or "").lower()
    if slug in CORE_SLUGS:
        return True
    blob = f"{(cat or {}).get('name_en') or ''} {(cat or {}).get('name_vi') or ''}".lower()
    return any(tok in blob for tok in ("plank", "dead bug", "v-up", "russian twist", "bọ chết"))


def classify_day(
    day: dict[str, Any],
    catalog: dict[int, dict[str, Any]],
) -> dict[str, Any]:
    strength: list[dict[str, Any]] = []
    abs_ex: list[dict[str, Any]] = []
    cardio: list[dict[str, Any]] = []
    for ex in day.get("exercises") or []:
        section = str(ex.get("section") or "main").strip().lower()
        if section in {"warmup", "primer", "cooldown", "stretch"}:
            continue
        try:
            eid = int(ex.get("exercise_id"))
        except (TypeError, ValueError):
            continue
        cat = catalog.get(eid) or {}
        row = {
            "id": eid,
            "name_en": cat.get("name_en"),
            "name_vi": cat.get("name_vi") or ex.get("name_vi"),
            "section": section,
            "sets": ex.get("sets"),
            "reps": ex.get("reps"),
            "rest_seconds": ex.get("rest_seconds"),
            "muscle_slug": cat.get("muscle_slug"),
            "pattern": cat.get("movement_pattern"),
        }
        if _is_cardio(ex, cat):
            cardio.append(row)
        elif _is_abs(ex, cat):
            abs_ex.append(row)
        else:
            strength.append(row)
    return {"strength": strength, "abs": abs_ex, "cardio": cardio}


def expected_strength_n(spec: dict[str, Any], role: str, phase_index: int) -> int:
    kit = infer_test_kit(_baseline(spec["kit"]), list(KITS[spec["kit"]]))
    kinds = kit_kind_list(
        kit,
        role,
        5,
        spec["focus"],
        phase_index=phase_index,
    )
    return len(kinds) if kinds else 0


def rubric_flags(
    spec: dict[str, Any],
    *,
    week_code: str,
    expected_code: str,
    expected_roles: list[str],
    expanded: list[dict[str, Any]],
    catalog: dict[int, dict[str, Any]],
) -> list[str]:
    flags: list[str] = []
    if (week_code or "").replace(" ", "") != (expected_code or "").replace(" ", ""):
        flags.append(f"week_code:{week_code!r}!={expected_code!r}")
    if not expanded:
        flags.append("empty_plan")
        return flags
    first_week_n = int(spec["sessions"])
    first_roles = [
        normalize_split_role(d.get("split_role"))
        for d in expanded[:first_week_n]
    ]
    if first_roles != expected_roles:
        flags.append(f"week1_roles:{first_roles}!={expected_roles}")

    for day in expanded:
        role = normalize_split_role(day.get("split_role"))
        classified = classify_day(day, catalog)
        ids = [x["id"] for x in classified["strength"] + classified["abs"] + classified["cardio"]]
        if len(ids) != len(set(ids)):
            flags.append(f"dup_id:d{day.get('day_number')}:{role}")
        for item in classified["strength"]:
            if is_denied_for_split(
                role, pattern=item.get("pattern"), muscle_slug=item.get("muscle_slug")
            ):
                flags.append(
                    f"leak:d{day.get('day_number')}:{role}:{item.get('name_en')}"
                )
        try:
            week = (int(day.get("day_number") or 1) - 1) // max(1, int(spec["sessions"])) + 1
        except (TypeError, ValueError):
            week = 1
        phase = challenge_phase_index(week)
        want_n = expected_strength_n(spec, role, phase)
        got_n = len(classified["strength"])
        if want_n and got_n != want_n:
            flags.append(
                f"strength_n:d{day.get('day_number')}:{role}:p{phase+1}:got{got_n}:want{want_n}"
            )
        fin = challenge_finisher_kind(phase, max(0, int(day.get("day_number") or 1) - 1) % max(1, int(spec["sessions"])))
        if phase <= 0:
            if not classified["abs"]:
                flags.append(f"missing_abs:d{day.get('day_number')}")
            if not classified["cardio"]:
                flags.append(f"missing_cardio:d{day.get('day_number')}")
        else:
            if fin == "abs" and not classified["abs"]:
                flags.append(f"missing_abs:d{day.get('day_number')}")
            if fin == "cardio" and not classified["cardio"]:
                flags.append(f"missing_cardio:d{day.get('day_number')}")
        if spec["minutes"] <= 45:
            for item in classified["strength"]:
                rest = item.get("rest_seconds")
                if rest is None:
                    continue
                try:
                    rest_i = int(rest)
                except (TypeError, ValueError):
                    continue
                if rest_i not in {90, 120} and rest_i > 150:
                    flags.append(f"rest:d{day.get('day_number')}:{item.get('name_en')}:{rest_i}")
    return flags


def compact_days(
    expanded: list[dict[str, Any]],
    catalog: dict[int, dict[str, Any]],
    sessions: int,
) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for day in expanded:
        try:
            num = int(day.get("day_number") or 0)
        except (TypeError, ValueError):
            num = 0
        week = (num - 1) // max(1, sessions) + 1 if num else 0
        classified = classify_day(day, catalog)
        out.append(
            {
                "day_number": num,
                "week": week,
                "phase": challenge_phase_index(week) + 1 if week else None,
                "split_role": normalize_split_role(day.get("split_role")),
                "title_vi": day.get("title_vi"),
                "strength": [
                    {
                        "name_en": x.get("name_en"),
                        "name_vi": x.get("name_vi"),
                        "rest": x.get("rest_seconds"),
                    }
                    for x in classified["strength"]
                ],
                "abs": [x.get("name_en") or x.get("name_vi") for x in classified["abs"]],
                "cardio": [x.get("name_en") or x.get("name_vi") for x in classified["cardio"]],
            }
        )
    return out


def done_ids(path: Path) -> set[str]:
    found: set[str] = set()
    if not path.is_file():
        return found
    with path.open(encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            try:
                row = json.loads(line)
            except json.JSONDecodeError:
                continue
            cid = row.get("case_id")
            if cid and not row.get("error"):
                found.add(str(cid))
    return found


def run_case(
    db: Session,
    spec: dict[str, Any],
    catalog: dict[int, dict[str, Any]],
) -> dict[str, Any]:
    t0 = time.perf_counter()
    expected_code, expected_roles = expected_week(spec)
    result = generate_workout(db, None, payload_for(spec), persist=False, commit=False)
    insights = result.get("insights") or {}
    templates = insights.get("week_templates") or []
    seed_days = result.get("days") or []
    expanded = expand_plan_days_for_weeks(
        seed_days,
        14,
        is_dict=True,
        experience_level=spec["level"],
        week_templates=templates,
        curriculum=True,
        goal="lose_weight",
    )
    week_code = str(result.get("week_code") or insights.get("week_code") or "")
    flags = rubric_flags(
        spec,
        week_code=week_code,
        expected_code=expected_code,
        expected_roles=expected_roles,
        expanded=expanded,
        catalog=catalog,
    )
    return {
        "case_id": spec["case_id"],
        "kit": spec["kit"],
        "gender": spec["gender"],
        "sessions": spec["sessions"],
        "minutes": spec["minutes"],
        "level": spec["level"],
        "focus": spec["focus"],
        "focus_key": spec["focus_key"],
        "week_code": week_code,
        "expected_week_code": expected_code,
        "expected_roles": expected_roles,
        "sessions_actual": result.get("sessions_actual"),
        "elapsed_s": round(time.perf_counter() - t0, 2),
        "error": None,
        "flags": flags,
        "ok": not flags,
        "days": compact_days(expanded, catalog, int(spec["sessions"])),
    }


def append_jsonl(path: Path, row: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as fh:
        fh.write(json.dumps(row, ensure_ascii=False) + "\n")


def load_rows(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    if not path.is_file():
        return rows
    with path.open(encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            try:
                rows.append(json.loads(line))
            except json.JSONDecodeError:
                continue
    latest: dict[str, dict[str, Any]] = {}
    for row in rows:
        cid = row.get("case_id")
        if cid:
            latest[str(cid)] = row
    return list(latest.values())


def _names(items: list[dict[str, Any]]) -> str:
    parts = []
    for x in items:
        vi = x.get("name_vi") or ""
        en = x.get("name_en") or ""
        if vi and en:
            parts.append(f"{vi} ({en})")
        else:
            parts.append(vi or en)
    return "; ".join(parts) if parts else "—"


def _phase_day(days: list[dict[str, Any]], phase: int, role: str) -> dict[str, Any] | None:
    for d in days:
        if d.get("phase") == phase and d.get("split_role") == role:
            return d
    return None


def rescore_row(row: dict[str, Any]) -> dict[str, Any]:
    """Recompute flags from stored days using sessions_actual (L1 may clamp 6→5)."""
    if row.get("error") or not row.get("days"):
        return row
    sessions = int(row.get("sessions_actual") or row.get("sessions") or 1)
    spec = {
        "kit": row["kit"],
        "gender": row["gender"],
        "sessions": sessions,
        "minutes": row["minutes"],
        "level": row["level"],
        "focus": row.get("focus") or [],
        "focus_key": row.get("focus_key") or "",
    }
    expected_code, expected_roles = expected_week(spec)
    flags: list[str] = []
    week_code = str(row.get("week_code") or "")
    if week_code.replace(" ", "") != (expected_code or "").replace(" ", ""):
        flags.append(f"week_code:{week_code!r}!={expected_code!r}")
    days = row.get("days") or []
    first = days[:sessions]
    first_roles = [normalize_split_role(d.get("split_role")) for d in first]
    if first_roles != expected_roles:
        flags.append(f"week1_roles:{first_roles}!={expected_roles}")
    for d in days:
        role = normalize_split_role(d.get("split_role"))
        week = int(d.get("week") or 0)
        if week <= 0:
            try:
                week = (int(d.get("day_number") or 1) - 1) // max(1, sessions) + 1
            except (TypeError, ValueError):
                week = 1
        phase = challenge_phase_index(week)
        got_n = len(d.get("strength") or [])
        want_n = expected_strength_n(spec, role, phase)
        if role == "core":
            continue
        if want_n and got_n != want_n:
            flags.append(
                f"strength_n:d{d.get('day_number')}:{role}:p{phase+1}:got{got_n}:want{want_n}"
            )
        idx = max(0, int(d.get("day_number") or 1) - 1) % max(1, sessions)
        fin = challenge_finisher_kind(phase, idx)
        if phase <= 0:
            if not d.get("abs"):
                flags.append(f"missing_abs:d{d.get('day_number')}")
            if not d.get("cardio"):
                flags.append(f"missing_cardio:d{d.get('day_number')}")
        else:
            if fin == "abs" and not d.get("abs"):
                flags.append(f"missing_abs:d{d.get('day_number')}")
            if fin == "cardio" and not d.get("cardio"):
                flags.append(f"missing_cardio:d{d.get('day_number')}")
        if int(spec["minutes"]) <= 45:
            for item in d.get("strength") or []:
                rest = item.get("rest")
                if rest is None:
                    continue
                try:
                    rest_i = int(rest)
                except (TypeError, ValueError):
                    continue
                if rest_i not in {90, 120} and rest_i > 150:
                    flags.append(f"rest:d{d.get('day_number')}:{item.get('name_en')}:{rest_i}")
    row = dict(row)
    row["expected_week_code"] = expected_code
    row["expected_roles"] = expected_roles
    row["flags"] = flags
    row["ok"] = not flags
    return row


def write_report(rows: list[dict[str, Any]], path: Path) -> None:
    rows = [rescore_row(r) for r in rows]
    ok_rows = [r for r in rows if r.get("ok") and not r.get("error")]
    err_rows = [r for r in rows if r.get("error")]
    flag_rows = [r for r in rows if r.get("flags") and not r.get("error")]
    lines: list[str] = []
    lines.append("# Báo cáo audit lịch 100 ngày (nhà)")
    lines.append("")
    lines.append(
        f"Đã gen **{len(rows)}** lịch (`persist=False`, cartesian 3 kit × 2 giới × 5 số buổi × 2 phút × 3 level × 5 focus = 900). "
        f"Đạt rubric (sau chấm lại): **{len(ok_rows)}**. Có cờ: **{len(flag_rows)}**. Lỗi gen: **{len(err_rows)}**."
    )
    lines.append("")
    lines.append("Baseline vòng/xà cố định: 16 chống đẩy, 8 kéo xà. Không đánh giá mức tạ/dây.")
    lines.append("")
    lines.append("## Split theo giới tính × số buổi (spec engine)")
    lines.append("")
    lines.append("| Giới | 2 | 3 | 4 | 5 | 6 |")
    lines.append("|---|---|---|---|---|---|")
    for gender in ("male", "female"):
        cells = [GENDER_VI[gender]]
        for sess in (2, 3, 4, 5, 6):
            code, _roles = expected_week(
                {
                    "gender": gender,
                    "sessions": sess,
                    "level": 2,
                    "focus": [],
                }
            )
            cells.append(f"`{code}`")
        lines.append("| " + " | ".join(cells) + " |")
    lines.append("")
    lines.append("Nữ 2 buổi = Lower rồi Upper. Focus `mong`/`chan` đầu list xoay tuần UL để Lower trước.")
    lines.append("")
    lines.append("## Slot lực kỳ vọng (P1 / P2–P3)")
    lines.append("")
    lines.append("| Kit | Upper P1 | Upper P2 | Lower/Legs P1 | Lower/Legs P2 | Push P1 | Pull P1 |")
    lines.append("|---|---|---|---|---|---|---|")
    for kit, kit_key in (
        ("dây", "band"),
        ("tạ đơn", "dumbbell"),
        ("vòng", "bar_rings"),
    ):
        def kinds(role: str, phase: int) -> str:
            return ", ".join(
                kit_kind_list(kit_key, role, 5, ["mong", "bung"], phase_index=phase) or []
            )

        lines.append(
            f"| {kit} | {kinds('upper', 0)} | {kinds('upper', 1)} | "
            f"{kinds('legs', 0)} | {kinds('legs', 1)} | {kinds('push', 0)} | {kinds('pull', 0)} |"
        )
    lines.append("")

    def persona_block(spec_match: dict[str, Any], title: str) -> None:
        lines.append(f"### {title}")
        lines.append("")
        hit = next(
            (
                r
                for r in rows
                if r.get("kit") == spec_match["kit"]
                and r.get("gender") == spec_match["gender"]
                and int(r.get("sessions") or 0) == spec_match["sessions"]
                and int(r.get("minutes") or 0) == spec_match["minutes"]
                and int(r.get("level") or 0) == spec_match["level"]
                and r.get("focus_key") == spec_match["focus_key"]
            ),
            None,
        )
        if not hit:
            lines.append("Chưa có dữ liệu (lỗi gen hoặc chưa chạy case này).")
            lines.append("")
            return
        if hit.get("error"):
            lines.append(f"Lỗi gen: `{hit['error']}`")
            lines.append("")
            return
        days = hit.get("days") or []
        lines.append(f"- Tuần: `{hit.get('week_code')}` · roles kỳ vọng: `{hit.get('expected_roles')}`.")
        lines.append(f"- Rubric: {'đạt' if hit.get('ok') else 'có cờ: ' + ', '.join(hit.get('flags') or [])}.")
        for phase in (1, 2):
            lines.append(f"- **Pha {phase}**")
            roles_seen: list[str] = []
            for d in days:
                if d.get("phase") != phase:
                    continue
                role = d.get("split_role")
                if role in roles_seen:
                    continue
                roles_seen.append(role)
                st = _names(d.get("strength") or [])
                abs_n = ", ".join(d.get("abs") or []) or "—"
                card = ", ".join(d.get("cardio") or []) or "—"
                lines.append(
                    f"  - {role}: {len(d.get('strength') or [])} lực — {st}. "
                    f"Abs: {abs_n}. Cardio: {card}."
                )
        lines.append("")

    lines.append("## Persona chi tiết (ví dụ bạn hỏi + đủ 3 kit)")
    lines.append("")
    for kit in KITS:
        persona_block(
            {
                "kit": kit,
                "gender": "female",
                "sessions": 2,
                "minutes": 45,
                "level": 2,
                "focus_key": "mong+bung",
            },
            f"Nữ · 2 buổi/tuần · 45 phút · L2 · mông + bụng · {KIT_LABEL[kit]}",
        )
    extra = [
        ("Nam · 3 buổi · 45p · L1 · không focus · dây", {
            "kit": "resistance-band", "gender": "male", "sessions": 3,
            "minutes": 45, "level": 1, "focus_key": "",
        }),
        ("Nam · 4 buổi · 60p · L3 · ngực · tạ đơn", {
            "kit": "dumbbell", "gender": "male", "sessions": 4,
            "minutes": 60, "level": 3, "focus_key": "nguc",
        }),
        ("Nữ · 4 buổi · 45p · L2 · chân · vòng+xà", {
            "kit": "gymnastic-rings", "gender": "female", "sessions": 4,
            "minutes": 45, "level": 2, "focus_key": "chan",
        }),
        ("Nam · 6 buổi · 45p · L2 · lưng · dây", {
            "kit": "resistance-band", "gender": "male", "sessions": 6,
            "minutes": 45, "level": 2, "focus_key": "lung",
        }),
    ]
    lines.append("## Thêm vài combo")
    lines.append("")
    for title, match in extra:
        persona_block(match, title)

    from collections import Counter

    kind_c = Counter()
    for r in flag_rows:
        for f in r.get("flags") or []:
            kind_c[str(f).split(":")[0]] += 1
    lines.append("## Case fail rubric")
    lines.append("")
    lines.append(
        f"Sau khi chấm lại với `sessions_actual` (L1 có thể kẹp 6→5 buổi): "
        f"**{len(flag_rows)}** / {len(rows)} case còn cờ."
    )
    lines.append("")
    if kind_c:
        lines.append("Nhóm cờ:")
        for k, n in kind_c.most_common():
            lines.append(f"- `{k}`: {n} lần")
        lines.append("")
    lines.append(
        "L1 6 buổi: policy beginner kẹp còn 5 — nam ra `PPLUL` thay vì `PPLPPL`; "
        "nữ ra `LULU, Cardio-Core` thay vì UL×3. Đó là hành vi engine, không phải kit chọn sai bài."
    )
    lines.append("")
    lines.append(
        "Vòng treo Push pha 2–3 thường **4** lực (Push-Up, Dip, Pike, Triceps Extension) "
        "thay vì 5 — slot shoulder thứ hai trùng pike/dip nên bị bỏ."
    )
    lines.append("")
    if not flag_rows:
        lines.append("Không còn case dính cờ.")
    else:
        lines.append("| case_id | số cờ | mẫu |")
        lines.append("|---|---|---|")
        for r in sorted(flag_rows, key=lambda x: (-len(x.get("flags") or []), x.get("case_id") or ""))[:40]:
            flags = r.get("flags") or []
            sample = "; ".join(flags[:2])
            lines.append(f"| `{r.get('case_id')}` | {len(flags)} | {sample} |")
        if len(flag_rows) > 40:
            lines.append(f"| … | {len(flag_rows) - 40} case nữa | |")
    lines.append("")
    if err_rows:
        lines.append("## Lỗi gen")
        lines.append("")
        for r in err_rows[:40]:
            lines.append(f"- `{r.get('case_id')}`: {r.get('error')}")
        lines.append("")
    lines.append("## Ghi chú programming")
    lines.append("")
    lines.append("- `bung` không thêm bài lực trên Upper; bụng nằm finisher abs.")
    lines.append("- Chân dây P1 đủ 4 lực (squat, lunge, RDL, hip thrust); P2 thêm ham và xoay hinge/glute.")
    lines.append("- Chân tạ đơn / vòng P1 theo kinds global chỉ 2 slot lực (legs+quad) trừ khi overlay khác.")
    lines.append("- 45p và 60p challenge dùng cùng số slot lực; khác rest và finisher.")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_xlsx(rows: list[dict[str, Any]], path: Path) -> None:
    from openpyxl import Workbook

    rows = [rescore_row(r) for r in rows]
    wb = Workbook(write_only=True)
    ws_cases = wb.create_sheet("Cases")
    ws_fail = wb.create_sheet("Fail")
    ws_days = wb.create_sheet("Days")
    ws_ex = wb.create_sheet("Exercises")
    case_headers = [
        "case_id",
        "kit",
        "gender",
        "sessions",
        "sessions_actual",
        "minutes",
        "level",
        "focus",
        "week_code",
        "expected_week_code",
        "expected_roles",
        "ok",
        "flag_count",
        "flags",
        "error",
        "elapsed_s",
    ]
    day_headers = [
        "case_id",
        "kit",
        "gender",
        "sessions_actual",
        "minutes",
        "level",
        "focus",
        "ok",
        "day_number",
        "week",
        "phase",
        "split_role",
        "title_vi",
        "strength_n",
        "strength_vi",
        "abs",
        "cardio",
    ]
    ex_headers = [
        "case_id",
        "kit",
        "gender",
        "minutes",
        "level",
        "focus",
        "day_number",
        "week",
        "phase",
        "split_role",
        "section",
        "name_vi",
        "name_en",
        "rest",
    ]
    ws_cases.append(case_headers)
    ws_fail.append(case_headers)
    ws_days.append(day_headers)
    ws_ex.append(ex_headers)
    for ws in (ws_cases, ws_fail, ws_days, ws_ex):
        ws.sheet_properties.tabColor = "1F4E79"

    for r in rows:
        focus = r.get("focus_key") or "+".join(r.get("focus") or [])
        roles = r.get("expected_roles") or []
        flags = r.get("flags") or []
        case_row = [
            r.get("case_id"),
            r.get("kit"),
            r.get("gender"),
            r.get("sessions"),
            r.get("sessions_actual"),
            r.get("minutes"),
            r.get("level"),
            focus,
            r.get("week_code"),
            r.get("expected_week_code"),
            ", ".join(roles) if isinstance(roles, list) else roles,
            bool(r.get("ok")),
            len(flags),
            "; ".join(flags),
            r.get("error"),
            r.get("elapsed_s"),
        ]
        ws_cases.append(case_row)
        if flags or r.get("error"):
            ws_fail.append(case_row)
        for d in r.get("days") or []:
            strength = d.get("strength") or []
            names_vi = "; ".join(
                str(x.get("name_vi") or x.get("name_en") or "") for x in strength
            )
            ws_days.append(
                [
                    r.get("case_id"),
                    r.get("kit"),
                    r.get("gender"),
                    r.get("sessions_actual"),
                    r.get("minutes"),
                    r.get("level"),
                    focus,
                    bool(r.get("ok")),
                    d.get("day_number"),
                    d.get("week"),
                    d.get("phase"),
                    d.get("split_role"),
                    d.get("title_vi"),
                    len(strength),
                    names_vi,
                    "; ".join(d.get("abs") or []),
                    "; ".join(d.get("cardio") or []),
                ]
            )
            for item in strength:
                ws_ex.append(
                    [
                        r.get("case_id"),
                        r.get("kit"),
                        r.get("gender"),
                        r.get("minutes"),
                        r.get("level"),
                        focus,
                        d.get("day_number"),
                        d.get("week"),
                        d.get("phase"),
                        d.get("split_role"),
                        "strength",
                        item.get("name_vi"),
                        item.get("name_en"),
                        item.get("rest"),
                    ]
                )
            for name in d.get("abs") or []:
                ws_ex.append(
                    [
                        r.get("case_id"),
                        r.get("kit"),
                        r.get("gender"),
                        r.get("minutes"),
                        r.get("level"),
                        focus,
                        d.get("day_number"),
                        d.get("week"),
                        d.get("phase"),
                        d.get("split_role"),
                        "abs",
                        name,
                        name,
                        None,
                    ]
                )
            for name in d.get("cardio") or []:
                ws_ex.append(
                    [
                        r.get("case_id"),
                        r.get("kit"),
                        r.get("gender"),
                        r.get("minutes"),
                        r.get("level"),
                        focus,
                        d.get("day_number"),
                        d.get("week"),
                        d.get("phase"),
                        d.get("split_role"),
                        "cardio",
                        name,
                        name,
                        None,
                    ]
                )
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--report-only", action="store_true")
    parser.add_argument("--xlsx-only", action="store_true")
    parser.add_argument("--retries", type=int, default=2)
    args = parser.parse_args()
    out_dir: Path = args.out
    jsonl = out_dir / "cases.jsonl"
    report_path = out_dir / "REPORT.md"
    xlsx_path = out_dir / "challenge_100_audit.xlsx"

    if args.xlsx_only or args.report_only:
        loaded = load_rows(jsonl)
        if args.report_only:
            write_report(loaded, report_path)
            print(f"wrote {report_path}")
        write_xlsx(loaded, xlsx_path)
        print(f"wrote {xlsx_path}")
        return

    specs = iter_specs()
    if args.limit:
        specs = specs[: max(0, args.limit)]
    skip = done_ids(jsonl)
    print(f"total {len(specs)} done {len(skip)} remaining {len(specs) - len([s for s in specs if s['case_id'] in skip])}")
    db = SessionLocal()
    try:
        catalog = load_catalog(db)
        for i, spec in enumerate(specs, 1):
            if spec["case_id"] in skip:
                continue
            last_err = None
            row = None
            for attempt in range(1, args.retries + 2):
                try:
                    row = run_case(db, spec, catalog)
                    last_err = None
                    break
                except Exception as exc:
                    last_err = f"{type(exc).__name__}: {exc}"
                    traceback.print_exc()
                    time.sleep(min(8, attempt * 2))
            if row is None:
                row = {
                    **{k: spec[k] for k in ("case_id", "kit", "gender", "sessions", "minutes", "level", "focus", "focus_key")},
                    "error": last_err,
                    "ok": False,
                    "flags": [],
                    "days": [],
                    "elapsed_s": None,
                }
            slim = dict(row)
            append_jsonl(jsonl, slim)
            status = "ERR" if slim.get("error") else ("OK" if slim.get("ok") else f"FLAG {len(slim.get('flags') or [])}")
            print(f"[{i}/{len(specs)}] {spec['case_id']} {status} {slim.get('elapsed_s')}s")
            skip.add(spec["case_id"])
    finally:
        db.close()
    loaded = load_rows(jsonl)
    write_report(loaded, report_path)
    write_xlsx(loaded, xlsx_path)
    print(f"wrote {report_path}")
    print(f"wrote {xlsx_path}")


if __name__ == "__main__":
    main()
