"""Export all Vietnamese cooking posts (cách nấu món) to Excel + Markdown.

Prefers the live database; falls back to seeds/cooking_posts if DB is empty
or unavailable.
"""

from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

API_DIR = Path(__file__).resolve().parents[1]
PROJECT_ROOT = API_DIR.parent
if str(API_DIR) not in sys.path:
    sys.path.insert(0, str(API_DIR))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.cooking_post_service import GROUP_SLUG_TO_VI

HEADER = [
    "id",
    "slug",
    "title_vi",
    "excerpt",
    "dish_slug",
    "group_file",
    "group_vi",
    "is_published",
    "sort_order",
    "servings",
    "yield_grams",
    "cover_image_url",
    "ingredients",
    "content_md",
    "url_path",
]

HEADER_FILL = PatternFill("FF1F4E3D", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFFFF")

GROUP_FILE_TO_VI = {
    "mon-com-gia-dinh.json": "Món Cơm Gia Đình",
    "dac-san-vung-mien.json": "Đặc sản vùng miền",
    "mon-nuoc-soi.json": "Món Nước & Sợi",
    "banh-mi-mon-cuon.json": "Bánh Mì & Món Cuốn",
}


def _iso(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat(timespec="seconds")
    return str(value)


def _json_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=False)


def _load_seed_rows() -> list[dict[str, Any]]:
    folder = PROJECT_ROOT / "seeds" / "cooking_posts"
    rows: list[dict[str, Any]] = []
    if not folder.is_dir():
        return rows
    for path in sorted(folder.glob("*.json")):
        data = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(data, list):
            continue
        group_vi = GROUP_FILE_TO_VI.get(path.name, path.stem)
        for item in data:
            if not isinstance(item, dict):
                continue
            rows.append(
                {
                    "id": None,
                    "slug": item.get("slug") or "",
                    "title_vi": item.get("title_vi") or "",
                    "excerpt": item.get("excerpt") or "",
                    "dish_slug": item.get("dish_slug") or "",
                    "group_file": path.name,
                    "group_vi": group_vi,
                    "is_published": bool(item.get("is_published", True)),
                    "sort_order": int(item.get("sort_order") or 0),
                    "servings": int(item.get("servings") or 1),
                    "yield_grams": item.get("yield_grams"),
                    "cover_image_url": item.get("cover_image_url") or "",
                    "ingredients": item.get("ingredients") or [],
                    "content_md": item.get("content_md") or "",
                    "source": "seed",
                }
            )
    rows.sort(key=lambda r: (r["group_file"], r["sort_order"] or 0, r["slug"]))
    return rows


def _load_db_rows() -> list[dict[str, Any]]:
    from sqlalchemy.orm import Session

    from app.core.database import SessionLocal
    from app.models.entities import CookingPost, Food

    db: Session = SessionLocal()
    try:
        posts = list(db.query(CookingPost).order_by(CookingPost.sort_order, CookingPost.id).all())
        if not posts:
            return []
        dish_slugs = [p.dish_slug for p in posts if p.dish_slug]
        foods: dict[str, Food] = {}
        if dish_slugs:
            for food in db.query(Food).filter(Food.slug.in_(dish_slugs)).all():
                foods[food.slug] = food
        rows: list[dict[str, Any]] = []
        for post in posts:
            dish = foods.get(post.dish_slug or "")
            group_slug = None
            group_vi = ""
            if dish is not None and isinstance(dish.tags, list):
                for raw in dish.tags:
                    tag = str(raw)
                    if tag.startswith("nhom_vi:"):
                        group_vi = tag[len("nhom_vi:") :].strip()
                    elif tag.startswith("nhom:"):
                        group_slug = tag[len("nhom:") :].strip()
                if group_slug and not group_vi:
                    group_vi = GROUP_SLUG_TO_VI.get(group_slug, group_slug)
            rows.append(
                {
                    "id": post.id,
                    "slug": post.slug,
                    "title_vi": post.title_vi,
                    "excerpt": post.excerpt or "",
                    "dish_slug": post.dish_slug or "",
                    "group_file": "",
                    "group_vi": group_vi,
                    "is_published": bool(post.is_published),
                    "sort_order": post.sort_order,
                    "servings": int(post.servings or 1),
                    "yield_grams": post.yield_grams,
                    "cover_image_url": post.cover_image_url or "",
                    "ingredients": post.ingredients or [],
                    "content_md": post.content_md or "",
                    "source": "db",
                    "published_at": _iso(post.published_at),
                }
            )
        return rows
    finally:
        db.close()


def _row_values(row: dict[str, Any]) -> list[Any]:
    return [
        row.get("id") or "",
        row.get("slug") or "",
        row.get("title_vi") or "",
        row.get("excerpt") or "",
        row.get("dish_slug") or "",
        row.get("group_file") or "",
        row.get("group_vi") or "",
        bool(row.get("is_published")),
        row.get("sort_order") or 0,
        row.get("servings") or 1,
        row.get("yield_grams") if row.get("yield_grams") is not None else "",
        row.get("cover_image_url") or "",
        _json_cell(row.get("ingredients")),
        row.get("content_md") or "",
        f"/cach-nau?mon={row.get('slug') or ''}",
    ]


def _write_header(ws) -> None:
    ws.append(HEADER)
    for col, _ in enumerate(HEADER, start=1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"


def _set_widths(ws) -> None:
    for col, title in enumerate(HEADER, start=1):
        letter = get_column_letter(col)
        if title == "content_md":
            width = 80
        elif title in {"title_vi", "excerpt", "ingredients"}:
            width = 40
        elif title in {"slug", "dish_slug", "url_path", "cover_image_url"}:
            width = 32
        else:
            width = max(12, min(18, len(title) + 2))
        ws.column_dimensions[letter].width = width
    content_col = HEADER.index("content_md") + 1
    for row_idx in range(2, ws.max_row + 1):
        cell = ws.cell(row_idx, content_col)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row_idx].height = 80


def _ingredient_md(ingredients: Any) -> str:
    lines: list[str] = []
    if not isinstance(ingredients, list):
        return ""
    for item in ingredients:
        if not isinstance(item, dict):
            continue
        label = str(item.get("amount_label") or "").strip()
        slug = str(item.get("food_slug") or "").strip()
        grams = item.get("grams")
        note = str(item.get("note") or "").strip()
        text = label or (f"{grams}g {slug}" if grams is not None else slug)
        if note:
            text = f"{text} ({note})"
        if text:
            lines.append(f"- {text}")
    return "\n".join(lines)


def _write_markdown(rows: list[dict[str, Any]], md_dir: Path) -> None:
    md_dir.mkdir(parents=True, exist_ok=True)
    index_lines = [
        "# Cách nấu món Việt (kho thực phẩm)",
        "",
        f"Tổng số bài: {len(rows)}",
        "",
    ]
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        grouped.setdefault(row.get("group_vi") or "Khác", []).append(row)
    for group, items in grouped.items():
        index_lines.append(f"## {group}")
        index_lines.append("")
        for row in items:
            slug = row.get("slug") or "bai"
            title = row.get("title_vi") or slug
            index_lines.append(f"- [{title}]({slug}.md)")
        index_lines.append("")
    (md_dir / "00-TAT-CA-BAI-VIET.md").write_text("\n".join(index_lines), encoding="utf-8")

    for row in rows:
        slug = row.get("slug") or "bai"
        parts = [
            f"# {row.get('title_vi') or slug}",
            "",
            f"- Nhóm: {row.get('group_vi') or ''}",
            f"- Món (dish_slug): {row.get('dish_slug') or ''}",
            f"- Khẩu phần: {row.get('servings') or 1}",
            f"- Tổng khối lượng: {row.get('yield_grams') or ''}",
            f"- Đường dẫn: `/cach-nau?mon={slug}`",
            "",
        ]
        excerpt = (row.get("excerpt") or "").strip()
        if excerpt:
            parts.extend([excerpt, ""])
        ingredients_md = _ingredient_md(row.get("ingredients"))
        if ingredients_md:
            parts.extend(["## Nguyên liệu", "", ingredients_md, ""])
        body = (row.get("content_md") or "").strip()
        if body:
            parts.extend([body, ""])
        (md_dir / f"{slug}.md").write_text("\n".join(parts), encoding="utf-8")


def main() -> Path:
    source = "db"
    try:
        rows = _load_db_rows()
        if not rows:
            source = "seed"
            rows = _load_seed_rows()
    except Exception as exc:
        print(f"DB unavailable ({exc}); falling back to seeds.")
        source = "seed"
        rows = _load_seed_rows()

    if not rows:
        raise SystemExit("No cooking posts found in DB or seeds.")

    out_dir = API_DIR / "exports"
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    out = out_dir / f"cooking_posts_{stamp}.xlsx"
    out_dir.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "cooking_posts"
    _write_header(ws)
    for row in rows:
        ws.append(_row_values(row))
    _set_widths(ws)
    wb.save(out)

    md_dir = out.with_suffix("")
    _write_markdown(rows, md_dir)
    print(f"Exported {len(rows)} cooking posts from {source}")
    print(out)
    print(md_dir)
    return out


if __name__ == "__main__":
    main()
