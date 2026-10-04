"""Export exercise catalog to Excel (api/exports/)."""

from __future__ import annotations

import json
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

API_DIR = Path(__file__).resolve().parents[1]
if str(API_DIR) not in sys.path:
    sys.path.insert(0, str(API_DIR))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.core.database import SessionLocal
from app.models.entities import Equipment, Exercise, ExerciseEquipment, MuscleGroup

HEADER = [
    "id",
    "is_active",
    "name_vi",
    "name_en",
    "muscle_slug",
    "muscle_vi",
    "secondary_muscles",
    "equipment_slugs",
    "equipment_vi",
    "exercise_type",
    "movement_role",
    "movement_pattern",
    "venue",
    "difficulty",
    "instruction_vi",
    "instruction_steps_vi",
    "common_mistakes_vi",
    "tips_vi",
    "notes_vi",
    "video_url",
    "gif_url",
    "image_url",
    "created_at",
    "updated_at",
]

HEADER_FILL = PatternFill("FF1F4E3D", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFFFF")


def _iso(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat(timespec="seconds")
    return str(value)


def _steps_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "\n".join(str(x).strip() for x in value if str(x).strip())
    if isinstance(value, str):
        raw = value.strip()
        if raw.startswith("["):
            try:
                parsed = json.loads(raw)
                if isinstance(parsed, list):
                    return "\n".join(str(x).strip() for x in parsed if str(x).strip())
            except json.JSONDecodeError:
                pass
        return raw
    return str(value)


def _json_list(value: Any) -> str:
    if not value:
        return ""
    if isinstance(value, list):
        return ", ".join(str(x) for x in value if str(x).strip())
    return str(value)


def _write_header(ws) -> None:
    ws.append(HEADER)
    for col, _ in enumerate(HEADER, start=1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"


def _set_widths(ws) -> None:
    wide = {
        "name_vi",
        "name_en",
        "instruction_vi",
        "instruction_steps_vi",
        "common_mistakes_vi",
        "tips_vi",
        "notes_vi",
        "video_url",
        "secondary_muscles",
        "equipment_vi",
    }
    for col, title in enumerate(HEADER, start=1):
        letter = get_column_letter(col)
        if title in wide:
            ws.column_dimensions[letter].width = 42
        else:
            ws.column_dimensions[letter].width = max(12, min(22, len(title) + 2))


def _wrap_text_columns(ws) -> None:
    wrap_names = {
        "instruction_vi",
        "instruction_steps_vi",
        "common_mistakes_vi",
        "tips_vi",
        "notes_vi",
    }
    wrap_cols = [HEADER.index(n) + 1 for n in wrap_names]
    for row_idx in range(2, ws.max_row + 1):
        for col in wrap_cols:
            ws.cell(row_idx, col).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row_idx].height = 64


def export_exercises(db: Session, out: Path) -> int:
    exercises = db.query(Exercise).order_by(Exercise.name_en, Exercise.id).all()
    muscles = {m.id: m for m in db.query(MuscleGroup).all()}
    equipment = {e.id: e for e in db.query(Equipment).all()}
    eq_by_ex: dict[int, list[int]] = defaultdict(list)
    for link in db.query(ExerciseEquipment).all():
        eq_by_ex[link.exercise_id].append(link.equipment_id)

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Bai_tap"
    _write_header(ws)

    for ex in exercises:
        mg = muscles.get(ex.muscle_group_id)
        eq_ids = eq_by_ex.get(ex.id, [])
        eq_rows = [equipment[i] for i in eq_ids if i in equipment]
        eq_rows.sort(key=lambda e: e.slug)
        ws.append(
            [
                ex.id,
                bool(ex.is_active),
                ex.name_vi,
                ex.name_en or "",
                mg.slug if mg else "",
                mg.name_vi if mg else "",
                _json_list(ex.secondary_muscles),
                ", ".join(e.slug for e in eq_rows),
                ", ".join(e.name_vi for e in eq_rows),
                ex.exercise_type or "",
                ex.movement_role or "",
                ex.movement_pattern or "",
                ex.venue or "",
                ex.difficulty,
                ex.instruction_vi or "",
                _steps_text(ex.instruction_steps_vi),
                ex.common_mistakes_vi or "",
                ex.tips_vi or "",
                ex.notes_vi or "",
                ex.video_url or "",
                ex.gif_url or "",
                ex.image_url or "",
                _iso(ex.created_at),
                _iso(ex.updated_at),
            ]
        )

    _set_widths(ws)
    _wrap_text_columns(ws)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    active = sum(1 for e in exercises if e.is_active)
    print(f"  active={active} inactive={len(exercises) - active}")
    return len(exercises)


def main() -> Path:
    import shutil

    out_dir = Path(__file__).resolve().parents[1] / "exports"
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    out = out_dir / f"exercises_catalog_{stamp}.xlsx"
    db = SessionLocal()
    try:
        n = export_exercises(db, out)
        desktop = Path.home() / "Desktop"
        desktop.mkdir(parents=True, exist_ok=True)
        dest = desktop / out.name
        shutil.copy2(out, dest)
        print(f"Exported {n} exercises")
        print(out)
        print(dest)
        return out
    finally:
        db.close()


if __name__ == "__main__":
    main()
