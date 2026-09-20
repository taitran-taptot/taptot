"""Export all familiarization curricula to an editable Excel workbook."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.core.database import SessionLocal, engine
from app.core.migrations import ensure_familiarization_exercises
from app.services.workout_generation.bmi import (
    BMI_LABEL_VI,
    bmi_band,
    compute_bmi,
    goal_from_bmi_band,
)
from app.services.workout_generation.familiarization_curriculum import (
    _PATH_LABELS,
    generate_familiarization_workout,
)
from app.services.workout_generation.fitness_standards import PATH_META

HEADER = [
    "Lộ trình",
    "Giới tính",
    "Tuổi",
    "Chiều cao (cm)",
    "Cân (kg)",
    "BMI",
    "Nhóm BMI",
    "Mục tiêu",
    "Địa điểm",
    "Dụng cụ",
    "Calo gợi ý (kcal/ngày)",
    "Ngày",
    "Tuần",
    "Tiêu đề ngày",
    "Vai trò",
    "Ghi chú ngày",
    "Thứ tự",
    "Phần",
    "Exercise ID",
    "Bài tập (VI)",
    "Bài tập (EN)",
    "Hiệp",
    "Số lần / thời gian",
    "Nghỉ (giây)",
    "Ghi chú bài",
]

WIDTHS = [
    28, 10, 8, 14, 10, 8, 22, 12, 10, 14, 18,
    8, 8, 28, 14, 40, 8, 10, 12, 36, 32, 8, 28, 12, 50,
]
HEADER_FILL = PatternFill("FF1F4E3D", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFFFF")
GENDERS = (("male", "Nam"), ("female", "Nữ"))
GOAL_VI = {
    "gain_weight": "Tăng cân",
    "maintain": "Giữ cân",
    "lose_weight": "Giảm cân",
}
PATH_SHEETS = {
    "first_push_pull": "Nhap_mon",
    "basic_foundation": "Xay_suc_manh",
    "advanced_foundation": "Nen_tang_NC",
}
BMI_SAMPLES: dict[str, tuple[int, tuple[tuple[str, float], ...]]] = {
    "female": (
        160,
        (
            ("underweight", 45),
            ("normal", 54),
            ("overweight", 62),
            ("obese_1", 70),
            ("obese_2", 80),
        ),
    ),
    "male": (
        170,
        (
            ("underweight", 50),
            ("normal", 62),
            ("overweight", 70),
            ("obese_1", 80),
            ("obese_2", 95),
        ),
    ),
}


def _style_header(ws) -> None:
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL


def _set_widths(ws) -> None:
    for i, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"


def _load_exercise_map(db: Session) -> dict[int, tuple[str, str | None]]:
    rows = db.execute(text("SELECT id, name_vi, name_en FROM exercises")).fetchall()
    return {int(r[0]): (r[1] or "", r[2]) for r in rows}


def _payload(path: str, gender: str, height_cm: float, weight_kg: float) -> dict:
    band = bmi_band(compute_bmi(weight_kg, height_cm) or 22.0)
    return {
        "generation_mode": "familiarization",
        "familiarization_path": path,
        "gender": gender,
        "age": 28,
        "height_cm": height_cm,
        "weight_kg": weight_kg,
        "activity": "light",
        "goal": goal_from_bmi_band(band),
        "sessions_per_week": 3,
        "session_minutes": 45,
        "experience_level": 1,
        "location": "home",
        "equipment_list": ["pull-up-bar"],
        "food_ids": [],
        "fitness_baseline": {},
    }


def _input_prefix(
    label: str,
    gender_label: str,
    payload: dict,
    insights: dict,
) -> list:
    bmi = compute_bmi(payload["weight_kg"], payload["height_cm"])
    band = bmi_band(bmi or 22.0)
    weight_goal = insights.get("weight_goal") or {}
    return [
        label,
        gender_label,
        payload["age"],
        payload["height_cm"],
        payload["weight_kg"],
        bmi,
        BMI_LABEL_VI.get(band, band),
        GOAL_VI.get(str(payload.get("goal") or ""), payload.get("goal") or ""),
        "Nhà",
        "Xà đơn",
        weight_goal.get("daily_kcal") or "",
    ]


def _empty_exercise_tail() -> list:
    return ["", "", "", "(Ngày nghỉ)", "", "", "", "", ""]


def main() -> Path:
    ensure_familiarization_exercises(engine)

    wb = Workbook()
    overview = wb.active
    overview.title = "Tổng_quan"
    overview.append(
        [
            "Lộ trình (key)",
            "Tên VI",
            "Giới tính",
            "Tuổi",
            "Chiều cao (cm)",
            "Cân (kg)",
            "BMI",
            "Nhóm BMI",
            "Mục tiêu",
            "Calo gợi ý",
            "Số ngày",
            "Số buổi tập",
            "Tóm tắt",
            "Chu kỳ",
        ]
    )
    _style_header(overview)

    master = wb.create_sheet("Tất_cả_lịch", 1)
    master.append(HEADER)
    _style_header(master)

    path_sheets: dict[str, object] = {}
    for meta in PATH_META:
        path = str(meta["key"])
        sheet = wb.create_sheet(PATH_SHEETS.get(path, path)[:31])
        sheet.append(HEADER)
        _style_header(sheet)
        path_sheets[path] = sheet

    db = SessionLocal()
    try:
        exercise_map = _load_exercise_map(db)
        for meta in PATH_META:
            path = str(meta["key"])
            label = _PATH_LABELS.get(path, path)
            sheet = path_sheets[path]
            for gender_key, gender_label in GENDERS:
                height_cm, samples = BMI_SAMPLES[gender_key]
                for _band_key, weight_kg in samples:
                    payload = _payload(path, gender_key, height_cm, weight_kg)
                    result = generate_familiarization_workout(
                        db, None, payload, persist=False
                    )
                    insights = result.get("insights") or {}
                    days = result.get("days") or []
                    train_count = sum(1 for day in days if day.get("exercises"))
                    overview_block = insights.get("overview") or {}
                    weight_goal = insights.get("weight_goal") or {}
                    prefix = _input_prefix(label, gender_label, payload, insights)
                    overview.append(
                        [
                            path,
                            label,
                            gender_label,
                            payload["age"],
                            payload["height_cm"],
                            payload["weight_kg"],
                            prefix[5],
                            prefix[6],
                            prefix[7],
                            weight_goal.get("daily_kcal") or "",
                            len(days),
                            train_count,
                            overview_block.get("summary_vi") or "",
                            overview_block.get("periodization_vi") or "",
                        ]
                    )

                    for day in days:
                        day_num = int(day.get("day_number") or 0)
                        week = ((day_num - 1) // 7) + 1 if day_num else ""
                        title = day.get("title_vi") or ""
                        role = day.get("split_role") or ""
                        day_notes = day.get("notes_vi") or ""
                        exercises = day.get("exercises") or []
                        day_core = [day_num, week, title, role, day_notes]
                        if not exercises:
                            row = [*prefix, *day_core, *_empty_exercise_tail()]
                            master.append(row)
                            sheet.append(row)
                            continue
                        for idx, ex in enumerate(exercises, start=1):
                            eid = int(ex.get("exercise_id") or 0)
                            name_vi, name_en = exercise_map.get(eid, (f"#{eid}", None))
                            row = [
                                *prefix,
                                *day_core,
                                idx,
                                ex.get("section") or "main",
                                eid,
                                name_vi,
                                name_en or "",
                                ex.get("sets"),
                                ex.get("reps"),
                                ex.get("rest_seconds"),
                                ex.get("notes_vi") or "",
                            ]
                            master.append(row)
                            sheet.append(row)

            _set_widths(sheet)

        overview.append([])
        overview.append(["Ghi chú"])
        overview.append(
            [f"Xuất lúc: {datetime.now().isoformat(timespec='seconds')}"]
        )
        overview.append(
            [
                "Ma trận: 3 lộ trình × Nam/Nữ × 5 nhóm BMI (gầy, bình thường, "
                "thừa cân, béo phì I, béo phì II+). Tuổi 28, nhà + xà đơn, 3 buổi × 45 phút."
            ]
        )
        overview.append(
            ["Bạn có thể sửa trực tiếp các cột Hiệp / Số lần / Nghỉ / Ghi chú bài."]
        )
        for col in range(1, 15):
            overview.column_dimensions[get_column_letter(col)].width = (
                22 if col < 10 else 50
            )
        overview.freeze_panes = "A2"
        _set_widths(master)

        out_dir = Path(__file__).resolve().parents[1] / "exports"
        out_dir.mkdir(parents=True, exist_ok=True)
        out = out_dir / f"lich_lam_quen_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        wb.save(out)
        return out
    finally:
        db.close()


if __name__ == "__main__":
    path = main()
    print(path)
