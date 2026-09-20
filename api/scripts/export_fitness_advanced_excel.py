"""Export all Thử thách thể lực nâng cao cases to Excel, with entry and graduation standards."""

from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.core.database import SessionLocal, engine
from app.core.migrations import ensure_familiarization_exercises
from app.services.workout_generation.familiarization_curriculum import (
    expand_familiarization_weeks,
)
from app.services.workout_generation.fitness_advanced_curriculum import (
    SESSION_MINUTES,
    STANDARDS,
    WEEKS,
    build_fitness_advanced_week_templates,
    resolve_catalog,
)
from app.services.workout_generation.fitness_standards import familiarization_catalog

HEADER = [
    "Case",
    "Giới tính",
    "Buổi/tuần",
    "Khối",
    "Ngày",
    "Tuần",
    "Tiêu đề ngày",
    "Vai trò",
    "Ghi chú ngày",
    "Thứ tự",
    "Phần",
    "Exercise ID",
    "Bài tập (VI)",
    "Bài tập (EN)",
    "Hiệp",
    "Số lần / thời gian",
    "Nghỉ (giây)",
    "Ghi chú bài",
]

WIDTHS = [18, 10, 12, 16, 8, 8, 28, 14, 48, 8, 10, 12, 36, 32, 8, 28, 12, 50]
HEADER_FILL = PatternFill("FF1F4E3D", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFFFF")
ENTRY_FILL = PatternFill("FFDEEAF6", fill_type="solid")
EXIT_FILL = PatternFill("FFE2F0D9", fill_type="solid")
GENDERS = (("male", "Nam"), ("female", "Nữ"))
SESSIONS = (4, 5, 6)
ROLE_VI = {
    "a": "Đẩy + plank",
    "b": "Chạy dễ / tempo",
    "c": "Kéo xà + squat",
    "d": "Chạy chất lượng",
    "e": "Volume phụ",
    "f": "Đẩy nhẹ + plank",
    "recovery": "Nghỉ",
    "test": "Tốt nghiệp",
}
METRIC_VI = {
    "push": "Chống đẩy",
    "pull": "Kéo xà",
    "squat": "Squat",
    "plank": "Plank",
    "run": "Chạy 10 phút",
}


def _style_header(ws) -> None:
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL


def _set_widths(ws, widths: list[int] | None = None) -> None:
    for i, width in enumerate(widths or WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"


def _plank_vi(seconds: int) -> str:
    return f"{seconds // 60}:{seconds % 60:02d}"


def _run_vi(meters: int) -> str:
    return f"{meters / 1000:.1f} km".replace(".", ",")


def _block_label(week: int) -> str:
    if week <= 4:
        return "1 · Volume"
    if week <= 8:
        return "2 · Mật độ"
    if week <= 11:
        return "3 · Gần max"
    return "12 · Giảm tải / tốt nghiệp"


def _sheet_name(gender_label: str, sessions: int) -> str:
    slug = "Nam" if gender_label == "Nam" else "Nu"
    return f"{slug}_{sessions}buoi"


def _case_label(gender_label: str, sessions: int) -> str:
    return f"{gender_label} · {sessions} buổi"


def _load_exercise_map(db: Session) -> dict[int, tuple[str, str | None]]:
    rows = db.execute(text("SELECT id, name_vi, name_en FROM exercises")).fetchall()
    return {int(r[0]): (r[1] or "", r[2]) for r in rows}


def _empty_exercise_tail(role: str) -> list:
    marker = "(Ngày tốt nghiệp — test camera)" if role == "test" else "(Ngày nghỉ)"
    return ["", "", "", marker, "", "", "", "", ""]


def _write_standards(ws) -> None:
    headers = [
        "Loại",
        "Giới tính",
        "Môn",
        "Chuẩn",
        "Đạt",
        "Khá",
        "Giỏi",
        "Ghi chú",
    ]
    ws.append(headers)
    _style_header(ws)
    catalog = familiarization_catalog()["standards"]
    for gender_key, gender_label in GENDERS:
        for row in catalog[gender_key]["advanced"]:
            ws.append(
                [
                    "Đầu vào — cửa ra nền tảng nâng cao",
                    gender_label,
                    row["label_vi"],
                    row["display_vi"],
                    "",
                    "",
                    "",
                    "Test camera trước khi tạo lịch. Pass = đạt sàn khoảng này.",
                ]
            )
            for cell in ws[ws.max_row]:
                cell.fill = ENTRY_FILL
        bands = STANDARDS[gender_key]
        dat, kha, gioi = bands["dat"], bands["kha"], bands["gioi"]
        rows = [
            ("push", dat["push"], kha["push"], gioi["push"], lambda n: str(n)),
            ("pull", dat["pull"], kha["pull"], gioi["pull"], lambda n: str(n)),
            ("squat", dat["squat"], kha["squat"], gioi["squat"], lambda n: str(n)),
            ("plank", dat["plank"], kha["plank"], gioi["plank"], _plank_vi),
            ("run", dat["run_m"], kha["run_m"], gioi["run_m"], _run_vi),
        ]
        for key, d, k, g, fmt in rows:
            ws.append(
                [
                    "Đầu ra — tốt nghiệp ngày cuối tuần 12",
                    gender_label,
                    METRIC_VI[key],
                    f"Đạt {fmt(d)} · Khá {fmt(k)} · Giỏi {fmt(g)}",
                    fmt(d),
                    fmt(k),
                    fmt(g),
                    "Test chính thức. Pass sản phẩm = đủ 5/5 Đạt. Nghỉ 2 phút giữa các bài camera.",
                ]
            )
            for cell in ws[ws.max_row]:
                cell.fill = EXIT_FILL
    ws.append([])
    ws.append(
        [
            "Pass thử thách = đủ 5/5 Đạt ở test tốt nghiệp. Đầu vào không dùng bảng Đạt này."
        ]
    )
    _set_widths(ws, [38, 12, 16, 48, 12, 12, 12, 70])


def main() -> Path:
    ensure_familiarization_exercises(engine)

    wb = Workbook()
    standards = wb.active
    standards.title = "Tieu_chuan"
    _write_standards(standards)

    overview = wb.create_sheet("Tong_quan", 1)
    overview.append(
        [
            "Case",
            "Giới tính",
            "Buổi/tuần",
            "Phút/buổi",
            "Số tuần",
            "Số ngày",
            "Số buổi tập",
            "Ngày tốt nghiệp",
            "Tóm tắt",
        ]
    )
    _style_header(overview)

    master = wb.create_sheet("Tat_ca_lich", 2)
    master.append(HEADER)
    _style_header(master)

    case_sheets: dict[tuple[str, int], object] = {}
    for gender_key, gender_label in GENDERS:
        for sessions in SESSIONS:
            sheet = wb.create_sheet(_sheet_name(gender_label, sessions))
            sheet.append(HEADER)
            _style_header(sheet)
            case_sheets[(gender_key, sessions)] = sheet

    db = SessionLocal()
    try:
        exercise_map = _load_exercise_map(db)
        catalog = resolve_catalog(db)
        for gender_key, gender_label in GENDERS:
            for sessions in SESSIONS:
                case = _case_label(gender_label, sessions)
                templates = build_fitness_advanced_week_templates(
                    db,
                    {"gender": gender_key, "sessions_per_week": sessions},
                    catalog=catalog,
                )
                days = [day.model_dump() for day in expand_familiarization_weeks(templates)]
                train_count = sum(1 for day in days if day.get("exercises"))
                test_day = next(
                    (day for day in days if (day.get("split_role") or "") == "test"),
                    {},
                )
                overview.append(
                    [
                        case,
                        gender_label,
                        sessions,
                        SESSION_MINUTES,
                        WEEKS,
                        len(days),
                        train_count,
                        test_day.get("title_vi") or "",
                        (
                            f"12 tuần cố định. Đầu vào = cửa ra nền tảng nâng cao. "
                            f"Ngày {test_day.get('day_number') or 84} test chính thức Đạt 5/5."
                        ),
                    ]
                )
                sheet = case_sheets[(gender_key, sessions)]
                for day in days:
                    day_num = int(day.get("day_number") or 0)
                    week = ((day_num - 1) // 7) + 1 if day_num else ""
                    role = str(day.get("split_role") or "").lower()
                    prefix = [
                        case,
                        gender_label,
                        sessions,
                        _block_label(int(week) if week else 1),
                        day_num,
                        week,
                        day.get("title_vi") or "",
                        ROLE_VI.get(role, role),
                        day.get("notes_vi") or "",
                    ]
                    exercises = day.get("exercises") or []
                    if not exercises:
                        row = [*prefix, *_empty_exercise_tail(role)]
                        master.append(row)
                        sheet.append(row)
                        continue
                    for idx, ex in enumerate(exercises, start=1):
                        eid = int(ex.get("exercise_id") or 0)
                        name_vi, name_en = exercise_map.get(eid, (f"#{eid}", None))
                        row = [
                            *prefix,
                            idx,
                            ex.get("section") or "main",
                            eid,
                            name_vi,
                            name_en or "",
                            ex.get("sets"),
                            ex.get("reps"),
                            ex.get("rest_seconds"),
                            ex.get("notes_vi") or "",
                        ]
                        master.append(row)
                        sheet.append(row)
                _set_widths(sheet)

        overview.append([])
        overview.append(["Ghi chú"])
        overview.append([f"Xuất lúc: {datetime.now().isoformat(timespec='seconds')}"])
        overview.append(
            [
                "Ma trận lịch tập: Nam/Nữ × 4–6 buổi/tuần = 6 case. "
                "Lịch không đổi theo BMI; BMI chỉ ráp thực đơn khi tạo lịch trên app."
            ]
        )
        overview.append(
            [
                "Sheet Tieu_chuan: đầu vào = cửa ra nền tảng nâng cao (cấp 3); "
                "đầu ra = test tốt nghiệp ngày cuối tuần 12 (Đạt / Khá / Giỏi)."
            ]
        )
        for col in range(1, 10):
            overview.column_dimensions[get_column_letter(col)].width = 22 if col < 8 else 70
        overview.freeze_panes = "A2"
        _set_widths(master)

        out_dir = Path(__file__).resolve().parents[1] / "exports"
        out_dir.mkdir(parents=True, exist_ok=True)
        out = out_dir / f"lich_the_luc_nang_cao_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        wb.save(out)
        return out
    finally:
        db.close()


if __name__ == "__main__":
    path = main()
    print(path)
