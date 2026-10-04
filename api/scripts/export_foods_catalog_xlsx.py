"""Export active public food catalog to Excel (api/exports/)."""

from __future__ import annotations

import json
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

API_DIR = Path(__file__).resolve().parents[1]
if str(API_DIR) not in sys.path:
    sys.path.insert(0, str(API_DIR))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.core.database import SessionLocal
from app.models.entities import Food, FoodAlias, FoodCategory, FoodPortion

HEADER = [
    "id",
    "slug",
    "name_vi",
    "name_en",
    "category_slug",
    "category_name_vi",
    "food_kind",
    "prep_state",
    "status",
    "is_common",
    "is_verified",
    "serving_size",
    "serving_grams",
    "calories",
    "protein_g",
    "carbs_g",
    "fat_g",
    "fiber_g",
    "sugar_g",
    "sodium_mg",
    "kcal_100g",
    "protein_100g",
    "carbs_100g",
    "fat_100g",
    "fiber_100g",
    "sugar_100g",
    "sodium_100mg",
    "alcohol_100g",
    "tags",
    "vitamins_json",
    "macro_roles",
    "meal_slots",
    "aliases",
    "portions",
    "source_ref",
    "confidence",
    "yield_factor",
    "density_g_per_ml",
    "ai_eligible",
    "ai_priority",
    "default_for_ai",
    "is_complete_meal",
    "region_slug",
    "province_id",
    "description_vi",
    "image_url",
    "created_at",
]

HEADER_FILL = PatternFill("FF1F4E3D", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFFFF")


def _json_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, dict)):
        if not value:
            return ""
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return str(value)


def _csv_list(value: Any) -> str:
    if not value:
        return ""
    if isinstance(value, list):
        return ", ".join(str(v) for v in value if v is not None and str(v).strip())
    return str(value)


def _format_portions(portions: list[FoodPortion]) -> str:
    if not portions:
        return ""
    parts: list[str] = []
    for p in sorted(portions, key=lambda x: (x.sort_order, x.id)):
        flag = "*" if p.is_default else ""
        parts.append(f"{p.label_vi}={p.grams:g}g{flag}")
    return "; ".join(parts)


def _write_header(ws) -> None:
    ws.append(HEADER)
    for col, _ in enumerate(HEADER, start=1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"


def _set_widths(ws) -> None:
    for col, title in enumerate(HEADER, start=1):
        letter = get_column_letter(col)
        if title in {"description_vi", "vitamins_json", "aliases", "portions", "tags"}:
            width = 40
        elif title in {"name_vi", "name_en", "slug", "image_url", "source_ref"}:
            width = 28
        elif title in {"category_name_vi", "category_slug"}:
            width = 22
        else:
            width = max(12, min(18, len(title) + 2))
        ws.column_dimensions[letter].width = width


def _row_for(
    food: Food,
    category: FoodCategory | None,
    aliases: list[str],
    portions: list[FoodPortion],
) -> list[Any]:
    created = food.created_at
    created_s = ""
    if created is not None:
        created_s = (
            created.isoformat(timespec="seconds")
            if hasattr(created, "isoformat")
            else str(created)
        )
    return [
        food.id,
        food.slug,
        food.name_vi,
        food.name_en or "",
        category.slug if category else "",
        category.name_vi if category else "",
        food.food_kind,
        food.prep_state or "",
        food.status,
        bool(food.is_common),
        bool(food.is_verified),
        food.serving_size,
        food.serving_grams,
        food.calories,
        food.protein_g,
        food.carbs_g,
        food.fat_g,
        food.fiber_g,
        food.sugar_g,
        food.sodium_mg,
        food.kcal_100g,
        food.protein_100g,
        food.carbs_100g,
        food.fat_100g,
        food.fiber_100g,
        food.sugar_100g,
        food.sodium_100mg,
        food.alcohol_100g,
        _csv_list(food.tags),
        _json_cell(food.vitamins_json),
        _csv_list(food.macro_roles),
        _csv_list(food.meal_slots),
        ", ".join(aliases),
        _format_portions(portions),
        food.source_ref or "",
        food.confidence or "",
        food.yield_factor,
        food.density_g_per_ml,
        bool(food.ai_eligible),
        food.ai_priority,
        bool(food.default_for_ai),
        food.is_complete_meal if food.is_complete_meal is not None else "",
        food.region_slug or "",
        food.province_id or "",
        food.description_vi or "",
        food.image_url or "",
        created_s,
    ]


def _write_foods_sheet(
    ws,
    foods: list[Food],
    categories: dict[int, FoodCategory],
    aliases_by_food: dict[int, list[str]],
    portions_by_food: dict[int, list[FoodPortion]],
) -> None:
    _write_header(ws)
    for food in foods:
        ws.append(
            _row_for(
                food,
                categories.get(food.category_id) if food.category_id else None,
                aliases_by_food.get(food.id, []),
                portions_by_food.get(food.id, []),
            )
        )
    _set_widths(ws)


def export_foods(db: Session, out: Path) -> int:
    foods = (
        db.query(Food)
        .filter(Food.owner_user_id.is_(None), Food.status == "active")
        .all()
    )
    foods = sorted(
        foods,
        key=lambda f: (
            f.category_id is None,
            f.category_id or 0,
            (f.name_vi or "").lower(),
            f.id,
        ),
    )

    cat_ids = {f.category_id for f in foods if f.category_id is not None}
    categories = (
        {
            c.id: c
            for c in db.query(FoodCategory).filter(FoodCategory.id.in_(cat_ids)).all()
        }
        if cat_ids
        else {}
    )

    food_ids = [f.id for f in foods]
    aliases_by_food: dict[int, list[str]] = defaultdict(list)
    portions_by_food: dict[int, list[FoodPortion]] = defaultdict(list)
    if food_ids:
        for a in (
            db.query(FoodAlias)
            .filter(FoodAlias.food_id.in_(food_ids))
            .order_by(FoodAlias.id)
            .all()
        ):
            aliases_by_food[a.food_id].append(a.alias)
        for p in (
            db.query(FoodPortion)
            .filter(FoodPortion.food_id.in_(food_ids))
            .order_by(FoodPortion.sort_order, FoodPortion.id)
            .all()
        ):
            portions_by_food[p.food_id].append(p)

    ingredients = [f for f in foods if (f.food_kind or "ingredient") != "dish"]
    dishes = [f for f in foods if (f.food_kind or "") == "dish"]

    wb = Workbook()
    ws_all = wb.active
    assert ws_all is not None
    ws_all.title = "Tat_ca"
    _write_foods_sheet(ws_all, foods, categories, aliases_by_food, portions_by_food)

    ws_ing = wb.create_sheet("Nguyen_lieu")
    _write_foods_sheet(ws_ing, ingredients, categories, aliases_by_food, portions_by_food)

    ws_dish = wb.create_sheet("Mon_truyen_thong_gia_dinh")
    _write_foods_sheet(ws_dish, dishes, categories, aliases_by_food, portions_by_food)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"  ingredients={len(ingredients)} dishes={len(dishes)}")
    return len(foods)


def main() -> Path:
    out_dir = Path(__file__).resolve().parents[1] / "exports"
    out = out_dir / f"foods_catalog_active_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    db = SessionLocal()
    try:
        n = export_foods(db, out)
        print(f"Exported {n} active public foods")
        return out
    finally:
        db.close()


if __name__ == "__main__":
    path = main()
    print(path)
