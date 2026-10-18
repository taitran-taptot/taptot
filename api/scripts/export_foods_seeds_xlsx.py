"""Export food warehouse seeds to Excel (no database)."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
SEEDS = ROOT / "seeds"

HEADER = [
    "seed_file",
    "slug",
    "name_vi",
    "name_en",
    "category_slug",
    "food_kind",
    "prep_state",
    "serving_size",
    "serving_grams",
    "calories",
    "protein_g",
    "carbs_g",
    "fat_g",
    "kcal_100g",
    "protein_100g",
    "carbs_100g",
    "fat_100g",
    "sugar_g",
    "sugar_100g",
    "sodium_mg",
    "sodium_100mg",
    "source_ref",
    "confidence",
    "tags",
    "macro_roles",
    "aliases",
    "portions",
    "region_slug",
    "province_id",
    "description_vi",
    "image_url",
]

HEADER_FILL = PatternFill("FF1F4E3D", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFFFF")


def _csv_list(value: Any) -> str:
    if not value:
        return ""
    if isinstance(value, list):
        if value and isinstance(value[0], dict):
            parts = []
            for row in value:
                label = row.get("label_vi") or ""
                grams = row.get("grams")
                parts.append(f"{label}={grams}g")
            return "; ".join(parts)
        return ", ".join(str(v) for v in value if v is not None and str(v).strip())
    return str(value)


def _foods_from(path: Path) -> list[dict]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(payload, list):
        return [x for x in payload if isinstance(x, dict)]
    if isinstance(payload, dict):
        foods = payload.get("foods")
        if isinstance(foods, list):
            default_cat = ""
            cat = payload.get("category")
            if isinstance(cat, dict):
                default_cat = str(cat.get("slug") or "")
            out = []
            for row in foods:
                if not isinstance(row, dict):
                    continue
                if default_cat and not row.get("category_slug"):
                    row = {**row, "category_slug": default_cat}
                out.append(row)
            return out
    return []


def _scale(row: dict) -> dict:
    serving = float(row.get("serving_grams") or 100) or 100.0
    scale = serving / 100.0
    kcal_100 = row.get("kcal_100g")
    protein_100 = row.get("protein_100g")
    carbs_100 = row.get("carbs_100g")
    fat_100 = row.get("fat_100g")
    if kcal_100 is None and row.get("calories") is not None:
        inv = 100.0 / serving
        kcal_100 = round(float(row["calories"]) * inv, 2)
        protein_100 = round(float(row.get("protein_g") or 0) * inv, 2)
        carbs_100 = round(float(row.get("carbs_g") or 0) * inv, 2)
        fat_100 = round(float(row.get("fat_g") or 0) * inv, 2)
    calories = row.get("calories")
    if calories is None and kcal_100 is not None:
        calories = round(float(kcal_100) * scale, 2)
        row = {
            **row,
            "calories": calories,
            "protein_g": round(float(protein_100 or 0) * scale, 2),
            "carbs_g": round(float(carbs_100 or 0) * scale, 2),
            "fat_g": round(float(fat_100 or 0) * scale, 2),
        }
    row["kcal_100g"] = kcal_100
    row["protein_100g"] = protein_100
    row["carbs_100g"] = carbs_100
    row["fat_100g"] = fat_100
    if not row.get("food_kind"):
        row["food_kind"] = "dish" if row.get("is_complete_meal") else "ingredient"
    return row


def _row(seed_file: str, food: dict) -> list[Any]:
    food = _scale(food)
    return [
        seed_file,
        food.get("slug") or "",
        food.get("name_vi") or "",
        food.get("name_en") or "",
        food.get("category_slug") or "",
        food.get("food_kind") or "",
        food.get("prep_state") or "",
        food.get("serving_size") or "",
        food.get("serving_grams"),
        food.get("calories"),
        food.get("protein_g"),
        food.get("carbs_g"),
        food.get("fat_g"),
        food.get("kcal_100g"),
        food.get("protein_100g"),
        food.get("carbs_100g"),
        food.get("fat_100g"),
        food.get("sugar_g"),
        food.get("sugar_100g"),
        food.get("sodium_mg"),
        food.get("sodium_100mg"),
        food.get("source_ref") or "",
        food.get("confidence") or "",
        _csv_list(food.get("tags")),
        _csv_list(food.get("macro_roles")),
        _csv_list(food.get("aliases")),
        _csv_list(food.get("portions")),
        food.get("region_slug") or "",
        food.get("province_id") or "",
        food.get("description_vi") or "",
        food.get("image_url") or "",
    ]


def _write_sheet(ws, rows: list[tuple[str, dict]]) -> None:
    ws.append(HEADER)
    for col, _ in enumerate(HEADER, start=1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"
    for seed_file, food in rows:
        ws.append(_row(seed_file, food))
    for col, title in enumerate(HEADER, start=1):
        letter = get_column_letter(col)
        width = 36 if title in {"description_vi", "aliases", "portions", "tags"} else 18
        ws.column_dimensions[letter].width = width


def main() -> Path:
    sources = [
        ("foods_catalog_v2.json", SEEDS / "foods_catalog_v2.json"),
        ("grain_nut_foods.json", SEEDS / "grain_nut_foods.json"),
        ("foods_cooking_pantry.json", SEEDS / "foods_cooking_pantry.json"),
        ("foods_traditional_dishes.json", SEEDS / "foods_traditional_dishes.json"),
    ]
    all_rows: list[tuple[str, dict]] = []
    for name, path in sources:
        if not path.is_file():
            continue
        for food in _foods_from(path):
            if name == "foods_traditional_dishes.json":
                food = {**food, "food_kind": food.get("food_kind") or "dish"}
            all_rows.append((name, food))

    ingredients = [
        item for item in all_rows if str(item[1].get("food_kind") or "ingredient") != "dish"
    ]
    dishes = [item for item in all_rows if str(item[1].get("food_kind") or "") == "dish"]
    if not dishes:
        dishes = [item for item in all_rows if item[0] == "foods_traditional_dishes.json"]

    wb = Workbook()
    ws_all = wb.active
    assert ws_all is not None
    ws_all.title = "Tat_ca"
    _write_sheet(ws_all, all_rows)
    _write_sheet(wb.create_sheet("Nguyen_lieu"), ingredients)
    _write_sheet(wb.create_sheet("Mon_truyen_thong"), dishes)

    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    out = Path(r"C:\Users\Tran Tai\Desktop\backupdata") / f"kho_thuc_pham_{stamp}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"all={len(all_rows)} ingredients={len(ingredients)} dishes={len(dishes)}")
    print(out)
    return out


if __name__ == "__main__":
    main()
