"""Orthogonal home-plan audit export (4 waves). Engine-only, persist=False."""

from __future__ import annotations

import argparse
import json
import sys
import time
from itertools import product
from pathlib import Path
from typing import Any, Iterable

ROOT = Path(__file__).resolve().parents[2]
API = ROOT / "api"
if str(API) not in sys.path:
    sys.path.insert(0, str(API))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.database import SessionLocal
from app.schemas.plans import PlanDayIn, PlanExerciseIn
from app.services.workout_generation.meal_engine import generate_meals
from app.services.workout_generation.nutrition_targets import estimate_targets
from app.services.workout_generation.service import generate_workout
from app.services.workout_generation.session_duration import estimate_day_minutes

KITS: dict[str, tuple[str, ...]] = {
    "dumbbell": ("dumbbell",),
    "resistance-band": ("resistance-band",),
    "gymnastic-rings": ("gymnastic-rings", "pull-up-bar"),
}
FOCUS_ONE = ("nguc", "lung", "mong", "chan", "bung", "vai", "tay")
FOCUS_MULTI = (
    ("nguc", "vai"),
    ("lung", "tay"),
    ("chan", "mong"),
    ("nguc", "tay"),
    ("nguc", "vai", "tay"),
    ("lung", "vai", "tay"),
)

CASE_HEADERS = [
    "wave",
    "case_id",
    "kit",
    "gender",
    "level",
    "sessions",
    "minutes",
    "focus",
    "amrap",
    "age",
    "height_cm",
    "weight_kg",
    "activity",
    "goal",
    "kg_per_week",
    "tdee",
    "target_kcal",
    "week_code",
    "split_roles",
    "elapsed_s",
    "error",
]
DAY_HEADERS = [
    "case_id",
    "kit",
    "gender",
    "level",
    "minutes",
    "focus",
    "day_number",
    "split_role",
    "main_count",
    "est_minutes",
]
EX_HEADERS = [
    "case_id",
    "kit",
    "gender",
    "level",
    "minutes",
    "focus",
    "day_number",
    "split_role",
    "section",
    "name_vi",
    "sets",
    "reps",
    "rest_seconds",
    "load_note_vi",
    "notes_vi",
]
MEAL_HEADERS = [
    "case_id",
    "day_number",
    "split_role",
    "meal_type",
    "food_id",
    "name_vi",
    "calories",
    "protein_g",
    "carbs_g",
    "fat_g",
    "day_kcal",
    "target_kcal",
]


def _profile(gender: str) -> dict[str, Any]:
    female = gender == "female"
    return {
        "age": 28,
        "height_cm": 162.0 if female else 175.0,
        "weight_kg": 62.0 if female else 82.0,
        "activity": "moderate",
    }


def _baseline(kit: str, amrap: int) -> dict[str, Any]:
    n = int(amrap)
    if kit == "dumbbell":
        return {
            "test_kit": "dumbbell",
            "db_press_reps": n,
            "db_press_kg": 10,
            "db_row_reps": n,
            "db_row_kg": 10,
            "goblet_reps": n,
            "goblet_kg": 16,
            "pushups_max": n,
            "squats_max": n,
        }
    if kit == "resistance-band":
        return {
            "test_kit": "band",
            "band_level": "blue",
            "band_level_push": "blue",
            "band_level_pull": "blue",
            "band_level_squat": "blue",
            "pushups_max": n,
            "pullups_max": n,
            "squats_max": n,
        }
    return {
        "test_kit": "bar_rings",
        "pushups_max": max(n, 12),
        "pullups_max": max(5, min(n, 12)),
        "squats_max": n,
    }


def _payload(spec: dict[str, Any]) -> dict[str, Any]:
    gender = spec["gender"]
    prof = _profile(gender)
    body = {
        "goal": "lose_weight",
        "gender": gender,
        "sessions_per_week": spec["sessions"],
        "session_minutes": spec["minutes"],
        "location": "home",
        "focus_areas": list(spec["focus"]),
        "extra_goals": [],
        "equipment_list": list(KITS[spec["kit"]]),
        "food_ids": [],
        "experience_level": spec["level"],
        "ai_suggest_equipment": False,
        "no_equipment": False,
        "ai_suggest_foods": True,
        "fitness_baseline": _baseline(spec["kit"], spec["amrap"]),
        "duration_weeks": 4,
        "kg_per_week": 0.5,
        "challenge_100_days": True,
        "curriculum_12_weeks": False,
        **prof,
    }
    return body


def iter_wave(wave: int) -> list[dict[str, Any]]:
    kits = tuple(KITS)
    out: list[dict[str, Any]] = []
    if wave == 1:
        for kit, gender, minutes in product(kits, ("male", "female"), (45, 60)):
            out.append(
                {
                    "wave": 1,
                    "kit": kit,
                    "gender": gender,
                    "level": 2,
                    "sessions": 4,
                    "minutes": minutes,
                    "focus": ("nguc",),
                    "amrap": 16,
                }
            )
        for kit, amrap in product(kits, (8, 25)):
            out.append(
                {
                    "wave": 1,
                    "kit": kit,
                    "gender": "male",
                    "level": 2,
                    "sessions": 4,
                    "minutes": 60,
                    "focus": ("nguc",),
                    "amrap": amrap,
                }
            )
    elif wave == 2:
        for kit, gender, focus in product(kits, ("male", "female"), FOCUS_ONE):
            out.append(
                {
                    "wave": 2,
                    "kit": kit,
                    "gender": gender,
                    "level": 2,
                    "sessions": 4,
                    "minutes": 60,
                    "focus": (focus,),
                    "amrap": 16,
                }
            )
    elif wave == 3:
        for kit, gender, level, sessions in product(
            kits, ("male", "female"), (1, 2, 3), (3, 4, 5)
        ):
            out.append(
                {
                    "wave": 3,
                    "kit": kit,
                    "gender": gender,
                    "level": level,
                    "sessions": sessions,
                    "minutes": 60,
                    "focus": ("nguc",),
                    "amrap": 16,
                }
            )
        for gender, sessions in product(("male", "female"), (2, 6)):
            out.append(
                {
                    "wave": 3,
                    "kit": "dumbbell",
                    "gender": gender,
                    "level": 2,
                    "sessions": sessions,
                    "minutes": 60,
                    "focus": ("nguc",),
                    "amrap": 16,
                }
            )
    elif wave == 4:
        for kit, focus in product(kits, FOCUS_MULTI):
            out.append(
                {
                    "wave": 4,
                    "kit": kit,
                    "gender": "male",
                    "level": 2,
                    "sessions": 4,
                    "minutes": 60,
                    "focus": focus,
                    "amrap": 16,
                }
            )
    else:
        raise SystemExit(f"unknown wave {wave}")
    for spec in out:
        spec["case_id"] = _case_id(spec)
    return out


def _case_id(spec: dict[str, Any]) -> str:
    focus = "+".join(spec["focus"]) or "none"
    return (
        f"w{spec['wave']}_{spec['kit']}_{spec['gender'][0]}"
        f"_L{spec['level']}_{spec['sessions']}d_{spec['minutes']}m"
        f"_{focus}_a{spec['amrap']}"
    )


def _done_ids(path: Path) -> set[str]:
    if not path.exists():
        return set()
    out: set[str] = set()
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line:
            continue
        try:
            rec = json.loads(line)
        except json.JSONDecodeError:
            continue
        cid = str(rec.get("case_id") or (rec.get("case") or {}).get("case_id") or "")
        if cid:
            out.add(cid)
    return out


def _load_note(ex: dict[str, Any]) -> str:
    for row in ex.get("set_prescriptions") or []:
        note = (row or {}).get("load_note_vi") if isinstance(row, dict) else None
        if note:
            return str(note)
    return ""


def _ex_reps(ex: dict[str, Any]) -> str:
    rx = ex.get("set_prescriptions") or []
    if rx and isinstance(rx[0], dict) and rx[0].get("reps") not in (None, ""):
        return str(rx[0]["reps"])
    return str(ex.get("reps") or "")


def _est_minutes(day: dict[str, Any]) -> float | None:
    try:
        exercises = []
        for raw in day.get("exercises") or []:
            data = {
                k: raw.get(k)
                for k in (
                    "exercise_id",
                    "sets",
                    "reps",
                    "rest_seconds",
                    "section",
                    "sort_order",
                    "set_prescriptions",
                )
            }
            exercises.append(PlanExerciseIn.model_validate(data))
        dummy = PlanDayIn(day_number=int(day.get("day_number") or 1), exercises=exercises)
        return round(estimate_day_minutes(dummy), 1)
    except Exception:
        return None


def _week1_days(days: list[dict[str, Any]], sessions: int) -> list[dict[str, Any]]:
    training = [
        d
        for d in days
        if str(d.get("split_role") or "").strip().lower() not in {"", "rest", "recovery"}
    ]
    return training[: max(1, int(sessions))]


def _append_ckpt(path: Path, rec: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as fh:
        fh.write(json.dumps(rec, ensure_ascii=False) + "\n")


def _write_xlsx(path: Path, cases: list[dict], days: list[dict], exs: list[dict], meals: list[dict]) -> None:
    wb = Workbook()
    fill = PatternFill(start_color="FF1F4E3D", end_color="FF1F4E3D", fill_type="solid")
    font = Font(bold=True, color="FFFFFFFF")

    def _sheet(title: str, headers: list[str], rows: list[dict]) -> None:
        ws = wb.create_sheet(title)
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
        for row in rows:
            ws.append([row.get(h) for h in headers])
        for i, _h in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(i)].width = 16

    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    _sheet("Cases", CASE_HEADERS, cases)
    _sheet("Days", DAY_HEADERS, days)
    _sheet("Exercises", EX_HEADERS, exs)
    _sheet("Meals", MEAL_HEADERS, meals)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def run_case(db, spec: dict[str, Any]) -> dict[str, Any]:
    payload = _payload(spec)
    nut = estimate_targets(payload)
    t0 = time.perf_counter()
    err = None
    result: dict[str, Any] = {}
    try:
        result = generate_workout(db, None, payload, persist=False)
    except Exception as exc:
        err = f"{type(exc).__name__}: {exc}"
    elapsed = round(time.perf_counter() - t0, 2)
    days = _week1_days(list(result.get("days") or []), spec["sessions"]) if not err else []
    insights = result.get("insights") or {}
    nutrition = insights.get("nutrition") or {}
    roles = [str(d.get("split_role") or "") for d in days]
    case_row = {
        "wave": spec["wave"],
        "case_id": spec["case_id"],
        "kit": spec["kit"],
        "gender": spec["gender"],
        "level": spec["level"],
        "sessions": spec["sessions"],
        "minutes": spec["minutes"],
        "focus": "+".join(spec["focus"]),
        "amrap": spec["amrap"],
        "age": payload["age"],
        "height_cm": payload["height_cm"],
        "weight_kg": payload["weight_kg"],
        "activity": payload["activity"],
        "goal": payload["goal"],
        "kg_per_week": payload["kg_per_week"],
        "tdee": (nut.tdee if nut else nutrition.get("tdee")),
        "target_kcal": (nut.target_calories if nut else nutrition.get("target_calories")),
        "week_code": result.get("week_code"),
        "split_roles": ",".join(roles),
        "elapsed_s": elapsed,
        "error": err,
    }
    day_rows: list[dict[str, Any]] = []
    ex_rows: list[dict[str, Any]] = []
    meal_rows: list[dict[str, Any]] = []
    target = case_row["target_kcal"]
    for day in days:
        mains = [e for e in (day.get("exercises") or []) if (e.get("section") or "main") == "main"]
        day_rows.append(
            {
                "case_id": spec["case_id"],
                "kit": spec["kit"],
                "gender": spec["gender"],
                "level": spec["level"],
                "minutes": spec["minutes"],
                "focus": case_row["focus"],
                "day_number": day.get("day_number"),
                "split_role": day.get("split_role"),
                "main_count": len(mains),
                "est_minutes": _est_minutes(day),
            }
        )
        for ex in day.get("exercises") or []:
            ex_rows.append(
                {
                    "case_id": spec["case_id"],
                    "kit": spec["kit"],
                    "gender": spec["gender"],
                    "level": spec["level"],
                    "minutes": spec["minutes"],
                    "focus": case_row["focus"],
                    "day_number": day.get("day_number"),
                    "split_role": day.get("split_role"),
                    "section": ex.get("section"),
                    "name_vi": ex.get("name_vi"),
                    "sets": ex.get("sets"),
                    "reps": _ex_reps(ex),
                    "rest_seconds": ex.get("rest_seconds"),
                    "load_note_vi": _load_note(ex),
                    "notes_vi": ex.get("notes_vi"),
                }
            )
    if not err:
        meal_payload = {**payload, "ai_suggest_foods": True}
        try:
            meal_result = generate_meals(
                db, meal_payload, nut, plan_days=None, goal="lose_weight"
            )
        except Exception:
            meal_result = None
        foods = getattr(meal_result, "foods_by_id", None) or {}
        for i, tmpl in enumerate(list(getattr(meal_result, "templates", None) or [])[:2]):
            day = days[i] if i < len(days) else {}
            totals = getattr(tmpl, "totals", None) or {}
            day_kcal = totals.get("calories")
            if day_kcal is None:
                day_kcal = 0.0
                for meal in tmpl.meals:
                    food = foods.get(int(meal.food_id))
                    serv = float(meal.servings or 1)
                    if food:
                        day_kcal += float(food.calories or 0) * serv
            for meal in tmpl.meals:
                food = foods.get(int(meal.food_id))
                serv = float(meal.servings or 1)
                meal_rows.append(
                    {
                        "case_id": spec["case_id"],
                        "day_number": day.get("day_number") or i + 1,
                        "split_role": day.get("split_role"),
                        "meal_type": meal.meal_type,
                        "food_id": meal.food_id,
                        "name_vi": getattr(food, "name_vi", None),
                        "calories": round(float(food.calories or 0) * serv, 1) if food else None,
                        "protein_g": round(float(food.protein_g or 0) * serv, 1) if food else None,
                        "carbs_g": round(float(food.carbs_g or 0) * serv, 1) if food else None,
                        "fat_g": round(float(food.fat_g or 0) * serv, 1) if food else None,
                        "day_kcal": round(float(day_kcal or 0), 1),
                        "target_kcal": target,
                    }
                )
    return {"case": case_row, "days": day_rows, "exercises": ex_rows, "meals": meal_rows}


def _load_rows(ckpt: Path) -> tuple[list[dict], list[dict], list[dict], list[dict]]:
    cases: list[dict] = []
    days: list[dict] = []
    exs: list[dict] = []
    meals: list[dict] = []
    if not ckpt.exists():
        return cases, days, exs, meals
    for line in ckpt.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        rec = json.loads(line)
        if rec.get("case"):
            cases.append(rec["case"])
            days.extend(rec.get("days") or [])
            exs.extend(rec.get("exercises") or [])
            meals.extend(rec.get("meals") or [])
    return cases, days, exs, meals


def main(argv: Iterable[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--wave", default="all", help="1,2,3,4,all")
    parser.add_argument(
        "--out",
        default=str(ROOT / "docs" / "audit_home_plans.xlsx"),
    )
    parser.add_argument(
        "--ckpt",
        default=str(ROOT / "docs" / "audit_home_plans_checkpoint.jsonl"),
    )
    args = parser.parse_args(list(argv) if argv is not None else None)
    waves: list[int]
    if str(args.wave).strip().lower() == "all":
        waves = [1, 2, 3, 4]
    else:
        waves = [int(args.wave)]
    specs = [s for w in waves for s in iter_wave(w)]
    ckpt = Path(args.ckpt)
    out_path = Path(args.out)
    done = _done_ids(ckpt)
    cases, days, exs, meals = _load_rows(ckpt)
    db = SessionLocal()
    try:
        pending = [s for s in specs if s["case_id"] not in done]
        print(f"{len(specs)} cases, {len(pending)} pending, {len(done)} done", flush=True)
        for i, spec in enumerate(pending, start=1):
            print(f"[{i}/{len(pending)}] {spec['case_id']}", flush=True)
            rec = run_case(db, spec)
            _append_ckpt(ckpt, rec)
            cases.append(rec["case"])
            days.extend(rec["days"])
            exs.extend(rec["exercises"])
            meals.extend(rec["meals"])
            if i % 5 == 0 or i == len(pending):
                _write_xlsx(out_path, cases, days, exs, meals)
        _write_xlsx(out_path, cases, days, exs, meals)
    finally:
        db.close()
    print(f"wrote {out_path}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
