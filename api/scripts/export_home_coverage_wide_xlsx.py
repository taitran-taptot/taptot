"""Wide 3-kit challenge coverage -> one Excel, 14 weeks. persist=False."""

from __future__ import annotations

import json
import sys
import time
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[2]
API = ROOT / "api"
if str(API) not in sys.path:
    sys.path.insert(0, str(API))

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import text

from app.core.database import SessionLocal
from app.services.periodization import expand_plan_days_for_weeks
from app.services.workout_generation.load_estimate import (
    BAND_LEVEL_NAME_VI,
    default_band_level,
    default_db_kg,
)
from app.services.workout_generation.service import generate_workout

OUT_XLSX = ROOT / "docs" / "testing" / "home-coverage-3kit-2026-10-28.xlsx"
CKPT = ROOT / "docs" / "testing" / "home-coverage-3kit.ckpt.jsonl"

KITS = ("band", "db", "rings")
KIT_VI = {"band": "dây kháng lực", "db": "tạ đơn", "rings": "xà + vòng"}
KIT_EQUIP = {
    "band": ["resistance-band"],
    "db": ["dumbbell"],
    "rings": ["pull-up-bar", "gymnastic-rings"],
}
KIT_TEST = {"band": "band", "db": "dumbbell", "rings": "bar_rings"}

BODIES = {
    "female": ((155, 58), (160, 68), (165, 72)),
    "male": ((170, 68), (175, 75), (180, 85)),
}
MID_BODY = {"female": (160, 68), "male": (175, 75)}

CASE_HEADERS = [
    "kit",
    "case_id",
    "gioi",
    "trinh_do",
    "muc_tieu",
    "height_cm",
    "weight_kg",
    "kg_100_days",
    "tdee",
    "target_kcal",
    "buoi",
    "phut",
    "focus",
    "pushups",
    "pullups",
    "squats",
    "plank_s",
    "mau_day",
    "press_kg",
    "row_kg",
    "goblet_kg",
    "week_code",
    "giai_thich",
    "error",
]
EX_HEADERS = [
    "kit",
    "case_id",
    "tuan",
    "pha",
    "buoi",
    "split",
    "section",
    "ten",
    "sets",
    "reps",
    "rest",
    "load_note_vi",
]

GENDER_VI = {"male": "Nam", "female": "Nữ"}
LEVEL_VI = {1: "người mới", 2: "đã có thói quen", 3: "tập đều"}
GOAL_VI = {"lose_weight": "giảm cân", "gain_weight": "tăng cân"}
SPLIT_WHY = {
    ("male", 3): "PPL vì nam 3 buổi",
    ("male", 4): "ULUL vì nam 4 buổi",
    ("male", 5): "PPLUL vì nam 5 buổi",
    ("female", 3): "LUL vì nữ 3 buổi",
    ("female", 4): "LPPL vì nữ 4 buổi",
    ("female", 5): "LULU kèm cardio-core vì nữ 5 buổi",
}


def _kg_week(goal: str) -> float:
    return 0.5 if goal == "lose_weight" else 0.35


def _kg_100(goal: str) -> float:
    return round(_kg_week(goal) * 14, 1)


def _fmt_kg(value: Any) -> str:
    if value is None or value == "":
        return ""
    try:
        num = float(value)
    except (TypeError, ValueError):
        return str(value)
    text = f"{num:.1f}".rstrip("0").rstrip(".")
    return text.replace(".", ",")


def _slot_kg(gender: str, level: int, slot: str) -> float:
    return float(default_db_kg(gender=gender, experience_level=level, slot=slot))


def _band_vi(color: str | None) -> str:
    if not color:
        return "mặc định theo giới+trình độ"
    return BAND_LEVEL_NAME_VI.get(color, color)


def _core_baseline(
    kit: str,
    gender: str,
    level: int,
    *,
    reps: int | None = 15,
    plank: int | None = 20,
    band_color: str | None = None,
    press_kg: float | None = None,
    row_kg: float | None = None,
    goblet_kg: float | None = None,
    pull: int | None = None,
) -> dict[str, Any]:
    base: dict[str, Any] = {"test_kit": KIT_TEST[kit]}
    if kit == "band":
        if band_color:
            base["band_level"] = band_color
            base["band_level_push"] = band_color
            base["band_level_pull"] = band_color
            base["band_level_squat"] = band_color
        if reps is not None:
            base["pushups_max"] = reps
            base["pullups_max"] = reps
            base["squats_max"] = reps
        if plank is not None:
            base["plank_seconds"] = plank
        return base
    if kit == "db":
        if press_kg is None and reps is not None:
            press_kg = _slot_kg(gender, level, "press")
        if row_kg is None and reps is not None:
            row_kg = _slot_kg(gender, level, "row")
        if goblet_kg is None and reps is not None:
            goblet_kg = _slot_kg(gender, level, "squat")
        if press_kg is not None:
            base["db_press_kg"] = press_kg
        if row_kg is not None:
            base["db_row_kg"] = row_kg
        if goblet_kg is not None:
            base["goblet_kg"] = goblet_kg
        if reps is not None:
            base["db_press_reps"] = reps
            base["db_row_reps"] = reps
            base["goblet_reps"] = reps
            base["pushups_max"] = reps
            base["squats_max"] = reps
        if plank is not None:
            base["plank_seconds"] = plank
        return base
    if reps is not None:
        base["pushups_max"] = reps
        base["squats_max"] = reps
        base["pullups_max"] = pull if pull is not None else reps
    elif pull is not None:
        base["pullups_max"] = pull
    if plank is not None:
        base["plank_seconds"] = plank
    return base


def _blank_spec(kit: str, **kw: Any) -> dict[str, Any]:
    gender = kw.get("gender", "female")
    height, weight = MID_BODY[gender]
    spec: dict[str, Any] = {
        "kit": kit,
        "wave": "overlay",
        "gender": gender,
        "height_cm": height,
        "weight_kg": weight,
        "level": 1,
        "goal": "lose_weight",
        "sessions": 3,
        "minutes": 45,
        "focus": [],
        "pushups": 15,
        "pullups": 15,
        "squats": 15,
        "plank_s": 20,
        "band_color": "",
        "press_kg": "",
        "row_kg": "",
        "goblet_kg": "",
    }
    spec.update(kw)
    return spec


def _fill_db_kg_fields(spec: dict[str, Any]) -> None:
    if spec["kit"] != "db":
        return
    if spec.get("press_kg") not in ("", None):
        return
    if spec.get("pushups") in ("", None):
        return
    spec["press_kg"] = _slot_kg(spec["gender"], int(spec["level"]), "press")
    spec["row_kg"] = _slot_kg(spec["gender"], int(spec["level"]), "row")
    spec["goblet_kg"] = _slot_kg(spec["gender"], int(spec["level"]), "squat")


def iter_specs() -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for kit in KITS:
        for gender in ("male", "female"):
            for height, weight in BODIES[gender]:
                for level in (1, 2, 3):
                    for goal in ("lose_weight", "gain_weight"):
                        for sessions in (3, 4, 5):
                            for minutes in (45, 60):
                                spec = {
                                    "kit": kit,
                                    "case_id": (
                                        f"core_{kit}_{gender[0]}_h{height}_"
                                        f"L{level}_{goal[:4]}_s{sessions}_m{minutes}"
                                    ),
                                    "wave": "core",
                                    "gender": gender,
                                    "height_cm": height,
                                    "weight_kg": weight,
                                    "level": level,
                                    "goal": goal,
                                    "sessions": sessions,
                                    "minutes": minutes,
                                    "focus": [],
                                    "pushups": 15,
                                    "pullups": 15,
                                    "squats": 15,
                                    "plank_s": 20,
                                    "band_color": "",
                                    "press_kg": "",
                                    "row_kg": "",
                                    "goblet_kg": "",
                                    "baseline": _core_baseline(kit, gender, level),
                                }
                                _fill_db_kg_fields(spec)
                                out.append(spec)
        _add_overlays(out, kit)
    return out


def _add_overlays(out: list[dict[str, Any]], kit: str) -> None:
    def add(**kw: Any) -> None:
        spec = _blank_spec(kit, **kw)
        gender = spec["gender"]
        level = int(spec["level"])
        reps = spec.get("pushups")
        reps_i = None if reps in ("", None) else int(reps)
        pull = spec.get("pullups")
        pull_i = None if pull in ("", None) else int(pull)
        plank = spec.get("plank_s")
        plank_i = None if plank in ("", None) else int(plank)
        color = spec.get("band_color") or None
        press = spec.get("press_kg")
        press_f = None if press in ("", None) else float(press)
        row = spec.get("row_kg")
        row_f = None if row in ("", None) else float(row)
        goblet = spec.get("goblet_kg")
        goblet_f = None if goblet in ("", None) else float(goblet)
        if "baseline" not in spec:
            spec["baseline"] = _core_baseline(
                kit,
                gender,
                level,
                reps=reps_i,
                plank=plank_i,
                band_color=color if kit == "band" else None,
                press_kg=press_f,
                row_kg=row_f,
                goblet_kg=goblet_f,
                pull=pull_i if kit == "rings" else None,
            )
        _fill_db_kg_fields(spec)
        out.append(spec)

    for n in (8, 15, 25, 35):
        add(
            case_id=f"ov_{kit}_amrap_F_L1_s3_m45_lose_{n}",
            gender="female",
            level=1,
            goal="lose_weight",
            sessions=3,
            minutes=45,
            pushups=n,
            pullups=n,
            squats=n,
            band_color="yellow" if kit == "band" else "",
        )
        add(
            case_id=f"ov_{kit}_amrap_M_L3_s4_m60_gain_{n}",
            gender="male",
            level=3,
            goal="gain_weight",
            sessions=4,
            minutes=60,
            pushups=n,
            pullups=n,
            squats=n,
            band_color="black" if kit == "band" else "",
        )
    for n in (8, 35):
        add(
            case_id=f"ov_{kit}_amrap_F_L2_s3_m45_{n}",
            gender="female",
            level=2,
            sessions=3,
            minutes=45,
            pushups=n,
            pullups=n,
            squats=n,
        )
        add(
            case_id=f"ov_{kit}_amrap_M_L2_s4_m60_{n}",
            gender="male",
            level=2,
            sessions=4,
            minutes=60,
            pushups=n,
            pullups=n,
            squats=n,
        )
    if kit == "band":
        for color in ("yellow", "red", "blue", "green", "black"):
            add(
                case_id=f"ov_{kit}_color_F_L2_s3_m45_{color}",
                gender="female",
                level=2,
                band_color=color,
            )
    elif kit == "db":
        for press in (2.5, 5.0, 7.5, 10.0, 12.5):
            label = str(press).replace(".", "p")
            add(
                case_id=f"ov_{kit}_kg_F_L2_s3_m45_{label}",
                gender="female",
                level=2,
                press_kg=press,
                row_kg=press,
                goblet_kg=press + 2.5,
            )
    else:
        for pull in (0, 3, 8, 15, 25):
            add(
                case_id=f"ov_{kit}_pull_F_L2_s3_m45_{pull}",
                gender="female",
                level=2,
                pullups=pull,
            )
    for sec in (20, 40, 60):
        add(
            case_id=f"ov_{kit}_plank_M_L2_s3_m45_{sec}s",
            gender="male",
            level=2,
            plank_s=sec,
        )
    for key, focus in (
        ("none", []),
        ("bcm", ["bung", "chan", "mong"]),
        ("nguc", ["nguc"]),
    ):
        add(
            case_id=f"ov_{kit}_focus_F_L2_s3_m45_{key}",
            gender="female",
            level=2,
            focus=focus,
        )
    add(
        case_id=f"ov_{kit}_notest_M_L1_s3_m45",
        gender="male",
        level=1,
        pushups="",
        pullups="",
        squats="",
        plank_s=None,
        press_kg="",
        row_kg="",
        goblet_kg="",
        band_color="",
        baseline={"test_kit": KIT_TEST[kit]},
    )


def payload_for(spec: dict[str, Any]) -> dict[str, Any]:
    goal = spec["goal"]
    return {
        "goal": goal,
        "gender": spec["gender"],
        "sessions_per_week": spec["sessions"],
        "session_minutes": spec["minutes"],
        "location": "home",
        "focus_areas": list(spec.get("focus") or []),
        "extra_goals": [],
        "equipment_list": list(KIT_EQUIP[spec["kit"]]),
        "food_ids": [],
        "experience_level": spec["level"],
        "ai_suggest_equipment": False,
        "no_equipment": False,
        "ai_suggest_foods": False,
        "fitness_baseline": dict(spec["baseline"]),
        "duration_weeks": 14,
        "kg_per_week": _kg_week(goal),
        "challenge_100_days": True,
        "curriculum_12_weeks": False,
        "age": 28,
        "height_cm": float(spec["height_cm"]),
        "weight_kg": float(spec["weight_kg"]),
        "activity": "moderate",
    }


def _amrap_shared(n: Any) -> str:
    if n in ("", None):
        return "Không điền test nên compound 8–12, isolation 12–18."
    n = int(n)
    if n < 10:
        return f"Test {n} cái (<10) nên hạ một nấc tải / cửa sổ rep hẹp."
    if n > 30:
        return f"Test {n} cái (>30) nên lên hai nấc tải / nới rep."
    if n > 20:
        return f"Test {n} cái (>20) nên lên một nấc tải / nới rep."
    return f"Test {n} cái nằm 10–20 nên giữ tải giữa vùng."


def _kit_note(spec: dict[str, Any]) -> str:
    kit = spec["kit"]
    gender = spec["gender"]
    level = int(spec["level"])
    n = spec.get("pushups")
    if kit == "band":
        color = spec.get("band_color") or None
        default = default_band_level(gender=gender, experience_level=level)
        if color:
            return f"Dây xuất phát {_band_vi(color)}. " + _amrap_shared(n)
        return (
            f"Không ghi màu test nên mặc định {_band_vi(default)} "
            f"(nữ mới vàng; nữ hoặc mới đỏ; còn lại xanh dương). "
            + _amrap_shared(n)
        )
    if kit == "db":
        d_press = spec.get("press_kg") or _slot_kg(gender, level, "press")
        d_squat = spec.get("goblet_kg") or _slot_kg(gender, level, "squat")
        if spec.get("pushups") in ("", None):
            return (
                f"Không ghi kg test nên mặc định đẩy/chèo {_fmt_kg(d_press)} kg, "
                f"goblet {_fmt_kg(d_squat)} kg "
                f"(nữ mới 2,5/5; nữ hoặc mới 5/7,5; còn lại 7,5/10). "
                "Isolation/warmup hoặc người mới có thể hạ thêm một nấc kg. "
                + _amrap_shared(None)
            )
        return (
            f"Kg test đẩy/chèo {_fmt_kg(d_press)} kg, goblet {_fmt_kg(d_squat)} kg. "
            "Isolation/warmup hoặc người mới có thể hạ thêm một nấc kg. "
            + _amrap_shared(n)
        )
    pull = spec.get("pullups")
    bits = ["Kit xà đơn + vòng treo, không stamp kg tạ.", _amrap_shared(n)]
    if pull in ("", None):
        bits.append("Không ghi kéo xà nên slot kéo theo thiếu test.")
    else:
        p = int(pull)
        if p <= 0:
            bits.append("Kéo xà 0 cái nên ưu tiên chèo vòng / biến thể dễ hơn.")
        elif p < 8:
            bits.append(f"Kéo xà {p} cái (thấp) nên dễ ra chèo vòng hơn kéo xà đầy đủ.")
        else:
            bits.append(f"Kéo xà {p} cái nên đủ điều kiện bài kéo xà.")
    return " ".join(bits)


def giai_thich(spec: dict[str, Any], *, tdee: Any, target: Any, week_code: str) -> str:
    gender = spec["gender"]
    level = int(spec["level"])
    goal = spec["goal"]
    sessions = int(spec["sessions"])
    minutes = int(spec["minutes"])
    kg100 = _kg_100(goal)
    huong = "giảm" if goal == "lose_weight" else "tăng"
    parts = [
        f"{GENDER_VI[gender]} {int(spec['height_cm'])} cm, {int(spec['weight_kg'])} kg, "
        f"kit {KIT_VI[spec['kit']]}, "
        f"{huong} {str(kg100).replace('.', ',')} kg / 100 ngày"
        + (f", TDEE {tdee} kcal, mục tiêu {target} kcal/ngày" if tdee and target else "")
        + ".",
        f"Split {SPLIT_WHY[(gender, sessions)]}.",
        (
            "Trần 3 hiệp vì người mới; bài test pha 2 được 4 hiệp, pha 1 và 3 giữ 3."
            if level <= 1
            else "Trần 4 hiệp vì đã có thói quen hoặc tập đều; bài test pha 1/2/3 là 3/4/3."
        ),
        "Tuần 4, 8, 14 giảm hiệp (làm tròn 0,6, sàn 2), nghỉ lâu hơn, giữ mức tải.",
    ]
    if minutes <= 45:
        if level <= 1:
            parts.append(
                "45 phút: pha 3 có thể cộng hiệp ưu tiên nhưng người mới không vượt trần 3."
            )
        else:
            parts.append("45 phút: pha 3 có thể cộng 1 hiệp cho vùng ưu tiên (trần 4).")
    else:
        parts.append("60 phút: không cộng hiệp ưu tiên trên lịch thử thách.")
    if goal == "lose_weight":
        parts.append("Giảm cân: giữ khoảng số cái, không ramp rep theo tuần.")
    else:
        parts.append("Tăng cân: cho phép tăng số cái theo tuần.")
    parts.append(_kit_note(spec))
    plank = spec.get("plank_s")
    if plank:
        parts.append(f"Plank khoảng 80% của {plank} giây.")
    else:
        parts.append("Không test plank nên giữ khoảng mặc định.")
    focus = spec.get("focus") or []
    if focus:
        parts.append("Ưu tiên " + ", ".join(focus) + ".")
    else:
        parts.append("Không chọn nhóm ưu tiên.")
    if week_code:
        parts.append(f"Mã tuần engine: {week_code}.")
    return " ".join(parts)


def load_catalog(db) -> dict[int, dict[str, Any]]:
    rows = db.execute(
        text("SELECT id, name_en, name_vi FROM exercises")
    ).mappings().all()
    return {int(r["id"]): dict(r) for r in rows}


def _exercise_name(ex: dict[str, Any], catalog: dict[int, dict[str, Any]]) -> str:
    try:
        eid = int(ex.get("exercise_id"))
    except (TypeError, ValueError):
        eid = None
    cat = catalog.get(eid) if eid is not None else None
    return str(
        (cat or {}).get("name_vi")
        or ex.get("name_vi")
        or (cat or {}).get("name_en")
        or ex.get("name_en")
        or ""
    )


def _load_note(ex: dict[str, Any]) -> str:
    rx = ex.get("set_prescriptions") or []
    if rx and isinstance(rx[0], dict):
        return str(rx[0].get("load_note_vi") or "")
    return str(ex.get("notes_vi") or "")


def _week_of(day_number: int, sessions: int) -> int:
    return (max(1, int(day_number)) - 1) // max(1, int(sessions)) + 1


def _phase_of(week: int) -> int:
    if week <= 4:
        return 1
    if week <= 8:
        return 2
    return 3


def load_ckpt(path: Path) -> dict[str, dict[str, Any]]:
    latest: dict[str, dict[str, Any]] = {}
    if not path.is_file():
        return latest
    with path.open(encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            try:
                row = json.loads(line)
            except json.JSONDecodeError:
                continue
            cid = row.get("case_id")
            if cid:
                latest[str(cid)] = row
    return latest


def append_ckpt(path: Path, row: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as fh:
        fh.write(json.dumps(row, ensure_ascii=False) + "\n")


def run_case(db, spec: dict[str, Any], catalog: dict[int, dict[str, Any]]) -> dict[str, Any]:
    payload = payload_for(spec)
    t0 = time.perf_counter()
    try:
        result = generate_workout(db, None, payload, persist=False, commit=False)
    except Exception as exc:  # noqa: BLE001
        return {
            "error": str(exc),
            "elapsed_s": round(time.perf_counter() - t0, 2),
            "case": None,
            "exercises": [],
        }
    insights = result.get("insights") or {}
    nut = insights.get("nutrition") or {}
    tdee = nut.get("tdee")
    target = nut.get("target_calories")
    week_code = str(result.get("week_code") or insights.get("week_code") or "")
    sessions = int(result.get("sessions_actual") or spec["sessions"])
    expanded = expand_plan_days_for_weeks(
        result.get("days") or [],
        14,
        is_dict=True,
        experience_level=spec["level"],
        week_templates=insights.get("week_templates"),
        curriculum=True,
        goal=spec["goal"],
    )
    kit = spec["kit"]
    cid = spec["case_id"]
    ex_rows: list[dict[str, Any]] = []
    for day in expanded:
        week = _week_of(int(day.get("day_number") or 1), sessions)
        session_in_week = ((int(day.get("day_number") or 1) - 1) % sessions) + 1
        for ex in day.get("exercises") or []:
            ex_rows.append(
                {
                    "kit": kit,
                    "case_id": cid,
                    "tuan": week,
                    "pha": _phase_of(week),
                    "buoi": session_in_week,
                    "split": day.get("split_role") or "",
                    "section": ex.get("section") or "main",
                    "ten": _exercise_name(ex, catalog),
                    "sets": ex.get("sets"),
                    "reps": ex.get("reps"),
                    "rest": ex.get("rest_seconds"),
                    "load_note_vi": _load_note(ex),
                }
            )
    case_row = {
        "kit": kit,
        "case_id": cid,
        "gioi": GENDER_VI[spec["gender"]],
        "trinh_do": LEVEL_VI[spec["level"]],
        "muc_tieu": GOAL_VI[spec["goal"]],
        "height_cm": int(spec["height_cm"]),
        "weight_kg": int(spec["weight_kg"]),
        "kg_100_days": _kg_100(spec["goal"]),
        "tdee": tdee,
        "target_kcal": target,
        "buoi": spec["sessions"],
        "phut": spec["minutes"],
        "focus": "+".join(spec.get("focus") or []) or "",
        "pushups": spec.get("pushups") if spec.get("pushups") != "" else "",
        "pullups": spec.get("pullups") if spec.get("pullups") != "" else "",
        "squats": spec.get("squats") if spec.get("squats") != "" else "",
        "plank_s": spec.get("plank_s") if spec.get("plank_s") is not None else "",
        "mau_day": spec.get("band_color") or "",
        "press_kg": spec.get("press_kg") if spec.get("press_kg") != "" else "",
        "row_kg": spec.get("row_kg") if spec.get("row_kg") != "" else "",
        "goblet_kg": spec.get("goblet_kg") if spec.get("goblet_kg") != "" else "",
        "week_code": week_code,
        "giai_thich": giai_thich(spec, tdee=tdee, target=target, week_code=week_code),
        "error": "",
    }
    return {
        "error": None,
        "elapsed_s": round(time.perf_counter() - t0, 2),
        "case": case_row,
        "exercises": ex_rows,
    }


def write_xlsx(cases: list[dict[str, Any]], exercises: list[dict[str, Any]], path: Path) -> None:
    wb = Workbook(write_only=True)
    fill = PatternFill(start_color="FF1F4E3D", end_color="FF1F4E3D", fill_type="solid")
    font = Font(bold=True, color="FFFFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")

    def header_row(ws, headers: list[str]) -> None:
        cells = []
        for title in headers:
            cell = WriteOnlyCell(ws, value=title)
            cell.fill = fill
            cell.font = font
            cells.append(cell)
        ws.append(cells)

    ws_c = wb.create_sheet("Cases")
    header_row(ws_c, CASE_HEADERS)
    g_idx = CASE_HEADERS.index("giai_thich")
    for row in cases:
        vals = [row.get(h, "") for h in CASE_HEADERS]
        cells = []
        for i, val in enumerate(vals):
            cell = WriteOnlyCell(ws_c, value=val)
            if i == g_idx:
                cell.alignment = wrap
            cells.append(cell)
        ws_c.append(cells)

    ws_e = wb.create_sheet("BaiTap")
    header_row(ws_e, EX_HEADERS)
    for row in exercises:
        ws_e.append([row.get(h, "") for h in EX_HEADERS])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    specs = iter_specs()
    print(f"specs {len(specs)}")
    done = load_ckpt(CKPT)
    cases: list[dict[str, Any]] = []
    exercises: list[dict[str, Any]] = []
    db = SessionLocal()
    try:
        catalog = load_catalog(db)
        for i, spec in enumerate(specs, start=1):
            cid = spec["case_id"]
            if cid in done and done[cid].get("case") and any(
                (ex or {}).get("ten") for ex in (done[cid].get("exercises") or [])
            ):
                rec = done[cid]
                cases.append(rec["case"])
                exercises.extend(rec.get("exercises") or [])
                print(f"[{i}/{len(specs)}] skip {cid}")
                continue
            rec = run_case(db, spec, catalog)
            slim = {
                "case_id": cid,
                "error": rec.get("error"),
                "elapsed_s": rec.get("elapsed_s"),
                "case": rec.get("case"),
                "exercises": rec.get("exercises") or [],
            }
            append_ckpt(CKPT, slim)
            done[cid] = slim
            if rec.get("case"):
                cases.append(rec["case"])
                exercises.extend(rec.get("exercises") or [])
            else:
                cases.append(
                    {
                        "kit": spec["kit"],
                        "case_id": cid,
                        "error": rec.get("error") or "fail",
                        "giai_thich": rec.get("error") or "fail",
                    }
                )
            print(f"[{i}/{len(specs)}] {cid} {rec.get('elapsed_s')}s err={rec.get('error')}")
    finally:
        db.close()
    write_xlsx(cases, exercises, OUT_XLSX)
    print(f"wrote {OUT_XLSX} cases={len(cases)} exercises={len(exercises)}")


if __name__ == "__main__":
    main()
