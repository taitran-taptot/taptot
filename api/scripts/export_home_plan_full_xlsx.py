"""Full home-plan dump: inputs, overview copy, 14-week exercises. persist=False."""

from __future__ import annotations

import json
import sys
import time
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[2]
API = ROOT / "api"
SCRIPTS = Path(__file__).resolve().parent
for p in (str(API), str(SCRIPTS)):
    if p not in sys.path:
        sys.path.insert(0, p)

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

from app.core.database import SessionLocal
from app.services.periodization import expand_plan_days_for_weeks
from app.services.workout_generation.familiarization_curriculum import (
    generate_familiarization_workout,
)
from app.services.workout_generation.service import generate_workout
from export_home_coverage_wide_xlsx import (
    GENDER_VI,
    GOAL_VI,
    LEVEL_VI,
    MID_BODY,
    _exercise_name,
    _kg_100,
    _phase_of,
    _week_of,
    giai_thich,
    iter_specs,
    load_catalog,
    payload_for,
)

OUT_XLSX = ROOT / "docs" / "testing" / "home-plan-full-2026-10-29.xlsx"
CKPT = ROOT / "docs" / "testing" / "home-plan-full.ckpt.jsonl"

LUUY = [
    (
        "Kỹ thuật chuẩn",
        "Thực hiện động tác đúng kỹ thuật, kiểm soát chuyển động và sử dụng biên độ phù hợp.",
    ),
    (
        "Dinh dưỡng phù hợp mục tiêu",
        "Cung cấp đủ năng lượng, protein và các chất dinh dưỡng cần thiết.",
    ),
    (
        "Ngủ đủ và phục hồi tốt",
        "Hướng đến 7–9 giờ ngủ mỗi đêm, bố trí thời gian nghỉ phù hợp và điều chỉnh tải tập khi xuất hiện dấu hiệu mệt mỏi tích lũy.",
    ),
    (
        "Khởi động và đảm bảo an toàn",
        "Khởi động trước khi tập, làm nóng với mức tạ tăng dần, sử dụng dụng cụ đúng cách.",
    ),
    (
        "Theo dõi tiến bộ",
        "Ghi chép số lần lặp, số hiệp. Đánh giá định kỳ sức mạnh, cân nặng, số đo cơ thể hoặc các chỉ số phù hợp với mục tiêu.",
    ),
    (
        "Kiên trì và nhất quán",
        "Duy trì kết hợp tập luyện, dinh dưỡng và phục hồi. Đánh giá hiệu quả dựa trên tiến bộ thực tế.",
    ),
]

CASE_HEADERS = [
    "kit",
    "case_id",
    "gioi",
    "trinh_do",
    "muc_tieu",
    "height_cm",
    "weight_kg",
    "kg_100_days",
    "tdee",
    "target_kcal",
    "protein_g",
    "carbs_g",
    "fat_g",
    "buoi",
    "phut",
    "focus",
    "pushups",
    "pullups",
    "squats",
    "plank_s",
    "mau_day",
    "press_kg",
    "row_kg",
    "goblet_kg",
    "week_code",
    "giai_thich",
    "error",
]
TONG_HEADERS = [
    "kit",
    "case_id",
    "buoi",
    "phut",
    "nhom_buoi",
    "tdee",
    "target_kcal",
    "protein_g",
    "carbs_g",
    "fat_g",
    "summary_vi",
    "periodization_vi",
    "nutrition_vi",
    "recap_vi",
    "mission_vi",
    "weight_copy",
    "weight_kcal",
    "pha1_label",
    "pha1_weeks",
    "pha1_blurb",
    "pha1_rationale",
    "pha1_knowledge",
    "pha2_label",
    "pha2_weeks",
    "pha2_blurb",
    "pha2_rationale",
    "pha2_knowledge",
    "pha3_label",
    "pha3_weeks",
    "pha3_blurb",
    "pha3_rationale",
    "pha3_knowledge",
]
EX_HEADERS = [
    "kit",
    "case_id",
    "tuan",
    "pha",
    "buoi",
    "split",
    "section",
    "ten",
    "sets",
    "reps",
    "rest",
    "notes_vi",
    "load_note_vi",
    "hiep",
]
FAM_HEADERS = [
    "kind",
    "case_id",
    "gioi",
    "mission_vi",
    "outcome_vi",
    "nutrition_vi",
    "periodization_vi",
    "weight_copy",
    "weight_kcal",
    "ngay",
    "tuan",
    "split",
    "section",
    "ten",
    "sets",
    "reps",
    "rest",
    "notes_vi",
    "load_note_vi",
    "hiep",
]


def _blank(value: Any) -> Any:
    if value is None:
        return ""
    return value


def _sets_bits(ex: dict[str, Any]) -> tuple[str, str]:
    rx = ex.get("set_prescriptions") or []
    notes: list[str] = []
    parts: list[str] = []
    if isinstance(rx, list):
        for i, row in enumerate(rx, start=1):
            if not isinstance(row, dict):
                continue
            note = str(row.get("load_note_vi") or "").strip()
            if note and note not in notes:
                notes.append(note)
            reps = row.get("reps")
            bit = note
            if reps not in (None, ""):
                bit = f"{bit} x{reps}".strip()
            parts.append(f"H{i}:{bit}" if bit else f"H{i}")
    return (notes[0] if notes else "", " | ".join(parts))


def _phase_cell(meso: dict[str, Any] | None, key: str) -> str:
    if not meso:
        return ""
    if key == "weeks":
        weeks = [int(w) for w in (meso.get("weeks") or []) if str(w).isdigit() or isinstance(w, int)]
        if not weeks:
            return ""
        return f"{weeks[0]}-{weeks[-1]}"
    if key == "knowledge":
        refs = meso.get("knowledge_refs") or []
        labels = [str(r.get("label") or r.get("slug") or "") for r in refs if isinstance(r, dict)]
        return "; ".join(x for x in labels if x)
    return str(meso.get(key) or "")


def _week1_splits(days: list[dict[str, Any]], sessions: int) -> str:
    roles: list[str] = []
    for day in days:
        week = _week_of(int(day.get("day_number") or 1), sessions)
        if week != 1:
            break
        role = str(day.get("split_role") or "").strip()
        if role and role not in roles:
            roles.append(role)
    return " · ".join(roles)


def _exercise_rows(
    *,
    kit: str,
    case_id: str,
    days: list[dict[str, Any]],
    sessions: int,
    catalog: dict[int, dict[str, Any]],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for day in days:
        day_no = int(day.get("day_number") or 1)
        week = _week_of(day_no, sessions)
        session_in_week = ((day_no - 1) % max(1, sessions)) + 1
        for ex in day.get("exercises") or []:
            if not isinstance(ex, dict):
                continue
            load_note, hiep = _sets_bits(ex)
            rows.append(
                {
                    "kit": kit,
                    "case_id": case_id,
                    "tuan": week,
                    "pha": _phase_of(week),
                    "buoi": session_in_week,
                    "split": day.get("split_role") or "",
                    "section": ex.get("section") or "main",
                    "ten": _exercise_name(ex, catalog),
                    "sets": ex.get("sets") if ex.get("sets") is not None else "",
                    "reps": ex.get("reps") if ex.get("reps") is not None else "",
                    "rest": ex.get("rest_seconds") if ex.get("rest_seconds") is not None else "",
                    "notes_vi": ex.get("notes_vi") or "",
                    "load_note_vi": load_note,
                    "hiep": hiep,
                }
            )
    return rows


def _tong_quan(
    spec: dict[str, Any],
    insights: dict[str, Any],
    *,
    sessions: int,
    days: list[dict[str, Any]],
) -> dict[str, Any]:
    overview = insights.get("overview") or {}
    nut = insights.get("nutrition") or {}
    weight = insights.get("weight_goal") or {}
    wizard = insights.get("inputs") or {}
    mesos = list((insights.get("curriculum") or {}).get("mesocycles") or [])
    phases = (mesos + [{}, {}, {}])[:3]
    row: dict[str, Any] = {
        "kit": spec["kit"],
        "case_id": spec["case_id"],
        "buoi": sessions,
        "phut": spec["minutes"],
        "nhom_buoi": _week1_splits(days, sessions),
        "tdee": nut.get("tdee") or "",
        "target_kcal": nut.get("target_calories") or "",
        "protein_g": nut.get("protein_g") or "",
        "carbs_g": nut.get("carbs_g") or "",
        "fat_g": nut.get("fat_g") or "",
        "summary_vi": overview.get("summary_vi") or "",
        "periodization_vi": overview.get("periodization_vi") or "",
        "nutrition_vi": overview.get("nutrition_vi") or "",
        "recap_vi": wizard.get("recap_vi") or "",
        "mission_vi": overview.get("mission_vi") or "",
        "weight_copy": weight.get("copy_vi") or "",
        "weight_kcal": weight.get("daily_kcal") or "",
    }
    for i, meso in enumerate(phases, start=1):
        row[f"pha{i}_label"] = _phase_cell(meso, "label_vi")
        row[f"pha{i}_weeks"] = _phase_cell(meso, "weeks")
        row[f"pha{i}_blurb"] = _phase_cell(meso, "blurb_vi")
        row[f"pha{i}_rationale"] = _phase_cell(meso, "rationale_vi")
        row[f"pha{i}_knowledge"] = _phase_cell(meso, "knowledge")
    return row


def _case_row(spec: dict[str, Any], tong: dict[str, Any], *, week_code: str) -> dict[str, Any]:
    return {
        "kit": spec["kit"],
        "case_id": spec["case_id"],
        "gioi": GENDER_VI[spec["gender"]],
        "trinh_do": LEVEL_VI[int(spec["level"])],
        "muc_tieu": GOAL_VI[spec["goal"]],
        "height_cm": int(spec["height_cm"]),
        "weight_kg": int(spec["weight_kg"]),
        "kg_100_days": _kg_100(spec["goal"]),
        "tdee": tong.get("tdee") or "",
        "target_kcal": tong.get("target_kcal") or "",
        "protein_g": tong.get("protein_g") or "",
        "carbs_g": tong.get("carbs_g") or "",
        "fat_g": tong.get("fat_g") or "",
        "buoi": spec["sessions"],
        "phut": spec["minutes"],
        "focus": "+".join(spec.get("focus") or []) or "",
        "pushups": spec.get("pushups") if spec.get("pushups") != "" else "",
        "pullups": spec.get("pullups") if spec.get("pullups") != "" else "",
        "squats": spec.get("squats") if spec.get("squats") != "" else "",
        "plank_s": spec.get("plank_s") if spec.get("plank_s") is not None else "",
        "mau_day": spec.get("band_color") or "",
        "press_kg": spec.get("press_kg") if spec.get("press_kg") != "" else "",
        "row_kg": spec.get("row_kg") if spec.get("row_kg") != "" else "",
        "goblet_kg": spec.get("goblet_kg") if spec.get("goblet_kg") != "" else "",
        "week_code": week_code,
        "giai_thich": giai_thich(
            spec,
            tdee=tong.get("tdee") or None,
            target=tong.get("target_kcal") or None,
            week_code=week_code,
        ),
        "error": "",
    }


def run_challenge(db, spec: dict[str, Any], catalog: dict[int, dict[str, Any]]) -> dict[str, Any]:
    payload = payload_for(spec)
    result = generate_workout(db, None, payload, persist=False, commit=False)
    insights = result.get("insights") or {}
    sessions = int(result.get("sessions_actual") or spec["sessions"])
    expanded = expand_plan_days_for_weeks(
        result.get("days") or [],
        14,
        is_dict=True,
        experience_level=spec["level"],
        week_templates=insights.get("week_templates"),
        curriculum=True,
        goal=spec["goal"],
    )
    week_code = str(result.get("week_code") or insights.get("week_code") or "")
    tong = _tong_quan(spec, insights, sessions=sessions, days=expanded)
    return {
        "case": _case_row(spec, tong, week_code=week_code),
        "tong_quan": tong,
        "exercises": _exercise_rows(
            kit=spec["kit"],
            case_id=spec["case_id"],
            days=expanded,
            sessions=sessions,
            catalog=catalog,
        ),
        "familiarization": [],
    }


def _fam_payload(gender: str) -> dict[str, Any]:
    height, weight = MID_BODY[gender]
    return {
        "generation_mode": "familiarization",
        "familiarization_path": "first_push_pull",
        "gender": gender,
        "age": 28,
        "height_cm": float(height),
        "weight_kg": float(weight),
        "activity": "moderate",
        "goal": "lose_weight",
        "sessions_per_week": 3,
        "session_minutes": 35,
        "experience_level": 1,
        "location": "home",
        "equipment_list": [],
        "food_ids": [],
        "fitness_baseline": {},
        "ai_suggest_foods": False,
    }


def run_familiarization(db, gender: str, catalog: dict[int, dict[str, Any]]) -> dict[str, Any]:
    payload = _fam_payload(gender)
    result = generate_familiarization_workout(db, None, payload, persist=False, commit=False)
    insights = result.get("insights") or {}
    overview = insights.get("overview") or {}
    weight = insights.get("weight_goal") or {}
    case_id = f"nhapmon_{gender[0]}"
    days = result.get("days") or []
    sessions = int(result.get("sessions_actual") or 3)
    overview_row = {
        "kind": "tong_quan",
        "case_id": case_id,
        "gioi": GENDER_VI[gender],
        "mission_vi": overview.get("mission_vi") or "",
        "outcome_vi": overview.get("outcome_vi") or "",
        "nutrition_vi": overview.get("nutrition_vi") or "",
        "periodization_vi": overview.get("periodization_vi") or "",
        "weight_copy": weight.get("copy_vi") or "",
        "weight_kcal": weight.get("daily_kcal") or "",
        "ngay": "",
        "tuan": "",
        "split": "",
        "section": "",
        "ten": "",
        "sets": "",
        "reps": "",
        "rest": "",
        "notes_vi": overview.get("summary_vi") or "",
        "load_note_vi": "",
        "hiep": "",
    }
    exercise_rows: list[dict[str, Any]] = []
    for day in days:
        if not isinstance(day, dict):
            continue
        day_no = int(day.get("day_number") or 1)
        week = _week_of(day_no, sessions)
        for ex in day.get("exercises") or []:
            if not isinstance(ex, dict):
                continue
            load_note, hiep = _sets_bits(ex)
            exercise_rows.append(
                {
                    "kind": "bai",
                    "case_id": case_id,
                    "gioi": GENDER_VI[gender],
                    "mission_vi": "",
                    "outcome_vi": "",
                    "nutrition_vi": "",
                    "periodization_vi": "",
                    "weight_copy": "",
                    "weight_kcal": "",
                    "ngay": day_no,
                    "tuan": week,
                    "split": day.get("split_role") or "",
                    "section": ex.get("section") or "main",
                    "ten": _exercise_name(ex, catalog),
                    "sets": ex.get("sets") if ex.get("sets") is not None else "",
                    "reps": ex.get("reps") if ex.get("reps") is not None else "",
                    "rest": ex.get("rest_seconds") if ex.get("rest_seconds") is not None else "",
                    "notes_vi": ex.get("notes_vi") or "",
                    "load_note_vi": load_note,
                    "hiep": hiep,
                }
            )
    return {
        "case": {
            "kit": "nhapmon",
            "case_id": case_id,
            "gioi": GENDER_VI[gender],
            "trinh_do": LEVEL_VI[1],
            "muc_tieu": "nhập môn",
            "height_cm": int(payload["height_cm"]),
            "weight_kg": int(payload["weight_kg"]),
            "kg_100_days": "",
            "tdee": (insights.get("nutrition") or {}).get("tdee") or "",
            "target_kcal": (insights.get("nutrition") or {}).get("target_calories") or "",
            "protein_g": (insights.get("nutrition") or {}).get("protein_g") or "",
            "carbs_g": "",
            "fat_g": "",
            "buoi": sessions,
            "phut": 35,
            "focus": "",
            "pushups": "",
            "pullups": "",
            "squats": "",
            "plank_s": "",
            "mau_day": "",
            "press_kg": "",
            "row_kg": "",
            "goblet_kg": "",
            "week_code": "",
            "giai_thich": overview.get("summary_vi") or "",
            "error": "",
        },
        "tong_quan": None,
        "exercises": [],
        "familiarization": [overview_row, *exercise_rows],
    }


def load_ckpt(path: Path) -> dict[str, dict[str, Any]]:
    latest: dict[str, dict[str, Any]] = {}
    if not path.is_file():
        return latest
    with path.open(encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            try:
                row = json.loads(line)
            except json.JSONDecodeError:
                continue
            cid = row.get("case_id")
            if cid:
                latest[str(cid)] = row
    return latest


def append_ckpt(path: Path, row: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as fh:
        fh.write(json.dumps(row, ensure_ascii=False) + "\n")


def _ok(rec: dict[str, Any] | None) -> bool:
    if not rec or rec.get("error"):
        return False
    if rec.get("familiarization"):
        return True
    case = rec.get("case") or {}
    exercises = rec.get("exercises") or []
    return bool(case) and any((ex or {}).get("ten") for ex in exercises)


def _fail_case(spec: dict[str, Any], message: str) -> dict[str, Any]:
    return {
        "case_id": spec["case_id"],
        "error": message,
        "case": {
            "kit": spec["kit"],
            "case_id": spec["case_id"],
            "error": message,
            "giai_thich": message,
        },
        "tong_quan": {"kit": spec["kit"], "case_id": spec["case_id"]},
        "exercises": [],
        "familiarization": [],
    }


def _run_one(db, spec: dict[str, Any] | None, gender: str | None, catalog) -> dict[str, Any]:
    try:
        if gender:
            rec = run_familiarization(db, gender, catalog)
        else:
            assert spec is not None
            rec = run_challenge(db, spec, catalog)
        rec["error"] = None
        return rec
    except Exception as exc:  # noqa: BLE001
        db.rollback()
        if spec is None:
            spec = {
                "kit": "nhapmon",
                "case_id": f"nhapmon_{(gender or 'x')[0]}",
            }
        return _fail_case(spec, str(exc))


def _header(ws, headers: list[str], fill, font) -> None:
    cells = []
    for title in headers:
        cell = WriteOnlyCell(ws, value=title)
        cell.fill = fill
        cell.font = font
        cells.append(cell)
    ws.append(cells)


def _append_dict(ws, headers: list[str], row: dict[str, Any], wrap_cols: set[str], wrap) -> None:
    cells = []
    for key in headers:
        cell = WriteOnlyCell(ws, value=_blank(row.get(key, "")))
        if key in wrap_cols:
            cell.alignment = wrap
        cells.append(cell)
    ws.append(cells)


def write_xlsx(records: list[dict[str, Any]], path: Path) -> None:
    wb = Workbook(write_only=True)
    fill = PatternFill(start_color="FF1F4E3D", end_color="FF1F4E3D", fill_type="solid")
    font = Font(bold=True, color="FFFFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")
    wrap_case = {"giai_thich"}
    wrap_tong = {
        "summary_vi",
        "periodization_vi",
        "nutrition_vi",
        "recap_vi",
        "mission_vi",
        "weight_copy",
        "pha1_blurb",
        "pha1_rationale",
        "pha2_blurb",
        "pha2_rationale",
        "pha3_blurb",
        "pha3_rationale",
    }
    wrap_ex = {"notes_vi", "load_note_vi", "hiep"}
    wrap_fam = {
        "mission_vi",
        "outcome_vi",
        "nutrition_vi",
        "periodization_vi",
        "weight_copy",
        "notes_vi",
        "hiep",
    }

    ws_c = wb.create_sheet("CaNhanHoa")
    _header(ws_c, CASE_HEADERS, fill, font)
    ws_t = wb.create_sheet("TongQuan")
    _header(ws_t, TONG_HEADERS, fill, font)
    ws_l = wb.create_sheet("LuuY")
    _header(ws_l, ["title", "body"], fill, font)
    for title, body in LUUY:
        _append_dict(ws_l, ["title", "body"], {"title": title, "body": body}, {"body"}, wrap)
    ws_e = wb.create_sheet("BaiTap")
    _header(ws_e, EX_HEADERS, fill, font)
    ws_n = wb.create_sheet("NhapMon")
    _header(ws_n, FAM_HEADERS, fill, font)

    for rec in records:
        case = rec.get("case") or {}
        if case:
            _append_dict(ws_c, CASE_HEADERS, case, wrap_case, wrap)
        tong = rec.get("tong_quan")
        if tong:
            _append_dict(ws_t, TONG_HEADERS, tong, wrap_tong, wrap)
        for ex in rec.get("exercises") or []:
            _append_dict(ws_e, EX_HEADERS, ex, wrap_ex, wrap)
        for row in rec.get("familiarization") or []:
            _append_dict(ws_n, FAM_HEADERS, row, wrap_fam, wrap)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    specs = iter_specs()
    fam = ("female", "male")
    print(f"specs {len(specs)} familiarization {len(fam)}")
    done = load_ckpt(CKPT)
    ordered: list[dict[str, Any]] = []
    db = SessionLocal()
    try:
        catalog = load_catalog(db)
        jobs: list[tuple[str, dict[str, Any] | None, str | None]] = [
            (spec["case_id"], spec, None) for spec in specs
        ]
        jobs.extend((f"nhapmon_{gender[0]}", None, gender) for gender in fam)
        for i, (cid, spec, gender) in enumerate(jobs, start=1):
            if _ok(done.get(cid)):
                ordered.append(done[cid])
                print(f"[{i}/{len(jobs)}] skip {cid}")
                continue
            t0 = time.perf_counter()
            rec = _run_one(db, spec, gender, catalog)
            elapsed = round(time.perf_counter() - t0, 2)
            slim = {
                "case_id": cid,
                "error": rec.get("error"),
                "elapsed_s": elapsed,
                "case": rec.get("case"),
                "tong_quan": rec.get("tong_quan"),
                "exercises": rec.get("exercises") or [],
                "familiarization": rec.get("familiarization") or [],
            }
            append_ckpt(CKPT, slim)
            done[cid] = slim
            ordered.append(slim)
            print(
                f"[{i}/{len(jobs)}] {cid} {elapsed}s "
                f"ex={len(slim['exercises'])} fam={len(slim['familiarization'])} "
                f"err={slim.get('error')}"
            )
    finally:
        db.close()
    write_xlsx(ordered, OUT_XLSX)
    print(f"wrote {OUT_XLSX} records={len(ordered)}")


if __name__ == "__main__":
    main()
