"""Export all knowledge articles from DB to Excel (api/exports/)."""

from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path
from typing import Any

API_DIR = Path(__file__).resolve().parents[1]
if str(API_DIR) not in sys.path:
    sys.path.insert(0, str(API_DIR))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.core.database import SessionLocal
from app.models.entities import KnowledgeArticle

HEADER = [
    "id",
    "slug",
    "title_vi",
    "level",
    "level_vi",
    "sort_order",
    "read_time_min",
    "is_published",
    "published_at",
    "seo_title",
    "seo_description",
    "content_md",
    "url_path",
]

LEVEL_ORDER = {
    "beginner": 0,
    "basic": 0,
    "intermediate": 1,
    "advanced": 2,
}

LEVEL_VI = {
    "beginner": "Cơ bản",
    "basic": "Cơ bản",
    "intermediate": "Trung cấp",
    "advanced": "Nâng cao",
}

HEADER_FILL = PatternFill("FF1F4E3D", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFFFF")


def _iso(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat(timespec="seconds")
    return str(value)


def _write_header(ws) -> None:
    ws.append(HEADER)
    for col, _ in enumerate(HEADER, start=1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"


def _set_widths(ws) -> None:
    for col, title in enumerate(HEADER, start=1):
        letter = get_column_letter(col)
        if title == "content_md":
            width = 80
        elif title in {"title_vi", "seo_description", "seo_title"}:
            width = 40
        elif title in {"slug", "url_path"}:
            width = 32
        else:
            width = max(12, min(18, len(title) + 2))
        ws.column_dimensions[letter].width = width


def _wrap_content_column(ws) -> None:
    """Make full markdown readable in Excel (wrap + taller rows)."""
    content_col = HEADER.index("content_md") + 1
    for row_idx in range(2, ws.max_row + 1):
        cell = ws.cell(row_idx, content_col)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row_idx].height = 80


def _row_for(article: KnowledgeArticle) -> list[Any]:
    level = (article.level or "").strip().lower()
    return [
        article.id,
        article.slug,
        article.title_vi,
        article.level,
        LEVEL_VI.get(level, article.level or ""),
        article.sort_order,
        article.read_time_min,
        bool(article.is_published),
        _iso(article.published_at),
        article.seo_title or "",
        article.seo_description or "",
        article.content_md or "",
        f"/kien-thuc/{article.slug}",
    ]


def _write_markdown_files(articles: list[KnowledgeArticle], md_dir: Path) -> None:
    """One .md file per article so full content is easy to open outside Excel."""
    md_dir.mkdir(parents=True, exist_ok=True)
    index_lines = [
        "# Kho kiến thức TAPTOT",
        "",
        f"Xuất {len(articles)} bài.",
        "",
    ]
    for article in articles:
        level = (article.level or "").strip().lower()
        level_vi = LEVEL_VI.get(level, article.level or "")
        title = (article.title_vi or article.slug).strip()
        body = (article.content_md or "").strip()
        if not body.startswith("#"):
            body = f"# {title}\n\n{body}"
        header = (
            f"slug: {article.slug}\n"
            f"level: {level_vi}\n"
            f"url: /kien-thuc/{article.slug}\n"
            f"read_time_min: {article.read_time_min or ''}\n"
            f"published: {bool(article.is_published)}\n\n"
        )
        (md_dir / f"{article.slug}.md").write_text(header + body + "\n", encoding="utf-8")
        index_lines.append(f"- [{title}]({article.slug}.md) — {level_vi}")
    index_lines.append("")
    (md_dir / "00-index.md").write_text("\n".join(index_lines), encoding="utf-8")


# Thứ tự đọc trên trang kho kiến thức (frontend/src/lib/knowledgePath.ts).
_PATH_ARTICLES: list[tuple[str, str, str, str, str, str]] = [
    (
        "01",
        "Nền tảng tập luyện",
        "Đọc trước khi bắt đầu",
        "Hiểu cơ thể và chuẩn bị tập luyện",
        "1.0 Xác định mục tiêu",
        "Chọn một đích đo được (giảm mỡ, tăng cơ, giữ cân) trước khi đếm hiệp. Không có đích thì lịch tập và khẩu phần không có tiêu chí đúng/sai.",
    ),
    (
        "01",
        "Nền tảng tập luyện",
        "Đọc trước khi bắt đầu",
        "Hiểu cơ thể và chuẩn bị tập luyện",
        "1.1 Bản đồ nhóm cơ",
        "Biết cơ nào làm việc khi đẩy, kéo, squat để hiểu vì sao lịch gọi tên nhóm cơ, không phải để học giải phẫu trước buổi tập.",
    ),
    (
        "01",
        "Nền tảng tập luyện",
        "Đọc trước khi bắt đầu",
        "Hiểu cơ thể và chuẩn bị tập luyện",
        "1.2 Từ điển thuật ngữ",
        "Đọc được set, rep, RIR, warm-up trên lịch. Đặt sau mục tiêu và bản đồ cơ vì từ chỉ có nghĩa khi đã biết mình đang làm gì.",
    ),
    (
        "01",
        "Nền tảng tập luyện",
        "Đọc trước khi bắt đầu",
        "Hiểu cơ thể và chuẩn bị tập luyện",
        "1.3 Mẫu vận động và tăng/giảm độ khó",
        "Chọn biến thể vừa sức (tường, quỳ gối, sàn) thay vì cố bài khó. Nằm sau từ điển vì dùng đúng tên mẫu vận động.",
    ),
    (
        "01",
        "Nền tảng tập luyện",
        "Đọc trước khi bắt đầu",
        "Hiểu cơ thể và chuẩn bị tập luyện",
        "1.4 Tập tại nhà và dụng cụ an toàn",
        "Chốt sàn, tường, bàn, xà, cặp sách trước khi vào kỹ thuật. Người tập phòng gym vẫn đọc để biết bài nào không cần máy.",
    ),
    (
        "01",
        "Nền tảng tập luyện",
        "Đọc trước khi bắt đầu",
        "Học cách thực hiện bài tập",
        "1.5 Khởi động và vận động khớp",
        "Chuẩn bị đúng tầm vận động của buổi đó, không phải kéo giãn mọi khớp. Đặt trước form vì khớp lạnh làm form xấu dù đã biết lý thuyết.",
    ),
    (
        "01",
        "Nền tảng tập luyện",
        "Đọc trước khi bắt đầu",
        "Học cách thực hiện bài tập",
        "1.6 Hít thở, brace, kiểm soát thân",
        "Giữ bụng và nhịp thở để squat, plank, chống đẩy không gãy lưng. Học ngay trước bài form.",
    ),
    (
        "01",
        "Nền tảng tập luyện",
        "Đọc trước khi bắt đầu",
        "Học cách thực hiện bài tập",
        "1.7 Form và an toàn khớp",
        "Ba nguyên tắc kiểm soát (thân thẳng, hướng gối, khuỷu). Đặt sau khởi động và brace vì đó là điều kiện để giữ form.",
    ),
    (
        "01",
        "Nền tảng tập luyện",
        "Đọc trước khi bắt đầu",
        "Học cách thực hiện bài tập",
        "1.8 Theo dõi đau và dấu hiệu cần khám",
        "Phân biệt khó chịu khi tập với đau cần dừng. Nằm cuối khối kỹ thuật để người mới biết lúc nào không được 'cố thêm'.",
    ),
    (
        "01",
        "Nền tảng tập luyện",
        "Đọc trước khi bắt đầu",
        "Học cách thực hiện bài tập",
        "1.9 Cách đọc lịch tập",
        "Đọc set, rep, nghỉ, thứ tự bài trên lịch TAPTOT. Để cuối nền tảng vì lúc này từ vựng và form đã đủ để làm theo lịch.",
    ),
    (
        "02",
        "Xây dựng lịch tập",
        "Đọc khi cần",
        "Xây dựng và phát triển chương trình",
        "1.10 Volume, intensity, frequency",
        "Ba nút chỉnh lịch: bao nhiêu hiệp, nặng thế nào, mấy buổi. Chỉ đọc khi đã tập được vài tuần, không đọc trước buổi đầu.",
    ),
    (
        "02",
        "Xây dựng lịch tập",
        "Đọc khi cần",
        "Xây dựng và phát triển chương trình",
        "1.11 Progressive overload cơ bản",
        "Tăng tải khi form còn vững: thêm lần, rồi thêm tạ. Đặt sau ba nút vì overload là cách vặn các nút đó theo tuần.",
    ),
    (
        "02",
        "Xây dựng lịch tập",
        "Đọc khi cần",
        "Xây dựng và phát triển chương trình",
        "1.12 Nhật ký tập",
        "Ghi để biết tuần sau tăng gì. Không có số liệu thì overload chỉ là cảm giác.",
    ),
    (
        "02",
        "Xây dựng lịch tập",
        "Đọc khi cần",
        "Xây dựng và phát triển chương trình",
        "1.13 DOMS và chấn thương",
        "Mỏi cơ sau buổi mới là bình thường; đau khớp sắc, sưng, yếu đột ngột thì dừng. Đọc khi đã có buổi tập để so sánh.",
    ),
    (
        "02",
        "Xây dựng lịch tập",
        "Đọc khi cần",
        "Xây dựng và phát triển chương trình",
        "1.14 Phục hồi và giấc ngủ",
        "Cơ lớn lúc nghỉ, không lúc đang tập. Đặt sau DOMS vì người mới hay tưởng đau là phải tập tiếp.",
    ),
    (
        "02",
        "Xây dựng lịch tập",
        "Đọc khi cần",
        "Xây dựng và phát triển chương trình",
        "1.15 Deload cho người mới",
        "Tuần nhẹ để giữ kỹ thuật, không phải tuần thử bài lạ. Chỉ có ích sau khi đã có khối tập đều.",
    ),
    (
        "02",
        "Xây dựng lịch tập",
        "Đọc khi cần",
        "Xây dựng và phát triển chương trình",
        "1.16 Sai lầm phổ biến",
        "Tổng hợp lỗi sau khi đã biết form, overload và nghỉ. Đặt cuối khối để đối chiếu, không phải bài mở đầu.",
    ),
    (
        "03",
        "Dinh dưỡng và phục hồi",
        "Đọc khi cần",
        "Dinh dưỡng và thể lực",
        "1.17 Thâm hụt, thặng dư, cân bằng calo",
        "Cân nặng đi theo năng lượng cả tuần. Đọc khi mục tiêu là cân, không bắt buộc trước buổi tập đầu.",
    ),
    (
        "03",
        "Dinh dưỡng và phục hồi",
        "Đọc khi cần",
        "Dinh dưỡng và thể lực",
        "1.18 Cách tính TDEE",
        "Ước lượng năng lượng tiêu hao theo việc làm và buổi tập thật, rồi chỉnh theo cân sau 2 tuần. Nằm sau bài calo vì TDEE là cách ra con số.",
    ),
    (
        "03",
        "Dinh dưỡng và phục hồi",
        "Đọc khi cần",
        "Dinh dưỡng và thể lực",
        "1.19 Protein, carb, fat",
        "Chia năng lượng đã có thành đạm, tinh bột, béo. Không đọc trước khi biết mình đang thâm hụt hay thặng dư.",
    ),
    (
        "03",
        "Dinh dưỡng và phục hồi",
        "Đọc khi cần",
        "Dinh dưỡng và thể lực",
        "1.20 Ước lượng khẩu phần",
        "Đưa macro vào bữa ăn thật, không cần cân từng gram mãi. Nằm sau bài macro.",
    ),
    (
        "03",
        "Dinh dưỡng và phục hồi",
        "Đọc khi cần",
        "Dinh dưỡng và thể lực",
        "1.21 Nước, điện giải, rượu bia",
        "Những thứ làm buổi tập và cân sáng lệch mà không phải do cơ. Đọc khi đã theo ăn uống.",
    ),
    (
        "03",
        "Dinh dưỡng và phục hồi",
        "Đọc khi cần",
        "Dinh dưỡng và thể lực",
        "1.22 Cardio",
        "Đi bộ/chạy phục vụ sức khỏe và hỗ trợ thâm hụt, không thay tập sức mạnh. Đặt sau ăn uống vì giảm mỡ chủ yếu là calo.",
    ),
    (
        "03",
        "Dinh dưỡng và phục hồi",
        "Đọc khi cần",
        "Dinh dưỡng và thể lực",
        "1.23 Thực phẩm bổ sung",
        "Chỉ những thứ có bằng chứng, sau khi ăn và tập đã ổn. Cuối khối dinh dưỡng để không mua bột trước khi có bữa ăn.",
    ),
    (
        "04",
        "Tối ưu hóa tập luyện",
        "Dành cho người đã có kinh nghiệm",
        "Tối ưu hóa chương trình",
        "1.24 Overload nâng cao",
        "Nhiều cách tăng tải khi thêm lần đơn giản không còn đều. Cần nhật ký vài tuần.",
    ),
    (
        "04",
        "Tối ưu hóa tập luyện",
        "Dành cho người đã có kinh nghiệm",
        "Tối ưu hóa chương trình",
        "1.25 RPE và RIR",
        "Đo còn bao nhiêu lần trong hiệp. Người mới ước sai; để sau khi form đã ổn.",
    ),
    (
        "04",
        "Tối ưu hóa tập luyện",
        "Dành cho người đã có kinh nghiệm",
        "Tối ưu hóa chương trình",
        "1.26 Volume theo nhóm cơ",
        "Đếm hiệp trực tiếp và gián tiếp từng vùng. Chỉ có nghĩa khi đã ghi set và RIR.",
    ),
    (
        "04",
        "Tối ưu hóa tập luyện",
        "Dành cho người đã có kinh nghiệm",
        "Tối ưu hóa chương trình",
        "1.27 Deload đúng thời điểm",
        "Xả khi nhật ký cho thấy hiệu suất giảm kéo dài. Khác bài 1.15: bài này dành cho khối tập đã nặng.",
    ),
    (
        "04",
        "Tối ưu hóa tập luyện",
        "Dành cho người đã có kinh nghiệm",
        "Tối ưu hóa chương trình",
        "1.28 Chu kỳ hóa",
        "Xếp vài tuần nặng rồi một tuần nhẹ thành khối. Cần lịch đã ổn định.",
    ),
    (
        "04",
        "Tối ưu hóa tập luyện",
        "Dành cho người đã có kinh nghiệm",
        "Tối ưu hóa chương trình",
        "1.29 Kết nối não–cơ",
        "Công cụ khi đã an toàn nhưng khó cảm nhận cơ mục tiêu. Không phải điều kiện để tăng cơ.",
    ),
    (
        "04",
        "Tối ưu hóa tập luyện",
        "Dành cho người đã có kinh nghiệm",
        "Tối ưu hóa chương trình",
        "1.30 Drop set, superset, rest-pause",
        "Thêm mệt có chủ đích. Chỉ khi đã giữ form lúc gần thất bại với hiệp thường.",
    ),
    (
        "04",
        "Tối ưu hóa tập luyện",
        "Dành cho người đã có kinh nghiệm",
        "Tối ưu hóa chương trình",
        "1.31 Carb cycling",
        "Dồn carb vào ngày tập sau khi tổng calo tuần đã đúng. Không phải bước đầu giảm cân.",
    ),
    (
        "04",
        "Tối ưu hóa tập luyện",
        "Dành cho người đã có kinh nghiệm",
        "Tối ưu hóa chương trình",
        "1.32 Refeed và diet break",
        "Ngày hoặc tuần ăn nhiều hơn khi đã siết nhiều tuần. Cuối lộ trình vì dùng sai sẽ xóa thâm hụt.",
    ),
]

_WHY_PATH = [
    (
        "Bốn chặng, không phải 1.0 đến 1.32 theo số",
        "Số 1.x là mã bài. Thứ tự đọc là chặng: nền tảng, lịch tập, dinh dưỡng, tối ưu. Người mới không cần đọc hết trước buổi đầu.",
    ),
    (
        "Chặng 01 trước mọi thứ",
        "Mục tiêu, từ vựng, biến thể vừa sức, khởi động, thở, form, đau, cách đọc lịch. Thiếu chặng này thì volume và macro không áp dụng được.",
    ),
    (
        "Chặng 02 ghi là Đọc khi cần",
        "Volume và overload chỉ có ích sau vài tuần tập đều. Đọc sớm dễ tăng tạ khi form chưa vững.",
    ),
    (
        "Chặng 03 tách khỏi lịch tập",
        "Ăn uống phục vụ mục tiêu cân nặng. Người chỉ muốn học động tác vẫn tập được nếu chưa đọc TDEE.",
    ),
    (
        "Trong chặng 03: calo → TDEE → macro → khẩu phần → nước → cardio → bổ sung",
        "Mỗi bài dùng con số của bài trước. Bổ sung đặt cuối để không thay bữa ăn.",
    ),
    (
        "Chặng 04 để sau khi có nhật ký",
        "RIR, volume theo cơ, chu kỳ, drop set, carb cycling đều giả định kỹ thuật ổn và đã ghi được vài tuần. Bài viết ghi rõ 'Phù hợp khi'.",
    ),
    (
        "Vì sao bài viết có câu Phù hợp khi và lưu ý HLV",
        "Giới hạn ai nên dùng kỹ thuật đó, và câu nào không được hiểu thành quy luật sinh học (ví dụ hệ số 1/2 hiệp gián tiếp chỉ là điểm khởi đầu).",
    ),
]


def _write_path_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Lo_trinh", 0)
    headers = ["Chang", "Ten chang", "Khi nao doc", "Cum bai", "Bai", "Muc dich va vi sao dung o day"]
    ws.append(headers)
    for col, _ in enumerate(headers, start=1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"
    for row in _PATH_ARTICLES:
        ws.append(list(row))
    widths = [10, 28, 32, 36, 36, 72]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row_idx in range(2, ws.max_row + 1):
        for col in range(1, 7):
            ws.cell(row_idx, col).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row_idx].height = 48

    why = wb.create_sheet("Vi_sao_sap_xep", 1)
    why.append(["Y", "Giai thich"])
    for col in (1, 2):
        cell = why.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    why.freeze_panes = "A2"
    for row in _WHY_PATH:
        why.append(list(row))
    why.column_dimensions["A"].width = 62
    why.column_dimensions["B"].width = 88
    for row_idx in range(2, why.max_row + 1):
        for col in (1, 2):
            why.cell(row_idx, col).alignment = Alignment(wrap_text=True, vertical="top")
        why.row_dimensions[row_idx].height = 48


def export_articles(db: Session, out: Path) -> int:
    rows = list(db.query(KnowledgeArticle).all())
    rows.sort(
        key=lambda a: (
            LEVEL_ORDER.get((a.level or "").strip().lower(), 99),
            a.sort_order or 0,
            a.id or 0,
        )
    )

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "articles"
    _write_header(ws)
    for article in rows:
        ws.append(_row_for(article))
    _set_widths(ws)
    _wrap_content_column(ws)
    _write_path_sheet(wb)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)

    md_dir = out.with_suffix("")  # knowledge_articles_YYYYMMDD_HHMM/
    _write_markdown_files(rows, md_dir)
    return len(rows)


def main() -> Path:
    import shutil

    out_dir = Path(__file__).resolve().parents[1] / "exports"
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    out = out_dir / f"knowledge_articles_{stamp}.xlsx"
    db = SessionLocal()
    try:
        n = export_articles(db, out)
        md_dir = out.with_suffix("")
        desktop = Path.home() / "Desktop"
        desktop.mkdir(parents=True, exist_ok=True)
        xlsx_copy = desktop / out.name
        md_copy = desktop / md_dir.name
        shutil.copy2(out, xlsx_copy)
        if md_copy.exists():
            shutil.rmtree(md_copy)
        shutil.copytree(md_dir, md_copy)
        print(f"Exported {n} knowledge articles (full content_md in Excel + .md files)")
        print(out)
        print(md_dir)
        print(xlsx_copy)
        print(md_copy)
        return out
    finally:
        db.close()


if __name__ == "__main__":
    main()
