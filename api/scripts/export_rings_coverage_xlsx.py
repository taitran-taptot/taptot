"""One-shot: ~90 bar+rings challenge plans -> Excel. persist=False."""

from __future__ import annotations

import json
import sys
import time
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[2]
API = ROOT / "api"
if str(API) not in sys.path:
    sys.path.insert(0, str(API))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import text

from app.core.database import SessionLocal
from app.services.periodization import expand_plan_days_for_weeks
from app.services.workout_generation.service import generate_workout

OUT_XLSX = ROOT / "docs" / "testing" / "rings-coverage-90-2026-10-27.xlsx"
CKPT = ROOT / "docs" / "testing" / "rings-coverage-90.ckpt.jsonl"
SAMPLE_WEEKS = (1, 7, 14)

CASE_HEADERS = [
    "case_id",
    "gioi",
    "trinh_do",
    "muc_tieu",
    "height_cm",
    "weight_kg",
    "kg_100_days",
    "tdee",
    "target_kcal",
    "buoi",
    "phut",
    "focus",
    "test_cai",
    "pushups",
    "pullups",
    "squats",
    "plank_s",
    "week_code",
    "giai_thich",
    "error",
]
EX_HEADERS = [
    "case_id",
    "tuan",
    "pha",
    "buoi",
    "split",
    "section",
    "ten",
    "sets",
    "reps",
    "rest",
    "load_note_vi",
]

GENDER_VI = {"male": "Nam", "female": "Nữ"}
LEVEL_VI = {1: "người mới", 2: "đã có thói quen", 3: "tập đều"}
GOAL_VI = {"lose_weight": "giảm cân", "gain_weight": "tăng cân"}
SPLIT_WHY = {
    ("male", 3): "PPL vì nam 3 buổi",
    ("male", 4): "ULUL vì nam 4 buổi",
    ("male", 5): "PPLUL vì nam 5 buổi",
    ("female", 3): "LUL vì nữ 3 buổi",
    ("female", 4): "LPPL vì nữ 4 buổi",
    ("female", 5): "LULU kèm cardio-core vì nữ 5 buổi",
}


def _profile(gender: str) -> dict[str, Any]:
    female = gender == "female"
    return {
        "age": 28,
        "height_cm": 160.0 if female else 175.0,
        "weight_kg": 68.0 if female else 75.0,
        "activity": "moderate",
    }


def _kg_week(goal: str) -> float:
    return 0.5 if goal == "lose_weight" else 0.35


def _kg_100(goal: str) -> float:
    return round(_kg_week(goal) * 14, 1)


def _core_baseline(
    *,
    push: int | None = 15,
    pull: int | None = 15,
    squat: int | None = 15,
    plank: int | None = 20,
) -> dict[str, Any]:
    base: dict[str, Any] = {"test_kit": "bar_rings"}
    if push is not None:
        base["pushups_max"] = push
    if pull is not None:
        base["pullups_max"] = pull
    if squat is not None:
        base["squats_max"] = squat
    if plank is not None:
        base["plank_seconds"] = plank
    return base


def iter_specs() -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for gender in ("male", "female"):
        for level in (1, 2, 3):
            for goal in ("lose_weight", "gain_weight"):
                for sessions in (3, 4, 5):
                    for minutes in (45, 60):
                        cid = f"core_{gender[0]}_L{level}_{goal[:4]}_s{sessions}_m{minutes}"
                        out.append(
                            {
                                "case_id": cid,
                                "wave": "core",
                                "gender": gender,
                                "level": level,
                                "goal": goal,
                                "sessions": sessions,
                                "minutes": minutes,
                                "focus": [],
                                "baseline": _core_baseline(),
                                "test_reps": 15,
                                "pushups": 15,
                                "pullups": 15,
                                "squats": 15,
                                "plank_s": 20,
                            }
                        )

    def overlay(**kw: Any) -> dict[str, Any]:
        spec: dict[str, Any] = {
            "wave": "overlay",
            "gender": "female",
            "level": 1,
            "goal": "lose_weight",
            "sessions": 3,
            "minutes": 45,
            "focus": [],
            "test_reps": 15,
            "pushups": 15,
            "pullups": 15,
            "squats": 15,
            "plank_s": 20,
        }
        spec.update(kw)
        if "baseline" not in spec:
            spec["baseline"] = _core_baseline(
                push=spec.get("pushups"),
                pull=spec.get("pullups"),
                squat=spec.get("squats"),
                plank=spec.get("plank_s"),
            )
        return spec

    for n in (8, 15, 25, 35):
        out.append(
            overlay(
                case_id=f"amrap_F_L1_s3_m45_lose_{n}",
                test_reps=n,
                pushups=n,
                pullups=n,
                squats=n,
                baseline=_core_baseline(push=n, pull=n, squat=n, plank=20),
            )
        )
    for n in (8, 15, 25, 35):
        out.append(
            overlay(
                case_id=f"amrap_M_L3_s4_m60_gain_{n}",
                gender="male",
                level=3,
                goal="gain_weight",
                sessions=4,
                minutes=60,
                test_reps=n,
                pushups=n,
                pullups=n,
                squats=n,
                baseline=_core_baseline(push=n, pull=n, squat=n, plank=20),
            )
        )
    for pull in (0, 3, 8, 15, 25):
        out.append(
            overlay(
                case_id=f"pull_F_L2_s3_m45_{pull}",
                gender="female",
                level=2,
                pullups=pull,
                baseline=_core_baseline(push=15, pull=pull, squat=15, plank=20),
            )
        )
    for sec in (20, 60):
        out.append(
            overlay(
                case_id=f"plank_M_L2_s3_m45_{sec}s",
                gender="male",
                level=2,
                plank_s=sec,
                baseline=_core_baseline(push=15, pull=15, squat=15, plank=sec),
            )
        )
    out.append(
        overlay(
            case_id="focus_F_L2_s3_m45_none",
            gender="female",
            level=2,
            baseline=_core_baseline(),
        )
    )
    out.append(
        overlay(
            case_id="focus_F_L2_s3_m45_bcm",
            gender="female",
            level=2,
            focus=["bung", "chan", "mong"],
            baseline=_core_baseline(),
        )
    )
    out.append(
        overlay(
            case_id="notest_M_L1_s3_m45",
            gender="male",
            level=1,
            test_reps=None,
            pushups="",
            pullups="",
            squats="",
            plank_s=None,
            baseline={"test_kit": "bar_rings"},
        )
    )
    return out


def payload_for(spec: dict[str, Any]) -> dict[str, Any]:
    goal = spec["goal"]
    return {
        "goal": goal,
        "gender": spec["gender"],
        "sessions_per_week": spec["sessions"],
        "session_minutes": spec["minutes"],
        "location": "home",
        "focus_areas": list(spec["focus"] or []),
        "extra_goals": [],
        "equipment_list": ["pull-up-bar", "gymnastic-rings"],
        "food_ids": [],
        "experience_level": spec["level"],
        "ai_suggest_equipment": False,
        "no_equipment": False,
        "ai_suggest_foods": False,
        "fitness_baseline": dict(spec["baseline"]),
        "duration_weeks": 14,
        "kg_per_week": _kg_week(goal),
        "challenge_100_days": True,
        "curriculum_12_weeks": False,
        **_profile(spec["gender"]),
    }


def _amrap_note(n: int | None) -> str:
    if n is None:
        return "Không điền test nên compound 8–12, isolation 12–18; slot xà theo thiếu test."
    if n < 10:
        return f"Test {n} cái (<10) nên cửa sổ rep hẹp / skill khó hơn."
    if n > 20:
        return f"Test {n} cái (>20) nên nới cửa sổ rep."
    return f"Test {n} cái nằm 10–20 nên giữ cửa sổ rep giữa vùng."


def _pull_note(spec: dict[str, Any]) -> str:
    pull = spec.get("pullups")
    if pull == "" or pull is None:
        return "Không ghi kéo xà nên slot kéo theo thiếu test (thường chèo vòng / biến thể dễ hơn)."
    try:
        n = int(pull)
    except (TypeError, ValueError):
        return ""
    if n <= 0:
        return "Kéo xà 0 cái nên cổng bar_rings_slots ưu tiên chèo vòng / biến thể dễ hơn kéo xà."
    if n < 8:
        return f"Kéo xà {n} cái (thấp) nên dễ ra chèo vòng hoặc biến thể dễ hơn kéo xà đầy đủ."
    return f"Kéo xà {n} cái nên đủ điều kiện bài kéo xà trên xà/vòng."


def giai_thich(spec: dict[str, Any], *, tdee: Any, target: Any, week_code: str) -> str:
    gender = spec["gender"]
    level = int(spec["level"])
    goal = spec["goal"]
    sessions = int(spec["sessions"])
    minutes = int(spec["minutes"])
    prof = _profile(gender)
    kg100 = _kg_100(goal)
    huong = "giảm" if goal == "lose_weight" else "tăng"
    parts = [
        f"{GENDER_VI[gender]} {int(prof['height_cm'])} cm, {int(prof['weight_kg'])} kg, "
        f"{huong} {str(kg100).replace('.', ',')} kg / 100 ngày"
        + (f", TDEE {tdee} kcal, mục tiêu {target} kcal/ngày" if tdee and target else "")
        + ".",
        f"Split {SPLIT_WHY[(gender, sessions)]}.",
        (
            "Trần 3 hiệp vì người mới; bài test pha 2 được 4 hiệp, pha 1 và 3 giữ 3."
            if level <= 1
            else "Trần 4 hiệp vì đã có thói quen hoặc tập đều; bài test pha 1/2/3 là 3/4/3."
        ),
        "Tuần 4, 8, 14 giảm hiệp (làm tròn 0,6, sàn 2), nghỉ lâu hơn, giữ độ khó bài (không tăng skill).",
    ]
    if minutes <= 45:
        if level <= 1:
            parts.append(
                "45 phút: pha 3 có thể cộng hiệp ưu tiên nhưng người mới không vượt trần 3."
            )
        else:
            parts.append("45 phút: pha 3 có thể cộng 1 hiệp cho vùng ưu tiên (trần 4).")
    else:
        parts.append("60 phút: không cộng hiệp ưu tiên trên lịch thử thách.")
    if goal == "lose_weight":
        parts.append("Giảm cân: giữ khoảng số cái, không ramp rep theo tuần.")
    else:
        parts.append("Tăng cân: cho phép tăng số cái theo tuần.")
    parts.append("Kit xà đơn + vòng treo, không stamp kg tạ.")
    parts.append(_amrap_note(spec.get("test_reps")))
    pull_bit = _pull_note(spec)
    if pull_bit:
        parts.append(pull_bit)
    plank = spec.get("plank_s")
    if plank:
        parts.append(f"Plank khoảng 80% của {plank} giây.")
    else:
        parts.append("Không test plank nên giữ khoảng mặc định.")
    focus = spec.get("focus") or []
    if focus:
        parts.append("Ưu tiên " + ", ".join(focus) + ".")
    else:
        parts.append("Không chọn nhóm ưu tiên.")
    if week_code:
        parts.append(f"Mã tuần engine: {week_code}.")
    return " ".join(parts)


def load_catalog(db) -> dict[int, dict[str, Any]]:
    rows = db.execute(
        text("SELECT id, name_en, name_vi FROM exercises")
    ).mappings().all()
    return {int(r["id"]): dict(r) for r in rows}


def _exercise_name(ex: dict[str, Any], catalog: dict[int, dict[str, Any]]) -> str:
    try:
        eid = int(ex.get("exercise_id"))
    except (TypeError, ValueError):
        eid = None
    cat = catalog.get(eid) if eid is not None else None
    return str(
        (cat or {}).get("name_vi")
        or ex.get("name_vi")
        or (cat or {}).get("name_en")
        or ex.get("name_en")
        or ""
    )


def _load_note(ex: dict[str, Any]) -> str:
    rx = ex.get("set_prescriptions") or []
    if rx and isinstance(rx[0], dict):
        return str(rx[0].get("load_note_vi") or "")
    return str(ex.get("notes_vi") or "")


def _week_of(day_number: int, sessions: int) -> int:
    return (max(1, int(day_number)) - 1) // max(1, int(sessions)) + 1


def _phase_of(week: int) -> int:
    if week <= 4:
        return 1
    if week <= 8:
        return 2
    return 3


def load_ckpt(path: Path) -> dict[str, dict[str, Any]]:
    latest: dict[str, dict[str, Any]] = {}
    if not path.is_file():
        return latest
    with path.open(encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            try:
                row = json.loads(line)
            except json.JSONDecodeError:
                continue
            cid = row.get("case_id")
            if cid:
                latest[str(cid)] = row
    return latest


def append_ckpt(path: Path, row: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as fh:
        fh.write(json.dumps(row, ensure_ascii=False) + "\n")


def run_case(db, spec: dict[str, Any], catalog: dict[int, dict[str, Any]]) -> dict[str, Any]:
    payload = payload_for(spec)
    t0 = time.perf_counter()
    try:
        result = generate_workout(db, None, payload, persist=False, commit=False)
    except Exception as exc:  # noqa: BLE001
        return {
            "spec": spec,
            "error": str(exc),
            "elapsed_s": round(time.perf_counter() - t0, 2),
            "case": None,
            "exercises": [],
        }
    insights = result.get("insights") or {}
    nut = insights.get("nutrition") or {}
    tdee = nut.get("tdee")
    target = nut.get("target_calories")
    week_code = str(result.get("week_code") or insights.get("week_code") or "")
    sessions = int(result.get("sessions_actual") or spec["sessions"])
    seed = result.get("days") or []
    templates = insights.get("week_templates")
    expanded = expand_plan_days_for_weeks(
        seed,
        14,
        is_dict=True,
        experience_level=spec["level"],
        week_templates=templates,
        curriculum=True,
        goal=spec["goal"],
    )
    ex_rows: list[dict[str, Any]] = []
    for day in expanded:
        week = _week_of(int(day.get("day_number") or 1), sessions)
        if week not in SAMPLE_WEEKS:
            continue
        session_in_week = ((int(day.get("day_number") or 1) - 1) % sessions) + 1
        for ex in day.get("exercises") or []:
            ex_rows.append(
                {
                    "case_id": spec["case_id"],
                    "tuan": week,
                    "pha": _phase_of(week),
                    "buoi": session_in_week,
                    "split": day.get("split_role") or "",
                    "section": ex.get("section") or "main",
                    "ten": _exercise_name(ex, catalog),
                    "sets": ex.get("sets"),
                    "reps": ex.get("reps"),
                    "rest": ex.get("rest_seconds"),
                    "load_note_vi": _load_note(ex),
                }
            )
    prof = _profile(spec["gender"])
    case_row = {
        "case_id": spec["case_id"],
        "gioi": GENDER_VI[spec["gender"]],
        "trinh_do": LEVEL_VI[spec["level"]],
        "muc_tieu": GOAL_VI[spec["goal"]],
        "height_cm": int(prof["height_cm"]),
        "weight_kg": int(prof["weight_kg"]),
        "kg_100_days": _kg_100(spec["goal"]),
        "tdee": tdee,
        "target_kcal": target,
        "buoi": spec["sessions"],
        "phut": spec["minutes"],
        "focus": "+".join(spec.get("focus") or []) or "",
        "test_cai": spec.get("test_reps") if spec.get("test_reps") is not None else "",
        "pushups": spec.get("pushups") if spec.get("pushups") != "" else "",
        "pullups": spec.get("pullups") if spec.get("pullups") != "" else "",
        "squats": spec.get("squats") if spec.get("squats") != "" else "",
        "plank_s": spec.get("plank_s") if spec.get("plank_s") is not None else "",
        "week_code": week_code,
        "giai_thich": giai_thich(spec, tdee=tdee, target=target, week_code=week_code),
        "error": "",
    }
    return {
        "spec": spec,
        "error": None,
        "elapsed_s": round(time.perf_counter() - t0, 2),
        "case": case_row,
        "exercises": ex_rows,
    }


def write_xlsx(cases: list[dict[str, Any]], exercises: list[dict[str, Any]], path: Path) -> None:
    wb = Workbook()
    fill = PatternFill(start_color="FF1F4E3D", end_color="FF1F4E3D", fill_type="solid")
    font = Font(bold=True, color="FFFFFFFF")

    def sheet(title: str, headers: list[str], rows: list[dict[str, Any]], wrap: str | None = None) -> None:
        ws = wb.create_sheet(title)
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
        for row in rows:
            ws.append([row.get(h, "") for h in headers])
        for i, h in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(i)].width = 56 if h == wrap else 16
        if wrap:
            idx = headers.index(wrap) + 1
            for cells in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                cells[0].alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[1].height = 22

    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    sheet("Cases", CASE_HEADERS, cases, wrap="giai_thich")
    sheet("BaiTap", EX_HEADERS, exercises)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    specs = iter_specs()
    print(f"specs {len(specs)}")
    done = load_ckpt(CKPT)
    cases: list[dict[str, Any]] = []
    exercises: list[dict[str, Any]] = []
    db = SessionLocal()
    try:
        catalog = load_catalog(db)
        for i, spec in enumerate(specs, start=1):
            cid = spec["case_id"]
            if cid in done and done[cid].get("case") and any(
                (ex or {}).get("ten") for ex in (done[cid].get("exercises") or [])
            ):
                rec = done[cid]
                cases.append(rec["case"])
                exercises.extend(rec.get("exercises") or [])
                print(f"[{i}/{len(specs)}] skip {cid}")
                continue
            rec = run_case(db, spec, catalog)
            slim = {
                "case_id": cid,
                "error": rec.get("error"),
                "elapsed_s": rec.get("elapsed_s"),
                "case": rec.get("case"),
                "exercises": rec.get("exercises") or [],
            }
            append_ckpt(CKPT, slim)
            done[cid] = slim
            if rec.get("case"):
                cases.append(rec["case"])
                exercises.extend(rec.get("exercises") or [])
            else:
                cases.append(
                    {
                        "case_id": cid,
                        "error": rec.get("error") or "fail",
                        "giai_thich": rec.get("error") or "fail",
                    }
                )
            print(f"[{i}/{len(specs)}] {cid} {rec.get('elapsed_s')}s err={rec.get('error')}")
    finally:
        db.close()
    write_xlsx(cases, exercises, OUT_XLSX)
    print(f"wrote {OUT_XLSX} cases={len(cases)} exercises={len(exercises)}")


if __name__ == "__main__":
    main()
