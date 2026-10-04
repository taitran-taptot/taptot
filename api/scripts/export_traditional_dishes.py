"""Export traditional dishes list (seed recipe-BOM macros) to Excel + Markdown."""

from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

API_DIR = Path(__file__).resolve().parents[1]
PROJECT_ROOT = API_DIR.parent
if str(API_DIR) not in sys.path:
    sys.path.insert(0, str(API_DIR))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER = [
    "slug",
    "name_vi",
    "name_en",
    "group_vi",
    "region_slug",
    "province_id",
    "serving_size",
    "serving_grams",
    "calories",
    "protein_g",
    "carbs_g",
    "fat_g",
    "fiber_g",
    "prep_state",
    "source_ref",
    "cook_url",
    "description_vi",
]

HEADER_FILL = PatternFill("FF1F4E3D", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFFFF")


def _group_vi(tags: Any) -> str:
    for raw in tags or []:
        tag = str(raw)
        if tag.startswith("nhom_vi:"):
            return tag[len("nhom_vi:") :].strip()
    return ""


def load_dishes() -> list[dict[str, Any]]:
    path = PROJECT_ROOT / "seeds" / "foods_traditional_dishes.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, list):
        raise SystemExit("foods_traditional_dishes.json must be a list")
    rows = [d for d in data if isinstance(d, dict)]
    rows.sort(key=lambda d: (_group_vi(d.get("tags")), str(d.get("name_vi") or "")))
    return rows


def main() -> Path:
    dishes = load_dishes()
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    out_dir = API_DIR / "exports"
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx = out_dir / f"mon_truyen_thong_{stamp}.xlsx"

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "mon_truyen_thong"
    ws.append(HEADER)
    for col, _ in enumerate(HEADER, start=1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"

    md_lines = [
        "# Món truyền thống Việt",
        "",
        f"Tổng số món: {len(dishes)}",
        "",
        "Calo suất lấy từ BOM cách nấu (`source_ref: recipe-bom`).",
        "",
        "| Nhóm | Món | kcal/suất | Gram | Cách nấu |",
        "| --- | --- | ---: | ---: | --- |",
    ]

    for dish in dishes:
        slug = str(dish.get("slug") or "")
        group = _group_vi(dish.get("tags"))
        cook = f"/cach-nau?mon={slug}"
        ws.append(
            [
                slug,
                dish.get("name_vi") or "",
                dish.get("name_en") or "",
                group,
                dish.get("region_slug") or "",
                dish.get("province_id") or "",
                dish.get("serving_size") or "",
                dish.get("serving_grams"),
                dish.get("calories"),
                dish.get("protein_g"),
                dish.get("carbs_g"),
                dish.get("fat_g"),
                dish.get("fiber_g"),
                dish.get("prep_state") or "cooked",
                dish.get("source_ref") or "",
                cook,
                dish.get("description_vi") or "",
            ]
        )
        name = dish.get("name_vi") or slug
        kcal = dish.get("calories")
        grams = dish.get("serving_grams")
        md_lines.append(
            f"| {group} | {name} | {kcal} | {grams} | [{slug}]({cook}) |"
        )

    for col, title in enumerate(HEADER, start=1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = 42 if title in {"description_vi", "name_vi"} else 16
    wb.save(xlsx)

    md_path = out_dir / f"mon_truyen_thong_{stamp}.md"
    md_path.write_text("\n".join(md_lines) + "\n", encoding="utf-8")
    print(f"Exported {len(dishes)} traditional dishes")
    print(xlsx)
    print(md_path)
    return xlsx


if __name__ == "__main__":
    main()
