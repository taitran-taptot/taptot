"""Import knowledge articles from Excel into .md files + DB (upsert by slug)."""

from __future__ import annotations

import argparse
import re
import sys
from datetime import UTC, datetime
from pathlib import Path

API_DIR = Path(__file__).resolve().parents[1]
if str(API_DIR) not in sys.path:
    sys.path.insert(0, str(API_DIR))

from openpyxl import load_workbook
from sqlalchemy import select

from app.core.database import SessionLocal
from app.models.entities import KnowledgeArticle

ROOT = API_DIR / "app" / "data" / "knowledge"

DEFAULT_XLSX = Path.home() / "Downloads" / "TAPTOT_Kho_Kien_Thuc_Chuan_HLV_Full_28_Bai_v2.xlsx"

# (slug, image_key, alt_caption ≤7 words)
PLACEHOLDERS: list[tuple[str, str, str]] = [
    ("cach-doc-lich-tap-quy-uoc-buoi-tap", "set-decoder-card", "Giải mã 3×8–10"),
    ("ban-do-cac-nhom-co-chinh-co-che-chuyen-dong", "muscle-heat-map", "Bản đồ nhóm cơ"),
    ("ky-thuat-tap-chuan-form-an-toan-co-xuong-khop", "pushup-force-vectors", "Hít đất — hướng lực"),
    ("ky-thuat-tap-chuan-form-an-toan-co-xuong-khop", "plank-do-dont", "Plank: đúng vs sai"),
    ("ba-nut-chinh-khoi-luong-volume-do-nang-intensity-tan-suat-frequency", "vif-three-dials", "Volume · Intensity · Frequency"),
    ("nguyen-tac-qua-tai-luy-tien-progressive-overload-co-ban", "overload-pyramid", "Kim tự tháp overload"),
    ("dau-moi-co-doms-va-chan-thuong-cach-phan-biet-va-xu-ly", "squat-do-dont", "Squat: đúng vs sai"),
    ("nang-luong-va-can-nang-tham-hut-thang-du-va-can-bang-calo", "energy-balance-scale", "Cán cân năng lượng"),
    ("cach-tinh-tdee-theo-muc-van-dong-thuc-te", "tdee-pie", "Cấu trúc TDEE"),
    ("dinh-duong-da-luong-chat-dam-protein-tinh-bot-carb-va-chat-beo-fat", "hand-portion-guide", "Định lượng bàn tay"),
    ("rpe-va-rir-trong-tung-set", "rpe-rir-card", "Thang RPE / RIR"),
    ("mind-muscle-connection-nang-cao", "mmc-target-heat", "Cơ mục tiêu MMC"),
]

HEADER = [
    "id",
    "slug",
    "title_vi",
    "level",
    "level_vi",
    "sort_order",
    "read_time_min",
    "is_published",
    "published_at",
    "seo_title",
    "seo_description",
    "content_md",
    "url_path",
]


def _parse_published_at(value: object) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None) if value.tzinfo else value
    s = str(value).strip()
    if not s:
        return None
    try:
        dt = datetime.fromisoformat(s.replace("Z", "+00:00"))
        return dt.replace(tzinfo=None) if dt.tzinfo else dt
    except ValueError:
        return None


def _row_dict(headers: list[str], row: tuple) -> dict[str, object]:
    out: dict[str, object] = {}
    for i, key in enumerate(headers):
        if key:
            out[str(key)] = row[i] if i < len(row) else None
    return out


def strip_knowledge_figures(content: str) -> str:
    """Remove markdown knowledge images and italic figure captions."""
    lines = content.splitlines(keepends=True)
    out: list[str] = []
    i = 0
    img_re = re.compile(r"^!\[[^\]]*\]\((?:/media/knowledge/|pending:)[^)]+\)\s*$")
    caption_re = re.compile(r"^\*Hình[^*]*\*\s*$")
    while i < len(lines):
        if img_re.match(lines[i].strip()):
            i += 1
            if i < len(lines) and caption_re.match(lines[i].strip()):
                i += 1
            continue
        if caption_re.match(lines[i].strip()):
            i += 1
            continue
        out.append(lines[i])
        i += 1
    text = "".join(out)
    text = re.sub(r"\n{3,}", "\n\n", text)
    if content.endswith("\n") and not text.endswith("\n"):
        text += "\n"
    return text


def strip_image_suggestion_blocks(content: str) -> str:
    """Remove Excel 'Gợi ý hình ảnh minh họa' prompt blocks."""
    lines = content.splitlines(keepends=True)
    out: list[str] = []
    i = 0
    while i < len(lines):
        if "Gợi ý hình ảnh minh họa" in lines[i]:
            i += 1
            while i < len(lines):
                raw = lines[i]
                if raw.lstrip().startswith(">"):
                    i += 1
                    continue
                if raw.strip() == "":
                    j = i + 1
                    while j < len(lines) and lines[j].strip() == "":
                        j += 1
                    if j < len(lines) and lines[j].lstrip().startswith(">"):
                        i = j
                        continue
                    break
                break
            continue
        out.append(lines[i])
        i += 1
    text = "".join(out)
    text = re.sub(r"\n{3,}", "\n\n", text)
    if content.endswith("\n") and not text.endswith("\n"):
        text += "\n"
    return text


def _inject_placeholders(content: str, figures: list[tuple[str, str]]) -> str:
    """Insert ![alt](pending:key) blocks after the first --- (or after intro)."""
    blocks = [
        f"![{alt}](pending:{key})"
        for key, alt in figures
        if f"(pending:{key})" not in content and f"/{key}.png" not in content
    ]
    if not blocks:
        return content

    insert = "\n\n" + "\n\n".join(blocks) + "\n"
    # Prefer after first markdown thematic break
    m = re.search(r"\n---\s*\n", content)
    if m:
        pos = m.end()
        return content[:pos] + insert + content[pos:]
    # Else after first blockquote block
    m2 = re.search(r"(>.*\n(?:(?!>).*\n)*)", content)
    if m2:
        pos = m2.end()
        return content[:pos] + insert + content[pos:]
    # Fallback: after first blank line following H1
    m3 = re.search(r"^# .+\n+", content, re.M)
    if m3:
        return content[: m3.end()] + insert.lstrip("\n") + content[m3.end() :]
    return insert.lstrip("\n") + content


def load_rows(xlsx: Path) -> list[dict[str, object]]:
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb.active
    assert ws is not None
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter)
    headers = [str(h).strip() if h is not None else "" for h in header_row]
    missing = [h for h in ("slug", "title_vi", "content_md") if h not in headers]
    if missing:
        wb.close()
        raise SystemExit(f"Excel missing columns: {missing}")
    out: list[dict[str, object]] = []
    for row in rows_iter:
        d = _row_dict(headers, tuple(row))
        slug = str(d.get("slug") or "").strip()
        if not slug:
            continue
        out.append(d)
    wb.close()
    return out


def import_articles(xlsx: Path, *, with_placeholders: bool) -> int:
    rows = load_rows(xlsx)
    by_slug_figures: dict[str, list[tuple[str, str]]] = {}
    for slug, key, alt in PLACEHOLDERS:
        by_slug_figures.setdefault(slug, []).append((key, alt))

    now = datetime.now(UTC).replace(tzinfo=None)
    ROOT.mkdir(parents=True, exist_ok=True)
    db = SessionLocal()
    try:
        for d in rows:
            slug = str(d["slug"]).strip()
            title = str(d.get("title_vi") or "").strip()
            level = str(d.get("level") or "beginner").strip()
            content = str(d.get("content_md") or "").strip() + "\n"
            if with_placeholders and slug in by_slug_figures:
                content = _inject_placeholders(content, by_slug_figures[slug])
                if not content.endswith("\n"):
                    content += "\n"
            else:
                content = strip_knowledge_figures(content)
            content = strip_image_suggestion_blocks(content)
            if not content.endswith("\n"):
                content += "\n"

            md_path = ROOT / f"{slug}.md"
            md_path.write_text(content, encoding="utf-8")

            sort_order = int(d.get("sort_order") or 0)
            read_time = d.get("read_time_min")
            read_time_min = int(read_time) if read_time is not None and read_time != "" else None
            seo_title = str(d.get("seo_title") or "").strip() or None
            seo_description = str(d.get("seo_description") or "").strip() or None
            is_published = bool(d.get("is_published")) if d.get("is_published") is not None else True
            published_at = _parse_published_at(d.get("published_at")) or now

            row = db.scalar(select(KnowledgeArticle).where(KnowledgeArticle.slug == slug))
            if row is None:
                row = KnowledgeArticle(slug=slug, is_published=is_published, published_at=published_at)
                db.add(row)
            row.title_vi = title
            row.content_md = content
            row.level = level
            row.sort_order = sort_order
            row.read_time_min = read_time_min
            row.seo_title = seo_title
            row.seo_description = seo_description
            row.is_published = is_published
            if row.published_at is None:
                row.published_at = published_at
            print(f"import {slug}")
        db.commit()
        return len(rows)
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Import knowledge articles from Excel")
    parser.add_argument(
        "xlsx",
        nargs="?",
        default=str(DEFAULT_XLSX),
        help="Path to knowledge articles xlsx",
    )
    parser.add_argument(
        "--with-placeholders",
        action="store_true",
        help="Inject pending: image figure markers into content",
    )
    args = parser.parse_args()
    path = Path(args.xlsx)
    if not path.is_file():
        raise SystemExit(f"File not found: {path}")
    n = import_articles(path, with_placeholders=args.with_placeholders)
    print(f"Imported {n} articles from {path}")


if __name__ == "__main__":
    main()
