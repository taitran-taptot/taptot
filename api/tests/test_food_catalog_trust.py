"""Raw catalog rows must agree with their own per-100g figures."""

import json
from pathlib import Path

from app.services.food_catalog_trust import audit_ingredients, ingredient_issues

ROOT = Path(__file__).resolve().parents[2]


def test_piece_stored_as_100g_is_flagged():
    issues = ingredient_issues(
        {
            "slug": "trung-gia",
            "food_kind": "ingredient",
            "source_ref": "usda:example",
            "serving_size": "1 quả",
            "serving_grams": 100,
            "kcal_100g": 143,
            "calories": 143,
            "protein_100g": 12,
            "protein_g": 12,
            "carbs_100g": 1,
            "fat_100g": 10,
        }
    )
    assert "piece_as_100g" in issues


def test_real_egg_portion_is_not_piece_as_100g():
    catalog = json.loads((ROOT / "seeds" / "foods_catalog_v2.json").read_text(encoding="utf-8"))
    egg = next(row for row in catalog if row["slug"] == "trung-ga-ca-qua-song")
    issues = ingredient_issues(egg)
    assert "piece_as_100g" not in issues
    assert "serving_kcal_drift" not in issues


def test_chicken_breast_reference_row_is_internally_consistent():
    catalog = json.loads((ROOT / "seeds" / "foods_catalog_v2.json").read_text(encoding="utf-8"))
    row = next(item for item in catalog if item["slug"] == "uc-ga-khong-da-song")
    assert row["kcal_100g"] == 110
    assert row["source_ref"]
    assert ingredient_issues(row) == []


def test_catalog_audit_runs_and_chicken_stays_clean():
    rows: list[dict] = []
    for name in ("foods_catalog_v2.json",):
        payload = json.loads((ROOT / "seeds" / name).read_text(encoding="utf-8"))
        rows.extend(payload if isinstance(payload, list) else payload.get("foods") or [])
    pantry = json.loads((ROOT / "seeds" / "foods_cooking_pantry.json").read_text(encoding="utf-8"))
    rows.extend(pantry["foods"])
    found = audit_ingredients(rows)
    assert "uc-ga-khong-da-song" not in found
    assert "trung-ga-ca-qua-song" not in found


def test_same_food_rows_match_cited_vn_table():
    """Rows that repeated the cua-đồng pattern stay on the cited table."""
    from openpyxl import load_workbook

    wb = load_workbook(ROOT / "seeds" / "foods_catalog_perfect_v1.xlsx", read_only=True, data_only=True)
    ws = wb["Foods_Catalog_Master"]
    expected = {
        "hen-song-trung-truc": (45, 4.5, 5.1, 0.7, "vn-fct:hen-tuoi"),
        "luon-dong": (130, 18.4, 10.7, 1.5, "vn-fct:8038"),
        "bap-gio-heo-chan-gio-truoc": (230, 15.7, 0, 18.6, "vn-fct:chan-gio-lon-tuoi"),
        "gia-do-xanh": (52, 5.5, 7.17, 0.13, "vn-fct:gia-dau-xanh-tuoi"),
        "thit-lon-ba-chi-song": (260, 16.5, 0, 21.5, "vn-fct:7018"),
    }
    found = {}
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or row[1] not in expected:
            continue
        found[row[1]] = (row[14], row[15], row[16], row[17], row[24])
    assert found == expected


def test_boneless_skinless_chicken_thigh_removed_from_perfect_master():
    from openpyxl import load_workbook

    wb = load_workbook(ROOT / "seeds" / "foods_catalog_perfect_v1.xlsx", read_only=True, data_only=True)
    ws = wb["Foods_Catalog_Master"]
    slugs = {row[1] for i, row in enumerate(ws.iter_rows(values_only=True)) if i and row and row[1]}
    wb.close()
    assert "ma-dui-ga-rut-xuong-bo-da" not in slugs
    assert "ma-dui-ga-con-da" in slugs


def test_cua_dong_matches_vn_food_composition_8034():
    """Bảng thành phần thực phẩm Việt Nam, mã 8034, 100 g phần ăn được."""
    pantry = json.loads((ROOT / "seeds" / "foods_cooking_pantry.json").read_text(encoding="utf-8"))
    row = next(item for item in pantry["foods"] if item["slug"] == "cua-dong")
    assert row["kcal_100g"] == 87
    assert row["protein_100g"] == 12.3
    assert row["carbs_100g"] == 2.0
    assert row["fat_100g"] == 3.3
    assert row["sodium_100mg"] == 1484
    assert row["source_ref"] == "vn-fct:8034"
    assert ingredient_issues(row) == []
